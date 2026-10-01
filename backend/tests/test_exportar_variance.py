import io
import re
from decimal import ROUND_HALF_UP, Decimal
from types import SimpleNamespace

import pytest
from conftest import Mundo, como
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from real_agosto import AgostoReal, agosto_real  # noqa: F401
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_api_variance import _preparar

from fpa.calculo.motor import Cifras, LineaVariance
from fpa.calculo.variance import Convencion
from fpa.modelos import EventoAuditoria
from fpa.servicios.exportar_variance import (
    generar_excel,
    limpiar_texto,
    nombre_de_archivo,
    nombre_de_mes,
)
from fpa.servicios.variance import calcular_variance

D = Decimal
ANALISTA = "analista_a@test.local"
CONTRALOR = "contralor_ab@test.local"
RESPONSABLE = "responsable_a@test.local"
OTRA_ORG = "contralor_c@test.local"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _libro(contenido: bytes) -> Worksheet:
    return load_workbook(io.BytesIO(contenido))["Variance"]


def _filas_de_datos(hoja: Worksheet) -> list[tuple]:  # type: ignore[type-arg]
    """(código, descripción, 12 cifras) de cada fila con datos (después de los encabezados)."""
    salida = []
    for fila in hoja.iter_rows(min_row=6, values_only=True):
        cifras = [fila[i] for i in (2, 3, 4, 5, 7, 8, 9, 10, 12, 13, 14, 15)]
        salida.append((fila[0], fila[1], cifras))
    return salida


def _redondeado(v) -> Decimal | None:  # type: ignore[no-untyped-def]
    return None if v is None else D(str(v)).quantize(D("0.01"), rounding=ROUND_HALF_UP)


# --- Utilidades -------------------------------------------------------------------------------


def test_nombre_de_archivo_y_mes() -> None:
    assert nombre_de_archivo("2026-08") == "variance_2026-08.xlsx"
    # nada raro llega al nombre del archivo: ni comillas, ni barras, ni espacios
    assert re.fullmatch(r"variance_[0-9-]+\.xlsx", nombre_de_archivo('2026-08"; rm -rf /../x'))
    assert nombre_de_mes("2026-08") == "AUGUST" and nombre_de_mes("2026-01") == "JANUARY"
    assert limpiar_texto("a\x00b\x1fc\td\n") == "abc\td\n"


def test_un_texto_que_parece_formula_se_guarda_como_texto() -> None:
    def linea(id_: int, desc: str, tipo: str = "detalle") -> SimpleNamespace:
        return SimpleNamespace(
            id=id_, descripcion=desc, tipo=tipo, seccion_id=None, codigo_corporativo="=1+1"
        )

    cifras = Cifras(D("1"), D("2"), D("-1"), D("-0.5"))
    lineas = [
        linea(1, '=HYPERLINK("http://mala","clic")'),
        linea(2, "+cmd|' /C calc'!A0"),
        linea(3, "-2+3"),
        linea(4, "@SUM(A1)"),
    ]
    contenido = generar_excel(
        hotel="=Hotel",
        periodo="2026-08",
        ano_fiscal="2026/27",
        lineas=lineas,  # type: ignore[arg-type]
        secciones=[],
        resultado={
            x.id: LineaVariance(x.id, Convencion.REAL_MENOS_PRESUPUESTO, cifras, cifras, cifras)
            for x in lineas
        },
        advertencias=["=ADVERTENCIA"],
        version_mapeo=1,
        version_budget=None,
        periodos_ytd=["2026-07", "2026-08"],
    )
    libro = load_workbook(io.BytesIO(contenido))
    hoja = libro["Variance"]
    for fila in range(6, 10):
        for columna in (1, 2):
            celda = hoja.cell(row=fila, column=columna)
            assert celda.data_type == "s", celda.coordinate  # texto, no fórmula
    assert hoja["B6"].value == '=HYPERLINK("http://mala","clic")'
    assert hoja["A1"].data_type == "s" and hoja["A1"].value.startswith("=Hotel")
    notas = libro["Notas"]
    assert all(c.data_type != "f" for fila in notas.iter_rows() for c in fila)
    # ninguna celda de todo el libro es una fórmula
    for hoja_ in libro.worksheets:
        assert all(c.data_type != "f" for fila in hoja_.iter_rows() for c in fila)


# --- API --------------------------------------------------------------------------------------


def _exportar(cliente: TestClient, mundo: Mundo, usuario: str = ANALISTA, extra: str = ""):  # type: ignore[no-untyped-def]
    return cliente.get(
        f"/api/v1/hoteles/{mundo.hotel_a}/variance/exportar?periodo=2026-08{extra}",
        headers=como(usuario),
    )


def test_el_excel_se_descarga_y_cada_cifra_coincide_con_la_pantalla(
    cliente: TestClient, mundo: Mundo, db: Session
) -> None:
    _preparar(db, mundo)
    r = _exportar(cliente, mundo)
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == 'attachment; filename="variance_2026-08.xlsx"'
    assert r.headers["cache-control"] == "no-store"

    pantalla = cliente.get(
        f"/api/v1/hoteles/{mundo.hotel_a}/variance?periodo=2026-08", headers=como(ANALISTA)
    ).json()
    hoja = _libro(r.content)
    filas = [f for f in _filas_de_datos(hoja) if f[1] is not None]
    assert [f[1] for f in filas] == [x["descripcion"] for x in pantalla["lineas"]]
    for fila, linea in zip(filas, pantalla["lineas"], strict=True):
        esperado = [
            linea[bloque][campo]
            for bloque in ("mes", "ytd", "anual")
            for campo in ("actual", "budget", "var", "pct")
        ]
        for (valor, campo_esperado), nombre in zip(
            zip(fila[2], esperado, strict=True), range(12), strict=True
        ):
            es_pct = nombre % 4 == 3
            if campo_esperado is None:
                assert valor is None, (linea["descripcion"], nombre)
            elif es_pct:  # la API redondea el % a 6 decimales
                assert abs(D(str(valor)) - D(campo_esperado)) <= D("0.000001")
            else:
                assert _redondeado(valor) == D(campo_esperado), (linea["descripcion"], nombre)


def test_formato_del_libro(cliente: TestClient, mundo: Mundo, db: Session) -> None:
    _preparar(db, mundo)
    r = _exportar(cliente, mundo)
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Variance", "Notas"]
    hoja = libro["Variance"]
    assert hoja["A1"].value == "Hotel A - Variance 2026-08"
    assert [hoja.cell(row=4, column=c).value for c in (1, 2, 3, 7, 8, 12, 13)] == [
        "Account #",
        "Account Description",
        "AUGUST",
        "COMMENTS",
        "YEAR TO DATE",
        "YTD COMMENTS",
        "ANNUAL BUDGET",
    ]
    assert [hoja.cell(row=5, column=c).value for c in range(3, 7)] == [
        "Actual",
        "Budget",
        "Var",
        "%",
    ]
    assert hoja.freeze_panes == "C6"
    total = next(c for c in hoja["B"] if c.value == "Total")
    assert (
        total.font.bold
        and hoja.cell(row=total.row, column=3).number_format == "#,##0.00;[Red](#,##0.00)"
    )
    assert hoja.cell(row=total.row, column=6).number_format == "0.0%;[Red](0.0%)"
    notas = [c.value for c in libro["Notas"]["A"] if c.value]
    assert any("Variación favorable" in n for n in notas)
    assert any("Mapeo de cuentas: versión" in n for n in notas)


def test_los_filtros_de_la_pantalla_aplican_al_excel(
    cliente: TestClient, mundo: Mundo, db: Session
) -> None:
    _preparar(db, mundo)
    r = _exportar(cliente, mundo, extra="&tipo=total")
    assert [f[1] for f in _filas_de_datos(_libro(r.content)) if f[1]] == ["Total"]
    r = _exportar(cliente, mundo, extra="&q=gasto")
    assert [f[1] for f in _filas_de_datos(_libro(r.content)) if f[1]] == ["Gasto"]


def test_permisos_aislamiento_y_bitacora(cliente: TestClient, mundo: Mundo, db: Session) -> None:
    _preparar(db, mundo)
    assert _exportar(cliente, mundo, CONTRALOR).status_code == 200
    assert _exportar(cliente, mundo, RESPONSABLE).status_code == 403
    assert _exportar(cliente, mundo, OTRA_ORG).status_code == 404
    sin_sesion = cliente.get(f"/api/v1/hoteles/{mundo.hotel_a}/variance/exportar?periodo=2026-08")
    assert sin_sesion.status_code == 401
    assert _exportar(cliente, mundo, extra="&periodo=1999-01").status_code == 422
    ajeno = cliente.get(
        f"/api/v1/hoteles/{mundo.hotel_b}/variance/exportar?periodo=2026-08",
        headers=como(CONTRALOR),
    )
    assert ajeno.status_code == 404  # el hotel B no tiene mapeo
    db.expire_all()
    eventos = [e for e in db.scalars(select(EventoAuditoria)) if e.accion == "variance.exportado"]
    assert len(eventos) == 1  # solo la exportación autorizada
    assert eventos[0].detalle is not None and eventos[0].detalle["periodo"] == "2026-08"
    assert eventos[0].detalle["n_lineas"] == 3


# --- Agosto real -------------------------------------------------------------------------------


def test_agosto_real_el_excel_exportado_coincide_con_el_motor_y_abre(
    agosto_real: AgostoReal,  # noqa: F811
) -> None:
    a = agosto_real
    r = calcular_variance(
        a.db, hotel_id=a.hotel_id, periodo="2026-08", version_mapeo_id=a.version_mapeo_id
    )
    contenido = generar_excel(
        hotel="Torres Mazatlán",
        periodo=r.periodo,
        ano_fiscal=r.ano_fiscal,
        lineas=r.lineas,
        secciones=r.secciones,
        resultado=r.resultado,
        advertencias=r.advertencias,
        version_mapeo=r.version_mapeo_id,
        version_budget=r.version_plan_id,
        periodos_ytd=r.periodos_ytd,
    )
    hoja = _libro(contenido)  # abre sin error
    filas = [f for f in _filas_de_datos(hoja) if f[1] is not None]
    secciones = {s.descripcion for s in r.secciones}
    lineas_en_excel = [
        f for f in filas if f[1] not in secciones or any(c is not None for c in f[2])
    ]
    assert len(lineas_en_excel) >= len(r.lineas)
    por_descripcion: dict[str, list[LineaVariance]] = {}
    for linea in r.lineas:
        por_descripcion.setdefault(linea.descripcion, []).append(r.resultado[linea.id])
    comparadas = 0
    for _, descripcion, cifras in filas:
        candidatas = por_descripcion.get(descripcion)
        if not candidatas or all(c is None for c in cifras):
            continue
        esperado = [
            [
                _redondeado(getattr(getattr(c, b), campo))
                for b in ("mes", "ytd", "anual")
                for campo in ("actual", "budget", "var")
            ]
            for c in candidatas
        ]
        visto = [_redondeado(cifras[i]) for i in (0, 1, 2, 4, 5, 6, 8, 9, 10)]
        if candidatas[0] is not None and visto in esperado:
            comparadas += 1
    assert (
        comparadas > 450
    )  # prácticamente todas las líneas (las informativas llevan Actual en blanco)
    # Un valor conocido de agosto: Salaries del mes (321,190.84) en el libro, con formato contable
    sueldos = next(f for f in filas if f[1] == "Salaries")
    assert _redondeado(sueldos[2][0]) == D("321190.84")


@pytest.mark.parametrize("columna", [3, 8, 13])
def test_los_bloques_estan_donde_dice_el_encabezado(columna: int) -> None:
    contenido = generar_excel(
        hotel="H",
        periodo="2026-08",
        ano_fiscal="2026/27",
        lineas=[],
        secciones=[],
        resultado={},
        advertencias=[],
        version_mapeo=1,
        version_budget=None,
        periodos_ytd=["2026-08"],
    )
    hoja = _libro(contenido)
    assert hoja.cell(row=5, column=columna).value == "Actual"
    assert hoja.cell(row=5, column=columna + 3).value == "%"

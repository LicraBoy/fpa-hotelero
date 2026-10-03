"""Budget del año fiscal de Mazatlán (`Budget MZT 26-27.xlsx`): lectura, alineación, importación."""

import io
from decimal import Decimal
from pathlib import Path

import pytest
from conftest import Mundo, como
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.orm import Session
from test_mzt_integracion import Libro

from fpa.ingesta.budget_anual_mzt import (
    HOJA,
    alinear_con_lineas,
    es_budget_anual_mzt,
    leer_budget_anual,
    leer_budget_anual_filas,
    periodos_del_titulo,
)
from fpa.modelos.plan import LineaPlan
from fpa.repositorios.plan import RepositorioPlan
from fpa.repositorios.reporte import RepositorioReporte
from fpa.servicios.variance import calcular_variance
from fpa.servicios.versionado_mapeo import aprobar
from tests.mzt_sintetico import crear_libro_mzt

D = Decimal
ANALISTA = "analista_a@test.local"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MESES = ["JUL", "AUG", "SEP", "OCT", "NOV", "DEC", "JAN", "FEB", "MAR", "APR", "MAY", "JUN"]


def _fila(b: object, c: object, valores: list[object] | None, total: object = None) -> list[object]:
    fila: list[object] = [None, b, c]
    fila += list(valores) if valores is not None else [None] * 12
    fila += [total]
    return fila


def _filas_budget() -> list[list[object]]:
    """Mismo diseño que el archivo real: título, encabezado en la fila 7, tipo de cambio en la 8."""
    meses = list(range(1, 13))
    return [
        [None, "FY 26/27 Budget Performance"],
        [None, "TORRES MAZATLAN SA DE CV"],
        [None, "BUDGET"],
        [None],
        [None],
        [None, "Account #", "Account Description"],
        [None, None, None, *MESES, "TOTAL"],
        _fila("E x c h a n g e   R a t e", None, [18] * 12, 18),
        _fila("O t h e r   I n c o m e", None, None),
        _fila("028-0000-43080-000", "Rental Income", [100 + m for m in meses], 1278),
        _fila(
            "028-0111-60010-000", "Holydays", [40 + m for m in meses], 558
        ),  # ortografía distinta
        _fila("999", "Algo que el reporte no tiene", [1] * 12, 12),
    ]


def test_el_titulo_da_el_ano_fiscal_y_los_doce_meses() -> None:
    ano, periodos = periodos_del_titulo("FY 26/27 Budget Performance")
    assert ano == "2026/27" and periodos[0] == "2026-07" and periodos[-1] == "2027-06"
    assert periodos[5] == "2026-12" and periodos[6] == "2027-01" and len(periodos) == 12
    with pytest.raises(ValueError, match="año fiscal"):
        periodos_del_titulo("Budget sin año")
    with pytest.raises(ValueError, match="consecutivos"):
        periodos_del_titulo("FY 26/29 Budget")


def test_lee_los_meses_y_el_total_por_fila() -> None:
    leido = leer_budget_anual_filas(_filas_budget())
    assert leido.ano_fiscal == "2026/27" and len(leido.periodos) == 12
    por_fila = {f.fila: f for f in leido.filas}
    assert por_fila[8].meses[2] == D("18") and por_fila[8].total == D("18")  # el tipo de cambio
    assert por_fila[10].meses[0] == D("101") and por_fila[10].total == D("1278")
    assert 9 not in por_fila  # una fila de sección no trae cifras


def test_sin_encabezado_de_meses_se_rechaza() -> None:
    filas = _filas_budget()
    filas[6] = [None, None, None, *["X"] * 12, "TOTAL"]
    with pytest.raises(ValueError, match="JUL"):
        leer_budget_anual_filas(filas)


def test_se_alinea_por_fila_y_nombre_y_el_tipo_de_cambio_por_nombre() -> None:
    leido = leer_budget_anual_filas(_filas_budget())
    # el reporte: el tipo de cambio está en la fila 9 (el Budget lo trae en la 8)
    lineas = [(9, "Exchange Rate:"), (10, "Rental Income"), (11, "Holidays")]
    # las filas del Budget son 8, 10, 11, 12: la 12 no existe en el reporte
    al = alinear_con_lineas(leido, lineas)
    assert al.por_periodo["2026-09"][9].budget == D("18")  # tipo de cambio de septiembre
    assert al.por_periodo["2026-07"][10].budget == D("101")
    assert al.por_periodo["2027-06"][11].budget == D("52")  # junio: 40 + 12
    assert al.anual == {9: D("18"), 10: D("1278"), 11: D("558")}
    assert any("Holydays" in a for a in al.avisos)  # nombre distinto, alineada por la fila
    assert any("sin línea" in a for a in al.avisos)  # la fila 12 se ignora, con aviso


def test_una_fila_con_otro_nombre_en_la_misma_fila_no_se_alinea() -> None:
    leido = leer_budget_anual_filas(_filas_budget())
    al = alinear_con_lineas(
        leido, [(10, "Housekeeping Salaries")]
    )  # nada que ver con Rental Income
    assert 10 not in al.por_periodo["2026-07"] and 10 not in al.anual
    assert any("sin línea" in a for a in al.avisos)


def _libro_budget(ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = HOJA
    for fila in _filas_budget():
        ws.append(fila)
    wb.save(ruta)
    return ruta


def test_se_reconoce_el_libro_por_su_hoja(tmp_path: Path) -> None:
    assert es_budget_anual_mzt(_libro_budget(tmp_path / "b.xlsx"))
    assert not es_budget_anual_mzt(crear_libro_mzt(tmp_path / "v.xlsx"))
    assert leer_budget_anual(tmp_path / "b.xlsx").ano_fiscal == "2026/27"


# --- Importación por API sobre el libro sintético de Mazatlán -----------------------------------


def _libro_budget_para(tmp_path: Path) -> bytes:
    """Un Budget anual alineado con las líneas del libro sintético de Mazatlán (`Ago`)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = HOJA
    ws.append([None, "FY 26/27 Budget Performance"])
    for _ in range(5):
        ws.append([None])
    ws.append([None, None, None, *MESES, "TOTAL"])
    # el libro sintético: fila 9 tipo de cambio, 11 Rental Income, 12 Laundry, 14 Salaries
    for _ in range(8, 20):
        ws.append([None])
    datos = {
        8: ("E x c h a n g e   R a t e", None, [18] * 12, 18),
        11: ("028-0000-43080-000", "Rental Income", [10 * m for m in range(1, 13)], 780),
        14: ("028-0111-60010-000", "Salaries", [5 * m for m in range(1, 13)], 390),
    }
    for fila, (b, c, valores, total) in datos.items():
        for i, v in enumerate([None, b, c, *valores, total], start=1):
            ws.cell(row=fila, column=i, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_la_api_importa_el_budget_del_ano_y_el_variance_lo_usa(
    cliente: TestClient, mundo: Mundo, db: Session, tmp_path: Path
) -> None:
    libro = Libro(db, mundo, crear_libro_mzt(tmp_path / "mzt.xlsx"))
    assert libro.mapeo is not None
    aprobar(db, version_id=libro.mapeo.id, **libro.kw)  # type: ignore[arg-type]
    db.commit()
    base = f"/api/v1/hoteles/{mundo.hotel_a}"
    contenido = _libro_budget_para(tmp_path)
    r = cliente.post(
        f"{base}/budget/importar",
        files={"archivo": ("Budget MZT 26-27.xlsx", contenido, XLSX)},
        headers=como(ANALISTA),
    )
    assert r.status_code == 201, r.text
    assert r.json()["formato"] == "budget_mzt"
    assert r.json()["version"]["ano_fiscal"] == "2026/27"
    # septiembre (que el Excel de Contraloría no traía) ya tiene Budget y también el anual
    v = calcular_variance(db, hotel_id=mundo.hotel_a, periodo="2026-09")
    renta = next(v.resultado[x.id] for x in v.lineas if x.descripcion == "Rental Income")
    assert renta.mes.budget == D("30")  # septiembre = 3 * 10
    assert renta.anual.budget == D("780")  # el año completo
    assert renta.ytd.budget == D("60")  # julio + agosto + septiembre = 10 + 20 + 30
    assert not any("No hay Budget" in a for a in v.advertencias)
    cambio = next(v.resultado[x.id] for x in v.lineas if x.fila_origen == 9)
    assert cambio.mes.budget == D("18") and cambio.anual.budget == D("18")
    # reimportar el mismo archivo no cambia nada
    r2 = cliente.post(
        f"{base}/budget/importar",
        files={"archivo": ("Budget MZT 26-27.xlsx", contenido, XLSX)},
        headers=como(ANALISTA),
    )
    assert r2.status_code == 200
    plan = RepositorioPlan(db, mundo.hotel_a).vigente("budget", "2026/27")
    assert plan is not None
    estructura = RepositorioReporte(db, mundo.hotel_a).vigente()
    assert estructura is not None
    anuales = {x.linea_id: x.anual for x in db.query(LineaPlan).filter_by(version_plan_id=plan.id)}
    assert D("780") in anuales.values()


def test_la_api_rechaza_un_budget_anual_sin_estructura(
    cliente: TestClient, mundo: Mundo, tmp_path: Path
) -> None:
    r = cliente.post(
        f"/api/v1/hoteles/{mundo.hotel_a}/budget/importar",
        files={"archivo": ("b.xlsx", _libro_budget_para(tmp_path), XLSX)},
        headers=como(ANALISTA),
    )
    assert r.status_code == 409


# --- Archivo real (se salta si no está) ---------------------------------------------------------

_REAL = Path(__file__).resolve().parents[2] / "docs" / "adr" / "Budget MZT 26-27.xlsx"
_ENVIO = (
    Path(__file__).resolve().parents[2]
    / "docs"
    / "adr"
    / "02. MZT August 2026 Variance (Envio).xlsx"
)


@pytest.mark.skipif(not (_REAL.exists() and _ENVIO.exists()), reason="Faltan los Excel reales")
def test_budget_real_coincide_con_julio_y_agosto_del_excel_de_contraloria() -> None:
    from fpa.ingesta.variance_mzt import (
        leer_estructura_mzt,
        leer_hojas_mzt,
        leer_valores_mensuales_mzt,
    )

    hojas = leer_hojas_mzt(_ENVIO)
    estructura = leer_estructura_mzt(hojas)
    referencia = leer_valores_mensuales_mzt(hojas).por_periodo
    al = alinear_con_lineas(
        leer_budget_anual(_REAL), [(x.fila_origen, x.descripcion) for x in estructura.lineas]
    )
    assert al.ano_fiscal == "2026/27" and len(al.por_periodo) == 12
    comparadas = 0
    for periodo in ("2026-07", "2026-08"):
        for fila, esperado in referencia[periodo].items():
            if esperado.budget is None or fila not in al.por_periodo[periodo]:
                continue
            comparadas += 1
            assert abs(al.por_periodo[periodo][fila].budget - esperado.budget) <= D("0.01"), (
                periodo,
                fila,
            )  # type: ignore[operator]
    assert comparadas > 1000  # julio y agosto: 0 diferencias
    assert al.por_periodo["2026-09"][9].budget == D("18")  # el tipo de cambio de septiembre
    assert len(al.anual) > 600

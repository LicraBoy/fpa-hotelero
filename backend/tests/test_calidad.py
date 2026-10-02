import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from conftest import Mundo
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from fpa.calidad.evaluacion import Severidad, evaluar, periodo_del_encabezado
from fpa.ingesta.contrato import Cuenta, Fuente, ResultadoLectura
from fpa.ingesta.contrato import Movimiento as MovimientoLeido
from fpa.ingesta.novohit_con130 import leer_con130, leer_con130_filas
from fpa.ingesta.variance_base_datos import leer_variance_base_datos
from fpa.modelos import CargaArchivo, Movimiento
from fpa.servicios.cargar_archivo import Resolucion, cargar_archivo
from fpa.servicios.cargas import DatosArchivo

D = Decimal
_CARPETA_DOCUMENTOS = Path(
    os.environ.get(
        "FPA_DATOS_REALES", Path.home() / "Desktop" / "PLATAFORMA FP&A E INTELIGENCIA HOTELERA"
    )
)
CON130 = _CARPETA_DOCUMENTOS / "BASE DATOS.xlsx"
VARIANCE = _CARPETA_DOCUMENTOS / "02. VT August 2026 Variance (Envio) (1).xlsx"


def mov(
    fecha: date = date(2026, 9, 10),
    debito: str = "10.00",
    credito: str = "0",
    cuenta: str = "1000001.001",
    asiento: str = "9.2.1",
    fila: int = 1,
    concepto: str = "X",
) -> MovimientoLeido:
    return MovimientoLeido(
        cuenta_codigo=cuenta,
        cuenta_descripcion="D",
        asiento=asiento,
        fecha=fecha,
        concepto=concepto,
        debito=D(debito),
        credito=D(credito),
        fila=fila,
    )


def res(*movs: MovimientoLeido, balance: bool = False, **extra: object) -> ResultadoLectura:
    return ResultadoLectura(
        fuente=Fuente.CON130,
        controla_balance=balance,
        movimientos=list(movs),
        **extra,  # type: ignore[arg-type]
    )


def control(rep, codigo: str):  # type: ignore[no-untyped-def]
    return next((c for c in rep.controles if c.codigo == codigo), None)


# --- Periodo del encabezado -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("texto", "esperado"),
    [
        (
            "Periodo Septiembre 2026 a Septiembre 2026, Cuentas [1000001.001] hasta [6021081.001]",
            (date(2026, 9, 1), date(2026, 9, 30)),
        ),
        ("Periodo Julio 2026 a Agosto 2026", (date(2026, 7, 1), date(2026, 8, 31))),
        ("PERIODO FEBRERO 2028 A FEBRERO 2028", (date(2028, 2, 1), date(2028, 2, 29))),
        (
            "Fecha desde: 01/09/2026 hasta: 30/09/2026, Cuentas [1000001.001] hasta [6021081.001]",
            (date(2026, 9, 1), date(2026, 9, 30)),
        ),
        ("FECHA DESDE 1/8/2026 HASTA 15/9/2026", (date(2026, 8, 1), date(2026, 9, 15))),
        ("Fecha desde: 31/02/2026 hasta: 30/09/2026", None),  # fecha que no existe
        ("Fecha desde: 30/09/2026 hasta: 01/09/2026", None),  # al revés
        ("Periodo Agosto 2026 a Julio 2026", None),  # al revés
        ("Periodo Smarch 2026 a Abril 2026", None),
        ("otra cosa", None),
        (None, None),
    ],
)
def test_interpreta_el_periodo_del_encabezado(texto: str | None, esperado: object) -> None:
    assert periodo_del_encabezado(texto) == esperado


# --- Controles --------------------------------------------------------------------------------


def test_balance_ok_y_error() -> None:
    bien = evaluar(res(mov(debito="10"), mov(debito="0", credito="10", fila=2), balance=True))
    assert control(bien, "balance_debito_credito").severidad == Severidad.OK
    mal = evaluar(res(mov(debito="10.00"), balance=True))
    c = control(mal, "balance_debito_credito")
    assert c.severidad == Severidad.ERROR and "diferencia 10.00" in c.resumen
    assert not mal.aprobado


def test_sin_control_de_balance_no_genera_el_control() -> None:
    rep = evaluar(res(mov(debito="10")))
    assert control(rep, "balance_debito_credito") is None and rep.aprobado


def test_errores_de_lectura_rechazan_y_dan_ejemplos() -> None:
    rep = evaluar(res(mov(), errores=["Fila 9: monto inválido"]))
    assert not rep.aprobado
    assert "Fila 9: monto inválido" in rep.motivos_de_rechazo()


def test_total_por_cuenta_que_no_coincide_es_error_con_ejemplo() -> None:
    cuenta = Cuenta(
        codigo="1000001.001",
        descripcion="B",
        fila=1,
        suma_debito=D("10"),
        total_debito_reporte=D("9"),
        total_credito_reporte=D("0"),
        total_verificado=False,
    )
    rep = evaluar(res(mov(), cuentas={"1000001.001": cuenta}))
    c = control(rep, "totales_por_cuenta")
    assert c.severidad == Severidad.ERROR and "1000001.001" in c.ejemplos[0]


def test_total_por_cuenta_faltante_es_aviso() -> None:
    cuentas = {
        "A": Cuenta("A", "a", 1, total_verificado=True),
        "B": Cuenta("B", "b", 2, total_verificado=None),
    }
    rep = evaluar(res(mov(), cuentas=cuentas))
    assert control(rep, "totales_por_cuenta").severidad == Severidad.AVISO


def test_fechas_fuera_del_periodo_del_encabezado_son_error() -> None:
    rep = evaluar(
        res(
            mov(),
            mov(date(2026, 10, 2), fila=2),
            periodo_texto="Periodo Septiembre 2026 a Septiembre 2026, Cuentas",
        )
    )
    c = control(rep, "fechas_en_periodo_del_encabezado")
    assert c.severidad == Severidad.ERROR and "fila 2" in c.ejemplos[0]


def test_encabezado_ininteligible_es_aviso_no_error() -> None:
    rep = evaluar(res(mov(), periodo_texto="Periodo raro"))
    assert control(rep, "fechas_en_periodo_del_encabezado").severidad == Severidad.AVISO
    assert rep.aprobado


def test_periodos_declarados_por_el_usuario() -> None:
    movs = (mov(date(2026, 8, 1)), mov(date(2026, 9, 1), fila=2))
    assert not evaluar(res(*movs), periodos_esperados={"2026-09"}).aprobado
    assert evaluar(res(*movs), periodos_esperados={"2026-08", "2026-09"}).aprobado


def test_cuentas_nuevas_solo_avisan_y_no_en_la_primera_carga() -> None:
    cuentas = {"A": Cuenta("A", "a", 1), "B": Cuenta("B", "b", 2)}
    con_catalogo = evaluar(res(mov(), cuentas=cuentas), cuentas_conocidas={"A"})
    c = control(con_catalogo, "cuentas_nuevas")
    assert c.severidad == Severidad.AVISO and c.ejemplos == ["B"] and con_catalogo.aprobado
    primera = evaluar(res(mov(), cuentas=cuentas), cuentas_conocidas=set())
    assert control(primera, "cuentas_nuevas") is None
    todas = evaluar(res(mov(), cuentas=cuentas), cuentas_conocidas={"A", "B"})
    assert control(todas, "cuentas_nuevas").severidad == Severidad.OK


def test_filas_repetidas_y_negativos_son_avisos_que_no_rechazan() -> None:
    rep = evaluar(res(mov(), mov(fila=2), mov(debito="-5.00", fila=3)))
    assert control(rep, "filas_repetidas").severidad == Severidad.AVISO
    assert control(rep, "importes_negativos").severidad == Severidad.AVISO
    assert rep.aprobado


def test_debito_y_credito_a_la_vez_avisa() -> None:
    rep = evaluar(res(mov(debito="3", credito="2")))
    assert control(rep, "debito_y_credito_a_la_vez").severidad == Severidad.AVISO


def test_asiento_que_no_corresponde_al_mes_avisa_pero_los_de_inventario_no_aplican() -> None:
    rep = evaluar(res(mov(asiento="8.3.1"), mov(asiento="16102", fila=2)))
    c = control(rep, "asiento_y_mes")
    assert c.severidad == Severidad.AVISO and len(c.ejemplos) == 1
    assert (
        control(evaluar(res(mov(asiento="9.2.1"), mov(asiento="16102", fila=2))), "asiento_y_mes")
        is None
    )


def test_el_reporte_es_serializable_a_json() -> None:
    rep = evaluar(res(mov(), mov(fila=2)))
    datos = json.loads(json.dumps(rep.a_dict()))
    assert datos["aprobado"] is True and datos["fuente"] == "con130"
    assert {c["codigo"] for c in datos["controles"]} >= {"lectura_sin_errores", "filas_repetidas"}


# --- Integración con la carga -----------------------------------------------------------------


def cargar(db: Session, mundo: Mundo, lectura: ResultadoLectura, **kw: object):  # type: ignore[no-untyped-def]
    resultado = cargar_archivo(
        db,
        organizacion_id=mundo.org1,
        hotel_id=mundo.hotel_a,
        usuario_id=None,
        resultado=lectura,
        archivo=DatosArchivo("x.xlsx", "a" * 64),
        **kw,  # type: ignore[arg-type]
    )
    db.commit()
    return resultado


def test_la_carga_guarda_el_reporte_de_calidad_en_la_carga(db: Session, mundo: Mundo) -> None:
    r = cargar(db, mundo, res(mov(), mov(fila=2)))
    assert r.resolucion == Resolucion.CREADA
    controles = r.cargas[0].controles
    assert controles is not None and controles["aprobado"] is True
    assert "filas_repetidas" in {c["codigo"] for c in controles["controles"]}


def test_un_control_en_error_rechaza_sin_guardar_nada(db: Session, mundo: Mundo) -> None:
    r = cargar(db, mundo, res(mov(date(2026, 8, 1))), periodos_esperados={"2026-09"})
    assert r.resolucion == Resolucion.RECHAZADA
    assert any("fuera de los periodos declarados" in m for m in r.motivos)
    assert db.scalars(select(CargaArchivo)).first() is None
    assert db.scalars(select(Movimiento)).first() is None


# --- Con los archivos reales de Torres ---------------------------------------------------------


@pytest.mark.skipif(not CON130.exists(), reason="BASE DATOS.xlsx no está en la carpeta")
def test_con130_real_pasa_con_los_avisos_conocidos() -> None:
    rep = evaluar(leer_con130(CON130), cuentas_conocidas={"1000001.001"})
    assert rep.aprobado
    assert control(rep, "balance_debito_credito").severidad == Severidad.OK
    assert control(rep, "totales_por_cuenta").severidad == Severidad.OK
    assert control(rep, "fechas_en_periodo_del_encabezado").severidad == Severidad.OK
    assert control(rep, "filas_repetidas").resumen.startswith("14 ")
    assert control(rep, "importes_negativos").resumen.startswith("257 ")
    assert control(rep, "asiento_y_mes") is None  # 7,242 de 7,242 coinciden


@pytest.mark.skipif(not VARIANCE.exists(), reason="El Excel del Variance no está en la carpeta")
def test_variance_real_pasa_con_los_avisos_conocidos() -> None:
    rep = evaluar(leer_variance_base_datos(VARIANCE), periodos_esperados={"2026-07", "2026-08"})
    assert rep.aprobado
    assert control(rep, "filas_repetidas").resumen.startswith("5 ")
    assert control(rep, "importes_negativos").resumen.startswith("92 ")
    assert any("4389" in e for e in control(rep, "avisos_de_lectura").ejemplos)
    assert control(rep, "asiento_y_mes") is None


@pytest.mark.skipif(not VARIANCE.exists(), reason="El Excel del Variance no está en la carpeta")
def test_variance_declarado_solo_como_agosto_se_rechaza() -> None:
    rep = evaluar(leer_variance_base_datos(VARIANCE), periodos_esperados={"2026-08"})
    assert not rep.aprobado
    assert "fuera de los periodos declarados" in rep.errores[0].resumen


def _filas_reales_con130() -> list[tuple[object, ...]]:
    wb = load_workbook(CON130, read_only=True, data_only=True)
    try:
        return [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _primer_movimiento(filas: list[tuple[object, ...]]) -> int:
    # columnas del archivo real: G=# ASIENTO (6), H=Fecha (7), O=Débito (14)
    for i, f in enumerate(filas):
        if isinstance(f[7], str) and "/" in f[7] and f[14] is not None:
            return i
    raise AssertionError("no hay movimientos")


@pytest.mark.skipif(not CON130.exists(), reason="BASE DATOS.xlsx no está en la carpeta")
def test_archivo_real_alterado_a_proposito_un_debito_es_rechazado_con_mensaje_claro(
    db: Session, mundo: Mundo
) -> None:
    filas = _filas_reales_con130()
    i = _primer_movimiento(filas)
    original = filas[i]
    alterada = list(original)
    alterada[14] = float(original[14]) + 1.0  # type: ignore[arg-type]
    filas[i] = tuple(alterada)
    lectura = leer_con130_filas(filas)
    r = cargar(db, mundo, lectura)
    assert r.resolucion == Resolucion.RECHAZADA
    texto = " | ".join(r.motivos)
    assert "no coincide con la suma de sus movimientos" in texto  # total de la cuenta
    assert "distinto de Crédito total" in texto  # balance
    assert db.scalars(select(CargaArchivo)).first() is None


@pytest.mark.skipif(not CON130.exists(), reason="BASE DATOS.xlsx no está en la carpeta")
def test_archivo_real_con_una_fecha_fuera_del_periodo_es_rechazado() -> None:
    filas = _filas_reales_con130()
    i = _primer_movimiento(filas)
    alterada = list(filas[i])
    alterada[7] = "15/10/2026"
    filas[i] = tuple(alterada)
    rep = evaluar(leer_con130_filas(filas))
    c = control(rep, "fechas_en_periodo_del_encabezado")
    assert c.severidad == Severidad.ERROR and "2026-09-01 a 2026-09-30" in c.resumen

import io
import os
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from conftest import Mundo, como
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from fpa.calculo.ano_fiscal import etiqueta_ano_fiscal
from fpa.config import settings
from fpa.ingesta.variance_estructura import leer_variance_estructura_filas
from fpa.ingesta.variance_valores import a_decimal, leer_valores_filas, leer_variance_valores
from fpa.modelos import EventoAuditoria
from fpa.repositorios.plan import RepositorioPlan
from fpa.repositorios.reporte import RepositorioReporte
from fpa.servicios.budget import importar_budget
from fpa.servicios.cargar_archivo import Resolucion
from fpa.servicios.cargas import DatosArchivo
from fpa.servicios.estructura_reporte import importar_estructura_reporte

D = Decimal
_CARPETA = Path(
    os.environ.get(
        "FPA_DATOS_REALES", Path.home() / "Desktop" / "PLATAFORMA FP&A E INTELIGENCIA HOTELERA"
    )
)
VARIANCE = _CARPETA / "02. VT August 2026 Variance (Envio) (1).xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ANALISTA = "analista_a@test.local"
CONTRALOR = "contralor_ab@test.local"
RESPONSABLE = "responsable_a@test.local"
OTRA_ORG = "contralor_c@test.local"


@pytest.fixture(autouse=True)
def almacenamiento(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(settings, "storage_dir", str(tmp_path / "storage"))
    return tmp_path / "storage"


# --- Año fiscal -------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("periodo", "mes_inicio", "esperado"),
    [
        ("2026-08", 7, "2026/27"),
        ("2026-07", 7, "2026/27"),
        ("2026-06", 7, "2025/26"),
        ("2027-01", 7, "2026/27"),
        ("2099-12", 7, "2099/00"),
        ("2026-08", 1, "2026"),
    ],
)
def test_etiqueta_ano_fiscal(periodo: str, mes_inicio: int, esperado: str) -> None:
    assert etiqueta_ano_fiscal(periodo, mes_inicio) == esperado


@pytest.mark.parametrize("malo", ["2026-13", "2026-8", "agosto", ""])
def test_periodo_invalido(malo: str) -> None:
    with pytest.raises(ValueError):
        etiqueta_ano_fiscal(malo)


# --- Lector de valores ------------------------------------------------------------------------


def test_a_decimal_no_pasa_por_float_ni_redondea() -> None:
    assert a_decimal(20639.6489456203) == D("20639.6489456203")
    assert a_decimal(18) == D("18")
    assert a_decimal(None) is None and a_decimal("texto") is None and a_decimal(True) is None


def test_leer_valores_por_columna() -> None:
    fila = [None] * 17
    fila[3], fila[4], fila[9], fila[14] = 100, 80.5, 400.25, 1200.125
    valores = leer_valores_filas([[None] * 17, fila])
    assert list(valores) == [2]
    v = valores[2]
    assert (v.mes_actual, v.mes_budget, v.ytd_budget, v.anual_budget) == (
        D("100"),
        D("80.5"),
        D("400.25"),
        D("1200.125"),
    )
    assert v.tiene_budget and v.ytd_actual is None


# --- Servicio ---------------------------------------------------------------------------------


def fila(b=None, c=None, d=None, e=None, j=None, o=None):  # type: ignore[no-untyped-def]
    f: list[Any] = [None] * 20
    f[1], f[2], f[3], f[4], f[9], f[14] = b, c, d, e, j, o
    return tuple(f)


EJEMPLO = [
    fila(),
    fila("Account #", "Account Description", "AUGUST"),
    fila(None, None, "Actual"),
    fila("029-0000-1-000", "Ingreso", 100, 120.123456789, 230.5, 1400),
    fila("029-0000-2-000", "Gasto", 30, 25, 50, 300),
    fila(None, "T o t a l", 70, 95.123456789, 180.5, 1100),
]


def _estructura(db: Session, mundo: Mundo, hotel: int | None = None):  # type: ignore[no-untyped-def]
    est = leer_variance_estructura_filas(EJEMPLO)
    ver = importar_estructura_reporte(
        db,
        organizacion_id=mundo.org1,
        hotel_id=hotel or mundo.hotel_a,
        usuario_id=None,
        resultado=est,
        archivo=DatosArchivo(nombre="v.xlsx", sha256="a" * 64),
    ).version
    assert ver is not None
    return ver


def _importar(db: Session, mundo: Mundo, ver, sha: str = "b" * 64, **kw):  # type: ignore[no-untyped-def]
    return importar_budget(
        db,
        organizacion_id=mundo.org1,
        hotel_id=kw.pop("hotel_id", mundo.hotel_a),
        usuario_id=None,
        version_estructura=ver,
        valores=kw.pop("valores", leer_valores_filas(EJEMPLO)),
        archivo=DatosArchivo(nombre="v.xlsx", sha256=sha),
        periodo_corte="2026-08",
        ano_fiscal="2026/27",
        **kw,
    )


def test_importa_budget_tal_como_viene(db: Session, mundo: Mundo) -> None:
    ver = _estructura(db, mundo)
    res = _importar(db, mundo, ver)
    db.commit()
    assert res.resolucion is Resolucion.CREADA and res.version is not None
    assert (res.version.numero, res.version.estado, res.version.n_lineas) == (1, "vigente", 3)
    lineas = {x.id: x for x in RepositorioReporte(db, mundo.hotel_a).lineas(ver.id)}
    guardado = {
        lineas[x.linea_id].descripcion: (x.mes, x.ytd, x.anual)
        for x in RepositorioPlan(db, mundo.hotel_a).lineas(res.version.id)
    }
    assert guardado["Ingreso"] == (D("120.123457"), D("230.500000"), D("1400.000000"))
    assert guardado["Total"] == (D("95.123457"), D("180.500000"), D("1100.000000"))


def test_mismo_archivo_es_idempotente_y_otro_pide_confirmacion(db: Session, mundo: Mundo) -> None:
    ver = _estructura(db, mundo)
    _importar(db, mundo, ver)
    db.commit()
    assert _importar(db, mundo, ver).resolucion is Resolucion.DUPLICADA
    pide = _importar(db, mundo, ver, sha="c" * 64)
    assert pide.resolucion is Resolucion.REQUIERE_CONFIRMACION
    assert len(RepositorioPlan(db, mundo.hotel_a).versiones()) == 1


def test_reemplazar_deja_historial_y_no_toca_la_anterior(db: Session, mundo: Mundo) -> None:
    ver = _estructura(db, mundo)
    primera = _importar(db, mundo, ver).version
    assert primera is not None
    db.commit()
    antes = [(x.mes, x.ytd, x.anual) for x in RepositorioPlan(db, mundo.hotel_a).lineas(primera.id)]
    otro = [list(f) for f in EJEMPLO]
    otro[3][4] = 999
    res = _importar(
        db,
        mundo,
        ver,
        sha="c" * 64,
        valores=leer_valores_filas([tuple(f) for f in otro]),
        confirmar_reemplazo=True,
    )
    db.commit()
    assert res.resolucion is Resolucion.REEMPLAZADA
    repo = RepositorioPlan(db, mundo.hotel_a)
    assert [(v.numero, v.estado) for v in repo.versiones()] == [(1, "reemplazada"), (2, "vigente")]
    assert [(x.mes, x.ytd, x.anual) for x in repo.lineas(primera.id)] == antes  # inmutable
    acciones = [e.accion for e in db.scalars(select(EventoAuditoria))]
    assert {"budget.creado", "budget.reemplazado"} <= set(acciones)


def test_estructura_distinta_o_sin_budget_se_rechaza(db: Session, mundo: Mundo) -> None:
    ver = _estructura(db, mundo)
    fuera = {99: leer_valores_filas([(None,) * 3, (None,) * 3, (None,) * 3, (None,) * 4 + (5,)])[4]}
    assert _importar(db, mundo, ver, valores=fuera).resolucion is Resolucion.RECHAZADA
    assert _importar(db, mundo, ver, valores={}).resolucion is Resolucion.RECHAZADA
    assert RepositorioPlan(db, mundo.hotel_a).versiones() == []


def test_aislamiento_por_hotel(db: Session, mundo: Mundo) -> None:
    ver = _estructura(db, mundo)
    version = _importar(db, mundo, ver).version
    assert version is not None
    otro = RepositorioPlan(db, mundo.hotel_b)
    assert otro.versiones() == [] and otro.obtener(version.id) is None
    assert otro.lineas(version.id) == [] and otro.vigente("budget", "2026/27") is None
    with pytest.raises(ValueError):
        _importar(db, mundo, ver, hotel_id=mundo.hotel_b)


# --- API --------------------------------------------------------------------------------------


def libro_con_budget() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Ago"
    for r, f in enumerate(EJEMPLO, start=1):
        for c, v in enumerate(f, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    # La estructura la lee de las fórmulas; el total lleva su fórmula en S
    ws.cell(row=6, column=19, value="=S4-S5")
    ws.cell(row=4, column=19, value="=-'4001'!L1")
    ws.cell(row=5, column=19, value="='6001'!L1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def subir(
    cliente: TestClient, mundo: Mundo, usuario: str = ANALISTA, hotel: int | None = None, **datos
):  # type: ignore[no-untyped-def]
    return cliente.post(
        f"/api/v1/hoteles/{hotel or mundo.hotel_a}/budget/importar",
        files={"archivo": ("variance.xlsx", libro_con_budget(), XLSX)},
        data={"periodo_corte": "2026-08", **datos},
        headers=como(usuario),
    )


def importar_reporte(cliente: TestClient, mundo: Mundo) -> None:
    r = cliente.post(
        f"/api/v1/hoteles/{mundo.hotel_a}/reporte/importar",
        files={"archivo": ("variance.xlsx", libro_con_budget(), XLSX)},
        headers=como(ANALISTA),
    )
    assert r.status_code == 201, r.text


def test_api_exige_la_estructura_primero(cliente: TestClient, mundo: Mundo) -> None:
    assert subir(cliente, mundo).status_code == 409


def test_api_importa_lista_y_detalla(cliente: TestClient, mundo: Mundo, db: Session) -> None:
    importar_reporte(cliente, mundo)
    r = subir(cliente, mundo)
    assert r.status_code == 201, r.text
    version = r.json()["version"]
    assert (version["ano_fiscal"], version["periodo_corte"], version["n_lineas"]) == (
        "2026/27",
        "2026-08",
        3,
    )
    assert subir(cliente, mundo).status_code == 200  # mismo archivo
    lista = cliente.get(
        f"/api/v1/hoteles/{mundo.hotel_a}/budget/versiones", headers=como(ANALISTA)
    ).json()
    assert [v["numero"] for v in lista] == [1]
    detalle = cliente.get(
        f"/api/v1/hoteles/{mundo.hotel_a}/budget/versiones/{version['id']}", headers=como(ANALISTA)
    ).json()
    assert sorted(x["mes"] for x in detalle["lineas"]) == ["120.123457", "25.000000", "95.123457"]


def test_api_matriz_de_permisos_y_aislamiento(cliente: TestClient, mundo: Mundo) -> None:
    importar_reporte(cliente, mundo)
    assert subir(cliente, mundo, usuario=RESPONSABLE).status_code == 403
    assert subir(cliente, mundo, usuario=CONTRALOR).status_code == 201
    base = f"/api/v1/hoteles/{mundo.hotel_a}/budget/versiones"
    assert cliente.get(base, headers=como(RESPONSABLE)).status_code == 403
    assert cliente.get(base).status_code == 401
    assert cliente.get(base, headers=como(OTRA_ORG)).status_code == 404
    assert subir(cliente, mundo, usuario=OTRA_ORG).status_code == 404
    ajeno = f"/api/v1/hoteles/{mundo.hotel_b}/budget/versiones/1"
    assert cliente.get(ajeno, headers=como(CONTRALOR)).status_code == 404


def test_api_rechaza_archivos_malos_y_periodo(cliente: TestClient, mundo: Mundo) -> None:
    importar_reporte(cliente, mundo)
    r = cliente.post(
        f"/api/v1/hoteles/{mundo.hotel_a}/budget/importar",
        files={"archivo": ("falso.xlsx", b"no es excel", XLSX)},
        data={"periodo_corte": "2026-08"},
        headers=como(ANALISTA),
    )
    assert r.status_code == 415
    assert (
        subir(
            cliente,
            mundo,
        ).status_code
        == 201
    )
    r = cliente.post(
        f"/api/v1/hoteles/{mundo.hotel_a}/budget/importar",
        files={"archivo": ("v.xlsx", libro_con_budget(), XLSX)},
        data={"periodo_corte": "2026-13"},
        headers=como(ANALISTA),
    )
    assert r.status_code == 422


# --- Excel real -------------------------------------------------------------------------------


@pytest.mark.skipif(not VARIANCE.exists(), reason="No está el Excel real del Variance")
def test_budget_real_identico_al_excel_linea_por_linea(db: Session, mundo: Mundo) -> None:
    from fpa.ingesta.variance_estructura import leer_variance_estructura

    est = leer_variance_estructura(VARIANCE)
    ver = importar_estructura_reporte(
        db,
        organizacion_id=mundo.org1,
        hotel_id=mundo.hotel_a,
        usuario_id=None,
        resultado=est,
        archivo=DatosArchivo(nombre=VARIANCE.name, sha256="r" * 64),
    ).version
    assert ver is not None
    valores = leer_variance_valores(VARIANCE)
    res = _importar(db, mundo, ver, valores=valores, sha="r" * 64)
    db.commit()
    assert res.resolucion is Resolucion.CREADA and res.version is not None
    lineas = {x.id: x for x in RepositorioReporte(db, mundo.hotel_a).lineas(ver.id)}
    guardado = {
        lineas[x.linea_id].fila_origen: x
        for x in RepositorioPlan(db, mundo.hotel_a).lineas(res.version.id)
    }
    esperadas = {f: v for f, v in valores.items() if v.tiene_budget}
    esperadas = {f: v for f, v in esperadas.items() if f in guardado}
    assert len(guardado) > 400
    tolerancia = D("0.000001")  # la columna guarda 6 decimales
    for fila_, v in esperadas.items():
        g = guardado[fila_]
        for nuestro, suyo in (
            (g.mes, v.mes_budget),
            (g.ytd, v.ytd_budget),
            (g.anual, v.anual_budget),
        ):
            if suyo is None:
                assert nuestro is None, fila_
            else:
                assert nuestro is not None and abs(nuestro - suyo) <= tolerancia, (
                    fila_,
                    nuestro,
                    suyo,
                )
    # Un Budget conocido: Other Income de agosto (fila 12)
    assert (guardado[12].mes, guardado[12].anual) == (
        D("20639.648946"),
        D("168928.757564"),
    )

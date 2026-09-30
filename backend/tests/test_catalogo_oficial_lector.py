import collections
import os
import re
from pathlib import Path

import openpyxl
import pytest
from xls_sintetico import a_xls, catalogo_basico, cuenta, encabezado_de_pagina, titulo

from fpa.ingesta.novohit_con020 import leer_con020, leer_con020_filas, limpiar_nombre

_CARPETA_DOCUMENTOS = Path(
    os.environ.get(
        "FPA_DATOS_REALES", Path.home() / "Desktop" / "PLATAFORMA FP&A E INTELIGENCIA HOTELERA"
    )
)
CATALOGO = _CARPETA_DOCUMENTOS / "Catalogo de cuentas.xls"
VARIANCE = _CARPETA_DOCUMENTOS / "02. VT August 2026 Variance (Envio) (1).xlsx"
CON130 = _CARPETA_DOCUMENTOS / "BASE DATOS.xlsx"


def test_normaliza_nombres_con_letras_espaciadas() -> None:
    assert limpiar_nombre("B A N C O S") == "BANCOS"
    assert limpiar_nombre("P A S I V O   C I R C U L A N T E") == "PASIVO CIRCULANTE"
    assert limpiar_nombre("ACTIVO CIRCULANTE") == "ACTIVO CIRCULANTE"
    assert limpiar_nombre("IVA 16 %") == "IVA 16 %"  # no toca nombres normales


def test_lee_cuentas_titulos_y_datos_de_la_empresa() -> None:
    r = leer_con020_filas(catalogo_basico())
    assert r.errores == []
    assert (r.empresa, r.fecha_emision) == ("Torres Mazatlan S.A. de C.V.", "30/09/2026")
    assert [c.codigo for c in r.cuentas] == [
        "4001002",
        "4001002.001",
        "6009001",
        "6009001.001",
        "6009002",
    ]
    bar = r.cuentas[1]
    assert (bar.tipo, bar.grupo_ef, bar.depto_ef, bar.depto_ti) == (
        "Acreedor",
        "INCOME",
        "BAR",
        "CENTRO INGRESO",
    )
    assert bar.cuenta_sat == "401.01 Ventas" and bar.activa and not bar.auxiliar
    assert r.cuentas[0].auxiliar  # 'S' en Aux
    assert r.cuentas[0].grupo_ef is None  # celdas vacías -> None


def test_los_titulos_repetidos_por_pagina_se_ven_una_sola_vez() -> None:
    r = leer_con020_filas(catalogo_basico())
    assert [t.codigo for t in r.titulos] == ["4", "40", "4001", "6", "60", "6009"]
    assert r.titulos[-1].descripcion == "GERENCIA Y RECEPCION"
    assert r.filas_ignoradas > 0  # encabezados de página


def test_titulo_con_letras_espaciadas() -> None:
    filas = [*encabezado_de_pagina(), titulo("1000  B A N C O S", True), cuenta("1000001")]
    assert leer_con020_filas(filas).titulos[0].descripcion == "BANCOS"


def test_corrige_la_enie_que_novohit_exporta_como_yen() -> None:
    r = leer_con020_filas(catalogo_basico())
    assert r.cuentas[-1].descripcion == "SERVICIO A DUEÑOS"


def test_cuenta_inactiva_avisa_pero_se_lee() -> None:
    r = leer_con020_filas(catalogo_basico())
    assert r.cuentas[-1].activa is False
    assert any("6009002" in a and "inactiva" in a for a in r.avisos)


def test_tipo_invalido_es_error() -> None:
    filas = [*catalogo_basico(), cuenta("6009003", "", "X", tipo="Raro")]
    r = leer_con020_filas(filas)
    assert any("6009003" in e and "TIPO" in e for e in r.errores)
    assert not r.es_valido


def test_cuenta_repetida_con_datos_distintos_es_error_e_identica_se_ignora() -> None:
    base = catalogo_basico()
    identica = leer_con020_filas(
        [*base, cuenta("6009002", "", "SERVICIO A DUE¥OS", "Deudor", "", "N")]
    )
    assert identica.errores == [] and len(identica.cuentas) == 5
    distinta = leer_con020_filas([*base, cuenta("6009002", "", "OTRO NOMBRE", "Deudor", "", "N")])
    assert any("repetida con datos distintos" in e for e in distinta.errores)


def test_un_titulo_con_dos_nombres_es_error() -> None:
    filas = [*catalogo_basico(), titulo("6009  OTRO NOMBRE")]
    assert any("dos nombres" in e for e in leer_con020_filas(filas).errores)


def test_sin_encabezado_de_columnas_no_es_un_con020() -> None:
    with pytest.raises(ValueError, match="CON020"):
        leer_con020_filas([["hola", "mundo"], ["1", "2"]])


def test_archivo_que_no_es_xls_da_un_error_claro(tmp_path: Path) -> None:
    ruta = tmp_path / "raro.xls"
    ruta.write_bytes(b"esto no es un libro de Excel")
    with pytest.raises(ValueError, match=r"\.xls"):
        leer_con020(ruta)


def test_lee_un_xls_de_verdad_desde_disco(tmp_path: Path) -> None:
    ruta = tmp_path / "catalogo.xls"
    ruta.write_bytes(a_xls(catalogo_basico()))
    assert len(leer_con020(ruta).cuentas) == 5


# --- Con el archivo real de Torres ----------------------------------------------------------------


@pytest.mark.skipif(not CATALOGO.exists(), reason="El catálogo real no está en la carpeta")
def test_catalogo_real_de_torres() -> None:
    r = leer_con020(CATALOGO)
    assert r.errores == []
    assert len(r.cuentas) == 1238 and len(r.titulos) == 68
    assert collections.Counter(c.tipo for c in r.cuentas) == {"Deudor": 1060, "Acreedor": 178}
    assert sum(1 for c in r.cuentas if not c.activa) == 5
    nombres = {t.codigo: t.descripcion for t in r.titulos}
    assert nombres["1"] == "ACTIVO" and nombres["1000"] == "BANCOS"
    assert nombres["6009"] == "GERENCIA Y RECEPCION" and nombres["6007"] == "SERVICIO A DUEÑOS"
    por_codigo = {c.codigo: c for c in r.cuentas}
    assert por_codigo["1000001.001"].descripcion == "BANBAJIO PS 023240518 0201"
    assert por_codigo["1000001.001"].cuenta_sat == "102.01 Bancos nacionales"
    assert por_codigo["1000001"].auxiliar
    # los textos caben en las columnas de la base de datos
    from fpa.modelos.catalogo import CuentaLocal

    tabla = CuentaLocal.__table__.columns
    for campo, columna in (
        ("descripcion", "descripcion"),
        ("grupo_ef", "grupo_ef"),
        ("depto_ef", "depto_ef"),
        ("depto_ti", "depto_ti"),
        ("cuenta_terceros", "cuenta_terceros"),
        ("cuenta_sat", "cuenta_sat"),
    ):
        limite = tabla[columna].type.length  # type: ignore[attr-defined]
        mayor = max(len(getattr(c, campo) or "") for c in r.cuentas)
        assert mayor <= limite, f"{columna}: {mayor} > {limite}"


@pytest.mark.skipif(
    not (CATALOGO.exists() and CON130.exists()), reason="Faltan los archivos reales de Torres"
)
def test_el_catalogo_real_contiene_todas_las_cuentas_del_con130_de_septiembre() -> None:
    from fpa.ingesta.novohit_con130 import leer_con130

    oficial = {c.codigo for c in leer_con020(CATALOGO).cuentas}
    assert set(leer_con130(CON130).cuentas) <= oficial


@pytest.mark.skipif(
    not (CATALOGO.exists() and VARIANCE.exists()), reason="Faltan los archivos reales de Torres"
)
def test_el_signo_del_variance_es_la_naturaleza_de_la_cuenta() -> None:
    """Validación con el Excel de Contraloría: en las fórmulas del Variance toda cuenta Acreedor se
    resta y toda Deudor se suma. El signo es propiedad de la cuenta (TIPO), no de la línea."""
    tipos = {c.codigo: c.tipo for c in leer_con020(CATALOGO).cuentas}
    hoja = openpyxl.load_workbook(VARIANCE)["Ago"]
    conteo: collections.Counter[tuple[str, str]] = collections.Counter()
    for fila in range(1, hoja.max_row + 1):
        formula = hoja.cell(fila, 19).value  # columna S: total del mes desde las hojas de cuenta
        if isinstance(formula, str) and formula.startswith("="):
            for signo, cuenta_ref in re.findall(r"([+-]?)\s*'([^']+)'!L\d+", formula):
                if cuenta_ref in tipos:
                    conteo[(tipos[cuenta_ref], "resta" if signo == "-" else "suma")] += 1
    assert conteo[("Acreedor", "suma")] == 0 and conteo[("Deudor", "resta")] == 0
    assert conteo[("Acreedor", "resta")] >= 10 and conteo[("Deudor", "suma")] >= 100

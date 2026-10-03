"""El Excel exportado con fórmulas vivas: cada celda recalculada debe igualar al motor."""

import io
from decimal import Decimal
from pathlib import Path

import pytest
from conftest import Mundo, como
from evaluador_formulas import Evaluador, FormulaNoSoportada
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from real_agosto import AgostoReal, agosto_real  # noqa: F401
from sqlalchemy.orm import Session
from test_mzt_integracion import Libro

from fpa.servicios.exportar_variance import generar_excel
from fpa.servicios.exportar_variance_formulas import _suma, numero_de_periodo
from fpa.servicios.variance import ReporteVariance, calcular_variance
from tests.mzt_sintetico import crear_libro_mzt

D = Decimal
ANALISTA = "analista_a@test.local"
COLUMNAS = {"mes": 3, "ytd": 8, "anual": 13}
CAMPOS = ("actual", "budget", "var", "pct")


def _generar(r: ReporteVariance) -> Workbook:
    contenido = generar_excel(
        hotel="H",
        periodo=r.periodo,
        ano_fiscal=r.ano_fiscal,
        lineas=r.lineas,
        secciones=r.secciones,
        resultado=r.resultado,
        capturadas=r.capturadas,
        advertencias=r.advertencias,
        version_mapeo=r.version_mapeo_id,
        version_budget=r.version_plan_id,
        periodos_ytd=r.periodos_ytd,
        insumos=r.insumos,
    )
    return load_workbook(io.BytesIO(contenido))


def _comparar(libro: Workbook, r: ReporteVariance) -> tuple[int, int]:
    """Recalcula cada celda de cifras y la compara con el motor. Devuelve (comparadas, fórmulas)."""
    hoja = libro["Variance"]
    ev = Evaluador(libro)
    # fila de cada línea: la misma cuenta que el exportador (secciones consumen una fila)
    por_seccion = {s.id for s in r.secciones}
    fila, previa, mapa = 6, -1, {}
    for linea in r.lineas:
        if linea.seccion_id != previa:
            previa = linea.seccion_id
            if linea.seccion_id is not None and linea.seccion_id in por_seccion:
                fila += 1
        mapa[linea.id] = fila
        fila += 1
    comparadas = formulas = 0
    for linea in r.lineas:
        fila = mapa[linea.id]
        assert hoja.cell(row=fila, column=2).value == linea.descripcion
        for bloque, base in COLUMNAS.items():
            cifras = getattr(r.resultado[linea.id], bloque)
            for i, campo in enumerate(CAMPOS):
                esperado = getattr(cifras, campo)
                celda = hoja.cell(row=fila, column=base + i)
                if esperado is None:
                    # sin Budget no hay Variación ni %; el Actual de una línea informativa se omite
                    assert celda.value is None, (linea.descripcion, bloque, campo)
                    continue
                if celda.value is None:
                    # las líneas informativas no muestran cifras calculadas (como la pantalla)
                    assert linea.tipo == "informativa", (linea.descripcion, bloque, campo)
                    continue
                if celda.data_type == "f":
                    formulas += 1
                obtenido = ev.numero("Variance", base + i, fila)
                tolerancia = D("0.000001") if campo == "pct" else D("0.005")
                assert abs(obtenido - esperado) <= tolerancia, (
                    linea.descripcion,
                    bloque,
                    campo,
                    obtenido,
                    esperado,
                )
                comparadas += 1
    return comparadas, formulas


# --- Unidad -------------------------------------------------------------------------------------


def test_la_suma_con_signos_y_coeficientes() -> None:
    celdas = {1: "D10", 2: "D11", 3: "D12"}
    assert _suma([(1, 1), (2, -1), (3, D("0.160000"))], celdas) == "D10-D11+0.16*D12"
    assert _suma([(1, -1), (3, -D("0.5"))], celdas) == "-D10-0.5*D12"
    assert _suma([(99, 1)], celdas) == "0"  # un componente que no está en el reporte no suma
    assert numero_de_periodo("2026-09") == 202609


def test_el_evaluador_rechaza_lo_que_no_entiende() -> None:
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja["A1"] = "=VLOOKUP(1,2,3)"
    with pytest.raises(FormulaNoSoportada):
        Evaluador(libro).numero(hoja.title, 1, 1)


# --- Libro sintético de Mazatlán ----------------------------------------------------------------


@pytest.fixture
def libro(db: Session, mundo: Mundo, tmp_path: Path) -> Libro:
    return Libro(db, mundo, crear_libro_mzt(tmp_path / "mzt.xlsx"))


def test_cada_celda_recalculada_es_igual_al_motor(libro: Libro) -> None:
    libro.aprobar_y_cargar_budget()
    r = calcular_variance(
        libro.db, hotel_id=libro.mundo.hotel_a, periodo="2026-08", con_insumos=True
    )
    comparadas, formulas = _comparar(_generar(r), r)
    assert comparadas > 30 and formulas > 25


def test_las_hojas_de_apoyo_traen_las_cuentas_y_el_budget(libro: Libro) -> None:
    libro.aprobar_y_cargar_budget()
    r = calcular_variance(
        libro.db, hotel_id=libro.mundo.hotel_a, periodo="2026-08", con_insumos=True
    )
    wb = _generar(r)
    assert {"Variance", "Notas", "Saldos", "Mapeo", "Budget", "Capturas"} <= set(wb.sheetnames)
    assert wb["Variance"]["B3"].value == 202608 and wb["Variance"]["D3"].value == 202607
    claves = {c.value for c in wb["Saldos"]["A"][1:]}
    assert "K4002011" in claves  # la clave lleva prefijo: Excel no la toma por número
    tipos = {c.data_type for c in wb["Saldos"]["A"][1:]}
    assert tipos == {"s"}
    periodos_budget = {c.value for c in wb["Budget"]["C"][1:]}
    assert {202607, 202608} <= periodos_budget  # el Budget mes a mes


def test_el_tipo_de_cambio_se_promedia_con_formula(libro: Libro) -> None:
    libro.aprobar_y_cargar_budget()
    r = calcular_variance(
        libro.db, hotel_id=libro.mundo.hotel_a, periodo="2026-08", con_insumos=True
    )
    wb = _generar(r)
    hoja = wb["Variance"]
    fila = next(
        f
        for f in range(6, hoja.max_row + 1)
        if str(hoja.cell(row=f, column=2).value).startswith("Exchange")
    )
    assert "AVERAGEIFS" in str(hoja.cell(row=fila, column=8).value)
    ev = Evaluador(wb)
    assert abs(ev.numero("Variance", 8, fila) - (D("17.2175") + D("17.1337")) / 2) < D("0.0001")


# --- API ----------------------------------------------------------------------------------------


def test_la_api_entrega_el_reporte_completo_con_formulas_aunque_haya_filtros(
    cliente: TestClient, mundo: Mundo, libro: Libro
) -> None:
    libro.aprobar_y_cargar_budget()
    url = f"/api/v1/hoteles/{mundo.hotel_a}/variance/exportar?periodo=2026-08"
    completo = cliente.get(url, headers=como(ANALISTA))
    filtrado = cliente.get(url + "&tipo=total&q=zzz", headers=como(ANALISTA))
    assert completo.status_code == filtrado.status_code == 200
    a = load_workbook(io.BytesIO(completo.content))
    b = load_workbook(io.BytesIO(filtrado.content))
    assert a["Variance"].max_row == b["Variance"].max_row  # los filtros no recortan el libro
    assert any(c.data_type == "f" for fila in a["Variance"].iter_rows() for c in fila)


# --- Agosto real (se salta si faltan los Excel) ---------------------------------------------------


def test_agosto_real_cada_celda_recalculada_es_igual_al_motor(
    agosto_real: AgostoReal,  # noqa: F811
) -> None:
    a = agosto_real
    r = calcular_variance(
        a.db,
        hotel_id=a.hotel_id,
        periodo="2026-08",
        version_mapeo_id=a.version_mapeo_id,
        con_insumos=True,
    )
    comparadas, formulas = _comparar(_generar(r), r)
    assert comparadas > 4000 and formulas > 3000

"""Sentido de la Variación por línea y bloque, leído de las fórmulas G, L y Q del Excel."""

import io
from decimal import Decimal
from pathlib import Path

import pytest
from conftest import Mundo, como
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.orm import Session
from test_budget import importar_reporte, libro_con_budget

from fpa.calculo.variance import Convencion, variance, variance_pct
from fpa.datos_reales import archivo_real
from fpa.ingesta.variance_convenciones import (
    HOJA,
    leer_convenciones,
    leer_convenciones_filas,
    sentido_de_formula,
)
from fpa.modelos.reporte import ConvencionLinea
from fpa.repositorios.reporte import RepositorioReporte

D = Decimal
GASTO = Convencion.PRESUPUESTO_MENOS_REAL
INGRESO = Convencion.REAL_MENOS_PRESUPUESTO
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ANALISTA = "analista_a@test.local"


@pytest.mark.parametrize(
    ("formula", "esperado"),
    [
        ("=+F12-E12", GASTO),
        ("=F12-E12", GASTO),
        ("=-E12+F12", GASTO),
        ("=+E12-F12", INGRESO),
        ("=-F12+E12", INGRESO),
        ("=+E12+F12", None),  # suma, no resta
        ("=G414+G415", None),  # suma de otras Variaciones
        ("=SUM(G11:G31)", None),
        ("=+F12-E13", None),  # de dos filas distintas
        ("=+F12-E12*2", None),  # algo más que la resta
        ("=+F12-D12", None),  # no son las columnas del bloque
        (12, None),
        (None, None),
    ],
)
def test_el_sentido_sale_de_la_resta_escrita_en_la_formula(
    formula: object, esperado: Convencion | None
) -> None:
    assert sentido_de_formula(formula, "E", "F") == esperado


def test_acumulado_y_anual_se_leen_de_sus_columnas() -> None:
    assert sentido_de_formula("=+K36-J36", "J", "K") == GASTO
    assert sentido_de_formula("=-K9+J9", "J", "K") == INGRESO  # el tipo de cambio acumulado
    assert sentido_de_formula("=+P9-O9", "O", "P") == GASTO


def _fila(g: object = None, acumulado: object = None, anual: object = None) -> tuple[object, ...]:
    fila: list[object] = [None] * 17
    fila[6], fila[11], fila[16] = g, acumulado, anual
    return tuple(fila)


def test_cada_linea_puede_tener_un_sentido_distinto_en_cada_bloque() -> None:
    r = leer_convenciones_filas(
        [
            _fila(),
            _fila("=+F2-E2", "=-K2+J2", "=+P2-O2"),  # el tipo de cambio: el YTD va al revés
            _fila("=+E3-F3", "=+J3-K3"),
            _fila("=SUM(G1:G3)"),  # un total sin sentido propio
        ]
    )
    assert r.por_fila[2] == {"mes": GASTO, "ytd": INGRESO, "anual": GASTO}
    assert r.por_fila[3] == {"mes": INGRESO, "ytd": INGRESO}
    assert 4 not in r.por_fila and 1 not in r.por_fila


def test_una_hoja_sin_formulas_avisa() -> None:
    r = leer_convenciones_filas([_fila(5, 6, 7)])
    assert r.n_lineas == 0 and r.avisos


def test_el_porcentaje_del_excel_equivale_a_la_regla_del_motor() -> None:
    """H = SI(Var=0,0, SI(Actual=0, Var/Budget, SI(Budget=0, Var/Actual, Var/Budget)))."""

    def excel(var: Decimal, actual: Decimal, budget: Decimal) -> Decimal:
        if var == 0:
            return D(0)
        if actual == 0:
            return var / budget
        if budget == 0:
            return var / actual
        return var / budget

    casos = [
        (D("120"), D("100"), INGRESO),
        (D("30"), D("50"), GASTO),
        (D("0"), D("0"), GASTO),
        (D("5"), D("0"), GASTO),  # presupuesto 0: Excel divide entre el Actual
        (D("0"), D("7"), INGRESO),
        (D("40"), D("-20"), GASTO),  # presupuesto negativo
    ]
    for actual, budget, conv in casos:
        var = variance(actual, budget, conv)
        assert var is not None
        assert variance_pct(var, budget) == excel(var, actual, budget), (actual, budget, conv)


# --- Contra el Excel real (se salta si no está) -------------------------------------------------

_REAL = archivo_real("02. MZT August 2026 Variance.xlsx")


@pytest.mark.skipif(not _REAL.exists(), reason="El Excel de la Variación no está disponible")
def test_excel_real_mazatlan() -> None:
    r = leer_convenciones(_REAL)
    assert r.n_lineas >= 690
    assert r.por_fila[12]["mes"] == INGRESO  # Rental Income
    assert r.por_fila[36]["mes"] == GASTO  # Salaries
    assert r.por_fila[9] == {"mes": GASTO, "ytd": INGRESO, "anual": GASTO}  # Exchange Rate


def test_hoja_con_otro_nombre_se_rechaza(tmp_path: Path) -> None:
    wb = Workbook()
    assert wb.active is not None
    wb.active.title = "Otra"
    ruta = tmp_path / "x.xlsx"
    wb.save(ruta)
    with pytest.raises(ValueError, match=HOJA):
        leer_convenciones(ruta)


# --- API ----------------------------------------------------------------------------------------


def _libro_variance_con_g() -> bytes:
    """La hoja de la Variación con las fórmulas de 2 líneas (filas 4 y 5 del libro de ejemplo)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = HOJA
    ws.cell(row=4, column=7, value="=+E4-F4")  # Ingreso: Actual − Presupuesto
    ws.cell(row=5, column=7, value="=+F5-E5")  # Gasto: Presupuesto − Actual
    ws.cell(row=5, column=12, value="=-K5+J5")  # y en el acumulado, al revés
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_api_importa_las_convenciones_y_reimportar_reemplaza(
    cliente: TestClient, mundo: Mundo, db: Session
) -> None:
    importar_reporte(cliente, mundo)
    url = f"/api/v1/hoteles/{mundo.hotel_a}/reporte/convenciones"
    archivo = {"archivo": ("var.xlsx", _libro_variance_con_g(), XLSX)}
    r = cliente.post(url, files=archivo, headers=como(ANALISTA))
    assert r.status_code == 200, r.text
    assert r.json()["celdas"] == 3 and r.json()["lineas"] == 2
    repo = RepositorioReporte(db, mundo.hotel_a)
    estructura = repo.vigente()
    assert estructura is not None
    assert sorted(c.convencion for c in repo.convenciones(estructura.id)) == sorted(
        [INGRESO.value, GASTO.value, INGRESO.value]
    )
    assert cliente.post(url, files=archivo, headers=como(ANALISTA)).status_code == 200
    db.expire_all()
    assert db.query(ConvencionLinea).count() == 3  # reemplaza, no duplica


def test_api_convenciones_permisos_y_errores(cliente: TestClient, mundo: Mundo) -> None:
    url = f"/api/v1/hoteles/{mundo.hotel_a}/reporte/convenciones"
    archivo = {"archivo": ("var.xlsx", _libro_variance_con_g(), XLSX)}
    assert cliente.post(url, files=archivo, headers=como(ANALISTA)).status_code == 409
    importar_reporte(cliente, mundo)
    otro = como("responsable_a@test.local")
    assert cliente.post(url, files=archivo, headers=otro).status_code == 403
    assert cliente.post(url, files=archivo).status_code == 401
    assert (
        cliente.post(url, files=archivo, headers=como("contralor_c@test.local")).status_code == 404
    )
    falso = {"archivo": ("falso.xlsx", b"no es excel", XLSX)}
    assert cliente.post(url, files=falso, headers=como(ANALISTA)).status_code == 415
    sin_hoja = {"archivo": ("otro.xlsx", libro_con_budget(), XLSX)}
    assert cliente.post(url, files=sin_hoja, headers=como(ANALISTA)).status_code == 422

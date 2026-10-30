"""Budget por cuenta, entrega B: cuentas por PARTIDAS (Cant. × Costo, reparto mensual, inflación en
el
Budget).
"""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from conftest import Mundo, como
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.orm import Session
from test_budget_other_income import CONTRALOR, PERIODOS, XLSX, _url

from fpa.ingesta.budget_partidas import leer_hoja_partidas
from fpa.modelos import LineaReporte, VersionEstructuraReporte
from fpa.servicios.presupuesto_ciclo import crear_ano_fiscal, crear_ciclo, crear_periodo_fiscal
from fpa.servicios.presupuesto_partidas import (
    PartidaEntrada,
    calcular_partida,
    costo_extendido,
    distribucion_de,
    repartir_mensual,
)

D = Decimal
EXCEL_REAL = (
    Path(__file__).resolve().parents[2] / "docs" / "-1% - MZT Budget 26-27 TC 18.00 Final.xlsx"
)
CUENTAS = [
    ("Other Fringe", "028-0111-60055-000"),
    ("Telephone", "028-0111-61560-000"),
    ("Training", "028-0111-60110-000"),
    ("Insurance", "028-0111-60655-000"),
]


# --- Cálculo puro -----------------------------------------------------------------------------


def test_el_costo_extendido_es_cantidad_por_costo() -> None:
    assert costo_extendido(D("12"), D("735")) == D("8820")
    assert costo_extendido(D("3"), D("0.333333")) == D("0.999999")


def test_el_reparto_mensual_siempre_cierra_aunque_no_divida_exacto() -> None:
    for total in (D("100"), D("0.000007"), D("1000000.000001"), D("8820")):
        meses = repartir_mensual(total)
        assert len(meses) == 12 and sum(meses, D(0)) == total
    # el residuo va a los primeros meses
    assert repartir_mensual(D("0.000007"))[:7] == [D("0.000001")] * 7


def test_una_partida_estacional_dice_si_sus_meses_cuadran_con_el_extendido() -> None:
    meses = [D(0)] * 12
    meses[3] = meses[4] = meses[5] = D("10500")
    ok = calcular_partida(PartidaEntrada("Regalos", D(3), None, D("10500"), "estacional", meses), 1)
    assert ok.extendido == D("31500") and ok.diferencia == D(0)
    falta = calcular_partida(
        PartidaEntrada("Regalos", D(4), None, D("10500"), "estacional", meses), 1
    )
    assert falta.diferencia == D("-10500")  # faltan 10,500 por distribuir
    suma = distribucion_de(
        [ok, calcular_partida(PartidaEntrada("Tel", D(12), "mos", D(5), "mensual"), 2)]
    )
    assert suma[0] == D(5) and suma[3] == D("10505")


# --- Lector de la hoja de detalle --------------------------------------------------------------


def _hoja_detalle() -> list[tuple[object, ...]]:
    def fila(
        item: object,
        qty: object,
        unidad: object,
        costo: object,
        ext: object,
        marca: dict[int, str],
        meses: list[object],
    ) -> tuple[object, ...]:
        celdas: list[object] = [None] * 26
        celdas[1], celdas[3], celdas[4], celdas[5], celdas[6] = item, qty, unidad, costo, ext
        for col, valor in marca.items():
            celdas[col] = valor
        for i, m in enumerate(meses):
            celdas[12 + i] = m
        return tuple(celdas)

    encabezado = fila("Item", "Qty of Items", "Unit of Measure", "Cost", "Extended Cost", {}, [])
    estacional = [0.0] * 12
    estacional[5] = 3000.0
    inflacion: list[object] = [None] * 26
    inflacion[9] = "Mthly"
    for i in range(12):
        inflacion[12 + i] = 0.05
    return [
        (None,) * 26,
        (None,) * 26,
        (None,) * 26,
        (None,) * 26,
        (None,) * 26,
        (None,) * 26,
        (None,) * 26,
        encabezado,
        fila("LLAMADAS", 12, "mos", 100, 1200, {7: "X"}, [100.0] * 12),
        fila("BRINDIS", 1, "SVC", 3000, 3000, {11: "X"}, estacional),
        fila(None, None, None, None, 0, {}, [0.0] * 12),  # renglón vacío
        fila("Total Expense", None, None, None, 4200, {}, []),
        inflacion,
    ]


def test_el_lector_decide_mensual_o_estacional_por_los_datos_y_lee_la_inflacion() -> None:
    leida = leer_hoja_partidas("0111-61560 Telephone", _hoja_detalle())
    assert [(p.concepto, p.metodo) for p in leida.partidas] == [
        ("LLAMADAS", "mensual"),
        ("BRINDIS", "estacional"),
    ]
    assert leida.partidas[0].cantidad == D(12) and leida.partidas[0].unidad == "mos"
    assert sum(leida.distribucion, D(0)) == D(4200) and leida.total_expense == D(4200)
    assert leida.inflacion == [D("0.05")] * 12 and leida.avisos == []


def test_si_las_partidas_no_suman_el_total_expense_se_avisa() -> None:
    filas = _hoja_detalle()
    mala = list(filas)
    t = list(mala[11])
    t[6] = 9999
    mala[11] = tuple(t)
    leida = leer_hoja_partidas("0111-61560 Telephone", mala)
    assert any("Total Expense" in a for a in leida.avisos)


# --- API ---------------------------------------------------------------------------------------


def _ciclo(db: Session, mundo: Mundo) -> tuple[int, dict[str, int]]:
    version = VersionEstructuraReporte(
        hotel_id=mundo.hotel_a,
        nombre_archivo="e.xlsx",
        sha256="2" * 64,
        estado="vigente",
        n_lineas=6,
        n_detalle=4,
    )
    db.add(version)
    db.flush()
    seccion = LineaReporte(
        version_id=version.id, hotel_id=mundo.hotel_a, orden=1, tipo="seccion", nivel=0,
        descripcion="Mgmt & Check In", clave="mgmt", fila_origen=10,
    )  # fmt: skip
    db.add(seccion)
    db.flush()
    ids = {"seccion": seccion.id}
    for orden, (nombre, codigo) in enumerate(CUENTAS, start=2):
        linea = LineaReporte(
            version_id=version.id, hotel_id=mundo.hotel_a, orden=orden, tipo="detalle", nivel=1,
            seccion_id=seccion.id, codigo_corporativo=codigo, descripcion=nombre,
            clave=nombre.lower(), fila_origen=10 + orden,
        )  # fmt: skip
        db.add(linea)
        db.flush()
        ids[nombre] = linea.id
    ano = crear_ano_fiscal(db, hotel_id=mundo.hotel_a, etiqueta="2026/27", mes_inicio=7, mes_fin=6)
    for periodo in PERIODOS:
        crear_periodo_fiscal(db, ano_fiscal_id=ano.id, periodo=periodo)
    ciclo = crear_ciclo(
        db,
        organizacion_id=mundo.org1,
        hotel_id=mundo.hotel_a,
        usuario_id=None,
        ano_fiscal_id=ano.id,
    )
    db.commit()
    return ciclo.id, ids


def _cuerpo(inflacion_julio: str = "0.05") -> dict[str, object]:
    estacional = ["0"] * 12
    estacional[3] = estacional[4] = estacional[5] = "10500"
    return {
        "partidas": [
            {
                "concepto": "LLAMADAS",
                "cantidad": "12",
                "unidad": "mos",
                "costo": "4515",
                "metodo": "mensual",
            },
            {
                "concepto": "REGALOS",
                "cantidad": "3",
                "unidad": None,
                "costo": "10500",
                "metodo": "estacional",
                "meses": estacional,
            },
        ],
        "inflacion": {p: (inflacion_julio if i == 0 else "0") for i, p in enumerate(PERIODOS)},
    }


def test_guardar_partidas_mensualiza_la_cuenta_con_inflacion_en_el_budget(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    r = cliente.put(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/partidas"),
        json=_cuerpo(),
        headers=como(CONTRALOR),
    )
    assert r.status_code == 200, r.text
    h = r.json()
    assert (
        h["seccion"] == "Telephone"
        and h["padre"] == "Mgmt & Check In"
        and h["linea_id"] == ids["Telephone"]
    )
    assert [p["concepto"] for p in h["partidas"]] == ["LLAMADAS", "REGALOS"]
    assert D(h["partidas"][0]["extendido"]) == D(54180) and D(h["partidas"][1]["extendido"]) == D(
        31500
    )
    assert all(D(p["diferencia"]) == 0 for p in h["partidas"])
    totales = [D(x) for x in h["totales"]]
    assert totales[0] == D(4515) and totales[3] == D(4515 + 10500)  # distribución sin inflación
    # Budget = distribución × (1 + inflación): julio 4515 × 1.05
    assert D(h["total_con_inflacion"][0]) == D("4740.75") and D(h["total_con_inflacion"][1]) == D(
        4515
    )
    # la cuenta quedó mensualizada: lo ve la hoja de la sección y el detalle fijo
    detalle = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/detalle"), headers=como(CONTRALOR)
    ).json()
    assert detalle[0]["metodo"] == "fijo" and D(detalle[0]["monto_fijo"]) == D("4740.75")
    seccion = cliente.get(
        _url(mundo, ciclo_id, f"hojas/{ids['seccion']}"), headers=como(CONTRALOR)
    ).json()
    cuenta = next(c for c in seccion["cuentas"] if c["descripcion"] == "Telephone")
    assert D(cuenta["meses"][0]) == D("4740.75")
    # y se vuelve a leer igual
    g = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/hoja"), headers=como(CONTRALOR)
    ).json()
    assert g["partidas"] == h["partidas"] and g["total_con_inflacion"] == h["total_con_inflacion"]
    captura = next(
        c
        for c in cliente.get(_url(mundo, ciclo_id, "lineas"), headers=como(CONTRALOR)).json()
        if c["linea_id"] == ids["Telephone"]
    )
    assert captura["usa_partidas"] is True


def test_con_partidas_no_se_puede_editar_mes_a_mes(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    cliente.put(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/partidas"),
        json=_cuerpo(),
        headers=como(CONTRALOR),
    )
    celdas = [{"periodo": p, "metodo": "fijo", "monto_fijo": "1"} for p in PERIODOS]
    r = cliente.put(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/detalle"),
        json={"celdas": celdas},
        headers=como(CONTRALOR),
    )
    assert r.status_code == 409 and "partidas" in r.json()["detail"]


def test_cambiar_solo_la_inflacion_recalcula_el_budget(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    cliente.put(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/partidas"),
        json=_cuerpo(),
        headers=como(CONTRALOR),
    )
    r = cliente.put(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/hoja/inflacion"),
        json={"inflacion": dict.fromkeys(PERIODOS, "0.10")},
        headers=como(CONTRALOR),
    )
    assert r.status_code == 200, r.text
    assert D(r.json()["total_con_inflacion"][0]) == D("4966.5")  # 4515 × 1.10
    assert len(r.json()["partidas"]) == 2  # las partidas se conservan


def test_validaciones_de_las_partidas(cliente: TestClient, db: Session, mundo: Mundo) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    url = _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/partidas")
    base = _cuerpo()
    negativo = {
        **base,
        "partidas": [{"concepto": "X", "cantidad": "1", "costo": "-5", "metodo": "mensual"}],
    }
    assert cliente.put(url, json=negativo, headers=como(CONTRALOR)).status_code == 422
    sin_meses = {
        **base,
        "partidas": [{"concepto": "X", "cantidad": "1", "costo": "5", "metodo": "estacional"}],
    }
    r = cliente.put(url, json=sin_meses, headers=como(CONTRALOR))
    assert r.status_code == 422 and "12 meses" in r.text
    assert (
        cliente.put(url, json={**base, "partidas": []}, headers=como(CONTRALOR)).status_code == 422
    )
    incompleta = {**base, "inflacion": {PERIODOS[0]: "0.05"}}
    assert cliente.put(url, json=incompleta, headers=como(CONTRALOR)).status_code == 422
    assert (
        cliente.get(
            _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/hoja"), headers=como(CONTRALOR)
        ).status_code
        == 409
    )


def test_activar_conserva_lo_que_ya_tenia_y_quitar_vuelve_al_editor_mensual(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    celdas = [{"periodo": p, "metodo": "fijo", "monto_fijo": "200"} for p in PERIODOS]
    cliente.put(
        _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/detalle"),
        json={"celdas": celdas},
        headers=como(CONTRALOR),
    )
    r = cliente.post(
        _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/partidas/activar"),
        headers=como(CONTRALOR),
    )
    assert r.status_code == 200, r.text
    h = r.json()
    assert [p["concepto"] for p in h["partidas"]] == ["Importe actual"] and D(
        h["partidas"][0]["extendido"]
    ) == D(2400)
    assert [D(x) for x in h["total_con_inflacion"]] == [D(200)] * 12  # nada se perdió
    assert (
        cliente.delete(
            _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/partidas"), headers=como(CONTRALOR)
        ).status_code
        == 204
    )
    detalle = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/detalle"), headers=como(CONTRALOR)
    ).json()
    assert [D(c["monto_fijo"]) for c in detalle] == [D(200)] * 12
    assert (
        cliente.get(
            _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/hoja"), headers=como(CONTRALOR)
        ).status_code
        == 409
    )


# --- Importación desde el libro del Budget ------------------------------------------------------


def _libro_con_detalle() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Budget 26-27 MZT"
    ws.append([None, "Account #", "Account Description"])
    ws.append([None, "M g m t   &   C h e c k   I n"])
    ws.append([None, "028-0111-61560-000", "Telephone", *([4740.75] + [4515.0] * 11)])
    ws.append([None, "028-0111-60655-000", "Insurance", *([200.0] * 12)])
    ws.append([None, None, "T o t a l   M g m t   &   C h e c k   I n", *(["#REF!"] * 12)])
    for titulo, base in (("EstAct 25-26 MZT", 4000.0), ("MZT Budget 25-26", 3900.0)):
        hoja = wb.create_sheet(titulo)
        hoja.append([None, "M g m t   &   C h e c k   I n"])
        hoja.append([None, "028-0111-61560-000", "Telephone", *([base] * 12)])
        hoja.append([None, None, "T o t a l   M g m t   &   C h e c k   I n", *([base * 2] * 12)])
    detalle = wb.create_sheet("0111-61560 Telephone")
    for fila in _hoja_detalle():
        detalle.append(list(fila))
    # el detalle de la hoja: 12 × 100 + 3,000 estacional; con inflación 5 % no cuadra con la hoja
    # Budget
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def test_importar_mgmt_trae_las_partidas_y_las_referencias_por_cuenta(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    r = cliente.post(
        _url(mundo, ciclo_id, f"detalle/importar-other-income?seccion_id={ids['seccion']}"),
        files={"archivo": ("budget.xlsx", _libro_con_detalle(), XLSX)},
        headers=como(CONTRALOR),
    )
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["importadas"] == 2  # Telephone por partidas e Insurance como monto fijo
    h = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/hoja"), headers=como(CONTRALOR)
    ).json()
    assert [p["concepto"] for p in h["partidas"]] == ["LLAMADAS", "BRINDIS"]
    assert D(h["total_con_inflacion"][5]) == D("3255")  # diciembre: (100 + 3,000) × 1.05
    assert [D(x) for x in h["est_act"]] == [D(4000)] * 12 and [
        D(x) for x in h["budget_previo"]
    ] == [D(3900)] * 12
    assert h["origen"] == "excel"
    # avisa que las partidas no dan lo mismo que la hoja principal del Budget (se importan las
    # partidas)
    assert any("Telephone" in a and "hoja Budget" in a for a in cuerpo["avisos"])
    # una cuenta importada por partidas no se avisa como «sin línea» aunque su fila del Budget
    # exista
    assert not any("«Telephone»" in a and "no tiene línea" in a for a in cuerpo["avisos"])
    ins = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/detalle"), headers=como(CONTRALOR)
    ).json()
    assert [D(c["monto_fijo"]) for c in ins] == [D(200)] * 12


def test_importar_las_partidas_de_una_sola_cuenta_no_toca_las_demas(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    base = f"detalle/importar-other-income?seccion_id={ids['seccion']}"
    libro = {"archivo": ("budget.xlsx", _libro_con_detalle(), XLSX)}
    r = cliente.post(
        _url(mundo, ciclo_id, f"{base}&linea_id={ids['Telephone']}"),
        files=libro,
        headers=como(CONTRALOR),
    )
    assert r.status_code == 200, r.text
    assert r.json()["importadas"] == 1  # solo las partidas de Telephone
    h = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/hoja"), headers=como(CONTRALOR)
    ).json()
    assert [p["concepto"] for p in h["partidas"]] == ["LLAMADAS", "BRINDIS"]
    assert all(x is not None for x in h["est_act"])
    # Insurance (que el libro trae como monto fijo) quedó sin importar
    ins = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Insurance']}/detalle"), headers=como(CONTRALOR)
    ).json()
    assert all(c["monto_fijo"] in (None, "0", "0.000000") or D(c["monto_fijo"]) == 0 for c in ins)
    # una cuenta sin hoja de detalle en el libro: no cambia nada y lo dice
    r = cliente.post(
        _url(mundo, ciclo_id, f"{base}&linea_id={ids['Insurance']}"),
        files=libro,
        headers=como(CONTRALOR),
    )
    assert r.status_code == 200, r.text
    assert r.json()["importadas"] == 0
    assert any("no tiene partidas en el Excel" in a for a in r.json()["avisos"])
    # una línea que no es de la sección se rechaza
    r = cliente.post(
        _url(mundo, ciclo_id, f"{base}&linea_id={ids['seccion']}"),
        files=libro,
        headers=como(CONTRALOR),
    )
    assert r.status_code == 404, r.text


@pytest.mark.skipif(not EXCEL_REAL.exists(), reason="no está el Excel real del Budget en docs/")
def test_el_excel_real_importa_las_cuentas_por_partidas_con_los_totales_de_cada_hoja(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    ciclo_id, ids = _ciclo(db, mundo)
    r = cliente.post(
        _url(mundo, ciclo_id, f"detalle/importar-other-income?seccion_id={ids['seccion']}"),
        files={"archivo": (EXCEL_REAL.name, EXCEL_REAL.read_bytes(), XLSX)},
        headers=como(CONTRALOR),
    )
    assert r.status_code == 200, r.text
    fringe = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Other Fringe']}/hoja"), headers=como(CONTRALOR)
    ).json()
    assert len(fringe["partidas"]) == 16
    assert sum((D(x) for x in fringe["totales"]), D(0)) == D(358050)  # Total Expense de la hoja
    assert D(fringe["total_con_inflacion"][0]) == D("3087") and D(
        fringe["total_con_inflacion"][2]
    ) == D("26239.5")
    assert sum((D(x) for x in fringe["total_con_inflacion"]), D(0)) == D("375952.5")
    telefono = cliente.get(
        _url(mundo, ciclo_id, f"lineas/{ids['Telephone']}/hoja"), headers=como(CONTRALOR)
    ).json()
    assert [D(x) for x in telefono["total_con_inflacion"]] == [
        D(4515)
    ] * 12  # su hoja trae 0 % de inflación
    assert any(x is not None for x in fringe["est_act"])  # trae el Est-Act del año previo

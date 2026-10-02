"""El CON130 se acepta como .xlsx, .xls (el que exporta Novohit) o .csv, con el mismo resultado."""

import csv
import io
from decimal import Decimal
from pathlib import Path

import pytest
import xlwt
from conftest import Mundo, como
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from fpa.config import settings
from fpa.ingesta.novohit_con130 import leer_con130
from fpa.ingesta.tabla_cruda import formato_por_contenido
from fpa.modelos import CargaArchivo
from fpa.seguridad.archivos import FIRMA_OLE2

D = Decimal
ANALISTA = "analista_a@test.local"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# El reporte tal como lo exporta Novohit: encabezado, cuenta con "Saldo anterior", 2 movimientos
FILAS: list[list[object]] = [
    ["Torres Mazatlan S.A. de C.V."],
    ["Mayor Analítico (movimientos por cuentas)"],
    ["Periodo Septiembre 2026 a Septiembre 2026"],
    [
        "# ASIENTO",
        "Fecha",
        "Tp.Asiento",
        "Oper.",
        "Tp.Doc.",
        "Número",
        "Fecha",
        "Concepto",
        "Débito",
        "Crédito",
        "Saldo",
    ],
    ["1000"],
    ["1000001.001  BANCO ÑANDÚ", "", "", "", "", "", "", "Saldo anterior", "", "5.00", ""],
    [
        "9.2.1",
        "01/09/2026",
        "BANCOS/CAJA",
        "REGISTRO",
        "POLIZA",
        "T1",
        "01/09/2026",
        "DEP",
        "1,030.50",
        "",
        "1,035.50",
    ],
    [
        "9.2.2",
        "02/09/2026",
        "BANCOS/CAJA",
        "REGISTRO",
        "POLIZA",
        "T2",
        "02/09/2026",
        "RET",
        "",
        "30.50",
        "1,005.00",
    ],
    ["", "", "", "", "", "", "", "Total  BANCO ÑANDÚ", "1,030.50", "30.50", ""],
    ["4001002.001  INGRESOS DE BAR", "", "", "", "", "", "", "Saldo anterior", "", "", ""],
    [
        "9.2.3",
        "03/09/2026",
        "INGRESOS",
        "REGISTRO",
        "POLIZA",
        "T3",
        "03/09/2026",
        "VENTA",
        "",
        "1,000.00",
        "-1,000.00",
    ],
    ["", "", "", "", "", "", "", "Total  INGRESOS DE BAR", "", "1,000.00", ""],
]


def _vacios_a_none(fila: list[object]) -> list[object]:
    return [None if v == "" else v for v in fila]


def como_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for f in FILAS:
        ws.append(_vacios_a_none(f))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def como_xls() -> bytes:
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("CON130")
    for i, f in enumerate(FILAS):
        for j, v in enumerate(f):
            if v != "":
                hoja.write(i, j, v)
    buf = io.BytesIO()
    libro.save(buf)
    return buf.getvalue()


def como_csv(separador: str = ",", codificacion: str = "utf-8-sig") -> bytes:
    buf = io.StringIO(newline="")
    csv.writer(buf, delimiter=separador).writerows(FILAS)
    return buf.getvalue().encode(codificacion)


def _resumen(ruta: Path):  # type: ignore[no-untyped-def]
    r = leer_con130(ruta)
    return (
        [(m.cuenta_codigo, m.fecha, m.debito, m.credito, m.concepto) for m in r.movimientos],
        list(r.cuentas),
        r.errores,
        r.es_valido,
    )


@pytest.fixture
def esperado(tmp_path: Path):  # type: ignore[no-untyped-def]
    ruta = tmp_path / "base.xlsx"
    ruta.write_bytes(como_xlsx())
    return _resumen(ruta)


def test_el_resultado_base_es_el_esperado(esperado) -> None:  # type: ignore[no-untyped-def]
    movimientos, cuentas, errores, valido = esperado
    assert valido and not errores and cuentas == ["1000001.001", "4001002.001"]
    assert [(m[1].day, m[2], m[3]) for m in movimientos] == [
        (1, D("1030.50"), D("0")),
        (2, D("0"), D("30.50")),
        (3, D("0"), D("1000.00")),
    ]


@pytest.mark.parametrize(
    ("nombre", "contenido"),
    [
        ("libro.xls", como_xls()),
        ("coma.csv", como_csv(",")),
        ("punto_y_coma.csv", como_csv(";")),
        ("tabulador.csv", como_csv("\t")),
        ("barra.csv", como_csv("|")),
        ("windows.csv", como_csv(",", "cp1252")),  # Excel en español guarda así los CSV
    ],
)
def test_cada_formato_da_los_mismos_movimientos(
    tmp_path: Path,
    esperado,
    nombre: str,
    contenido: bytes,  # type: ignore[no-untyped-def]
) -> None:
    ruta = tmp_path / nombre
    ruta.write_bytes(contenido)
    assert _resumen(ruta) == esperado


def test_el_formato_se_reconoce_por_el_contenido_no_por_la_extension(tmp_path: Path) -> None:
    for nombre, contenido, formato in [
        ("a.dat", como_xlsx(), "xlsx"),
        ("b.dat", como_xls(), "xls"),
        ("c.dat", como_csv(), "csv"),
    ]:
        ruta = tmp_path / nombre
        ruta.write_bytes(contenido)
        assert formato_por_contenido(ruta) == formato
    assert como_xls()[:8] == FIRMA_OLE2


def test_un_xls_con_fechas_de_excel_se_lee_como_fechas(tmp_path: Path) -> None:
    from datetime import datetime

    libro = xlwt.Workbook()
    hoja = libro.add_sheet("CON130")
    estilo = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for i, f in enumerate(FILAS[:6]):
        for j, v in enumerate(f):
            if v != "":
                hoja.write(i, j, v)
    hoja.write(6, 0, "9.2.1")
    hoja.write(6, 1, datetime(2026, 9, 1), estilo)  # fecha de verdad, no texto
    hoja.write(6, 7, "DEP")
    hoja.write(6, 8, 10)
    hoja.write(7, 7, "Total  BANCO ÑANDÚ")
    hoja.write(7, 8, 10)
    hoja.write(7, 9, 0)
    buf = io.BytesIO()
    libro.save(buf)
    ruta = tmp_path / "fechas.xls"
    ruta.write_bytes(buf.getvalue())
    movimientos = leer_con130(ruta).movimientos
    assert [(m.fecha.isoformat(), m.debito) for m in movimientos] == [("2026-09-01", D("10"))]


def test_un_csv_binario_o_un_xls_roto_no_se_leen(tmp_path: Path) -> None:
    roto = tmp_path / "roto.xls"
    roto.write_bytes(FIRMA_OLE2 + b"basura")
    with pytest.raises(ValueError, match=r"(?i)no se pudo leer"):
        leer_con130(roto)
    binario = tmp_path / "binario.csv"
    binario.write_bytes(b"a,b\x00c\n1,2\n")
    with pytest.raises(ValueError, match="CSV"):
        leer_con130(binario)


def test_un_csv_sin_el_encabezado_de_novohit_se_rechaza_con_un_mensaje_claro(
    tmp_path: Path,
) -> None:
    ruta = tmp_path / "otro.csv"
    ruta.write_text("a,b,c\n1,2,3\n", encoding="utf-8")
    with pytest.raises(ValueError, match=r"ASIENTO"):
        leer_con130(ruta)


# --- API ----------------------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def almacenamiento(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(settings, "storage_dir", str(tmp_path / "storage"))
    return tmp_path / "storage"


def _subir(
    cliente: TestClient, mundo: Mundo, contenido: bytes, nombre: str, fuente: str = "con130"
):  # type: ignore[no-untyped-def]
    return cliente.post(
        f"/api/v1/hoteles/{mundo.hotel_a}/cargas",
        files={"archivo": (nombre, contenido, XLSX)},
        data={"fuente": fuente},
        headers=como(ANALISTA),
    )


def _cargas(db: Session) -> int:
    db.expire_all()
    return db.scalar(select(func.count()).select_from(CargaArchivo)) or 0


@pytest.mark.parametrize(
    ("nombre", "contenido"),
    [("novohit.xls", como_xls()), ("novohit.csv", como_csv()), ("novohit.xlsx", como_xlsx())],
)
def test_la_api_acepta_el_con130_en_los_tres_formatos(
    cliente: TestClient,
    db: Session,
    mundo: Mundo,
    almacenamiento: Path,
    nombre: str,
    contenido: bytes,
) -> None:
    r = _subir(cliente, mundo, contenido, nombre)
    assert r.status_code == 201, r.text
    assert r.json()["cargas"][0]["periodo"] == "2026-09"
    assert _cargas(db) == 1
    extension = Path(nombre).suffix
    assert list(almacenamiento.rglob(f"*{extension}"))  # se guarda con su extensión
    # subir otra vez el mismo archivo no cambia nada
    assert _subir(cliente, mundo, contenido, nombre).status_code == 200
    assert _cargas(db) == 1


@pytest.mark.parametrize(
    ("nombre", "contenido", "texto"),
    [
        ("falso.xls", b"esto no es un excel", ".xls"),
        ("falso.csv", como_xlsx(), "CSV"),  # un Excel renombrado a .csv
        ("binario.csv", b"a,b\x00\x01\x02", "CSV"),
        ("libro.xlsm", como_xlsx(), "macros"),
        ("libro.xlsb", como_xlsx(), ".xlsb"),
        ("libro.txt", como_csv(), ".xlsx, .xls o .csv"),
        ("libro", como_csv(), ".xlsx, .xls o .csv"),
    ],
)
def test_la_api_rechaza_lo_que_no_es_un_con130_valido(
    cliente: TestClient, db: Session, mundo: Mundo, nombre: str, contenido: bytes, texto: str
) -> None:
    r = _subir(cliente, mundo, contenido, nombre)
    assert r.status_code == 415 and texto in r.json()["detail"], r.text
    assert _cargas(db) == 0


def test_un_xls_con_macros_se_rechaza(cliente: TestClient, db: Session, mundo: Mundo) -> None:
    contenido = como_xls() + "_VBA_PROJECT".encode("utf-16-le")
    r = _subir(cliente, mundo, contenido, "macros.xls")
    assert r.status_code == 415 and "macros" in r.json()["detail"]


def test_un_csv_que_no_es_con130_da_un_error_claro(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    r = _subir(cliente, mundo, b"a,b\n1,2\n", "otro.csv")
    assert r.status_code == 422 and "ASIENTO" in r.json()["detail"]
    assert _cargas(db) == 0


def test_el_excel_del_variance_sigue_siendo_solo_xlsx(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    for nombre in ("variance.xls", "variance.csv"):
        r = _subir(cliente, mundo, como_xls(), nombre, fuente="variance_base_datos")
        assert r.status_code == 415
    assert _cargas(db) == 0


def _con_movimiento_de_cero() -> list[list[object]]:
    """El reporte de prueba más un movimiento de importe cero (Novohit lo deja en blanco)."""
    filas = [list(f) for f in FILAS]
    i = next(n for n, f in enumerate(filas) if f[0] == "9.2.2")
    filas.insert(
        i + 1,
        ["9.3.20", "30/09/2026", "DIARIO", "REGISTRO", "POLIZA", "", "30/09/2026", "AJUSTE EN CERO",
         "", "", "1,005.00"],
    )  # fmt: skip
    return filas


def test_un_movimiento_de_importe_cero_entra_como_aviso_en_xls_y_csv(tmp_path: Path) -> None:
    libro = xlwt.Workbook()
    hoja = libro.add_sheet("CON130")
    for i, f in enumerate(_con_movimiento_de_cero()):
        for j, v in enumerate(f):
            if v != "":
                hoja.write(i, j, v)
    buf = io.BytesIO()
    libro.save(buf)
    texto = io.StringIO(newline="")
    csv.writer(texto, delimiter=";").writerows(_con_movimiento_de_cero())
    for nombre, contenido in [
        ("cero.xls", buf.getvalue()),
        ("cero.csv", texto.getvalue().encode()),
    ]:
        ruta = tmp_path / nombre
        ruta.write_bytes(contenido)
        r = leer_con130(ruta)
        assert r.es_valido and r.errores == [], nombre
        assert len(r.movimientos) == 4 and len(r.avisos) == 1


_REAL = next(
    (
        p
        for p in (
            Path.home()
            / "Desktop"
            / "PLATAFORMA FP&A E INTELIGENCIA HOTELERA"
            / "SEPTIEMBRE 2026.xls",
            Path(__file__).resolve().parents[2] / "docs" / "SEPTIEMBRE 2026.xls",
        )
        if p.exists()
    ),
    None,
)


@pytest.mark.skipif(_REAL is None, reason="SEPTIEMBRE 2026.xls no está disponible")
def test_la_api_carga_el_con130_real_de_septiembre_en_xls(
    cliente: TestClient, db: Session, mundo: Mundo
) -> None:
    assert _REAL is not None
    r = _subir(cliente, mundo, _REAL.read_bytes(), "SEPTIEMBRE 2026.xls")
    assert r.status_code == 201, r.text
    carga = r.json()["cargas"][0]
    assert carga["periodo"] == "2026-09" and carga["n_movimientos"] == 8280
    assert carga["estado"] == "vigente"
    assert _subir(cliente, mundo, _REAL.read_bytes(), "SEPTIEMBRE 2026.xls").status_code == 200

import io
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from conftest import Mundo, como
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from fpa.config import settings
from fpa.ingesta.contrato import Fuente, ResultadoLectura
from fpa.ingesta.contrato import Movimiento as MovimientoLeido
from fpa.modelos import EventoAuditoria
from fpa.servicios.cargar_archivo import cargar_archivo
from fpa.servicios.cargas import DatosArchivo

D = Decimal
_CARPETA = Path(
    os.environ.get(
        "FPA_DATOS_REALES", Path.home() / "Desktop" / "PLATAFORMA FP&A E INTELIGENCIA HOTELERA"
    )
)
VARIANCE = _CARPETA / "02. VT August 2026 Variance (Envio) (1).xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

ANALISTA = "analista_a@test.local"
CONTRALOR = "contralor_ab@test.local"
RESPONSABLE = "responsable_a@test.local"
AUDITORIA_B = "auditoria_b@test.local"
CONTRALOR_C = "contralor_c@test.local"


@pytest.fixture(autouse=True)
def almacenamiento(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(settings, "storage_dir", str(tmp_path / "storage"))
    return tmp_path / "storage"


def libro_variance() -> bytes:
    """Un libro mínimo con la hoja `Ago`: 2 líneas de detalle y un total."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Ago"

    def poner(fila: int, b: Any, c: Any, d: Any, s: Any, t: Any = None) -> None:
        ws.cell(row=fila, column=2, value=b)
        ws.cell(row=fila, column=3, value=c)
        ws.cell(row=fila, column=4, value=d)
        ws.cell(row=fila, column=19, value=s)
        ws.cell(row=fila, column=20, value=t)

    poner(2, "Account #", "Account Description", "AUGUST", None)
    poner(3, None, None, "Actual", None)
    poner(4, "029-0000-1-000", "Ingreso", 100, "=-'4001002.001'!L5")
    poner(5, "029-0000-2-000", "Gasto", 30, "='6009001.001'!L9")
    poner(6, None, "T o t a l", 70, "=S4-S5")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def importar(
    cliente: TestClient,
    mundo: Mundo,
    contenido: bytes | None = None,
    usuario: str = ANALISTA,
    hotel: int | None = None,
    nombre: str = "variance.xlsx",
    **datos: str,
):  # type: ignore[no-untyped-def]
    return cliente.post(
        f"/api/v1/hoteles/{hotel or mundo.hotel_a}/reporte/importar",
        files={"archivo": (nombre, contenido or libro_variance(), XLSX)},
        data=datos,
        headers=como(usuario),
    )


def url(mundo: Mundo, ruta: str, hotel: int | None = None) -> str:
    return f"/api/v1/hoteles/{hotel or mundo.hotel_a}/{ruta}"


def cargar_movimientos(db: Session, mundo: Mundo, cuentas: list[tuple[str, str, str]]) -> None:
    r = ResultadoLectura(fuente=Fuente.VARIANCE_BASE_DATOS, controla_balance=False)
    r.movimientos = [
        MovimientoLeido(
            cuenta_codigo=c,
            cuenta_descripcion=c,
            asiento="8.1.1",
            fecha=date(2026, 8, 5),
            concepto="x",
            debito=D(deb),
            credito=D(cre),
            fila=i,
        )
        for i, (c, deb, cre) in enumerate(cuentas, 1)
    ]
    cargar_archivo(
        db,
        organizacion_id=mundo.org1,
        hotel_id=mundo.hotel_a,
        usuario_id=None,
        resultado=r,
        archivo=DatosArchivo(nombre="m.xlsx", sha256="m" * 64),
    )
    db.commit()


def acciones(db: Session) -> list[str]:
    db.expire_all()
    return [e.accion for e in db.scalars(select(EventoAuditoria).order_by(EventoAuditoria.id))]


# --- Importación ------------------------------------------------------------------------------


def test_importar_crea_estructura_reglas_y_borrador(
    cliente: TestClient, mundo: Mundo, db: Session, almacenamiento: Path
) -> None:
    r = importar(cliente, mundo, periodo_referencia="2026-08")
    assert r.status_code == 201, r.text
    cuerpo = r.json()
    assert (cuerpo["n_lineas"], cuerpo["n_detalle"], cuerpo["n_reglas_suma"]) == (3, 2, 1)
    assert (cuerpo["mapeo_numero"], cuerpo["n_vinculos"], cuerpo["lineas_vinculadas"]) == (1, 2, 2)
    assert list((almacenamiento / "variance" / str(mundo.hotel_a)).glob("*.xlsx"))

    estructura = cliente.get(url(mundo, "reporte/estructura"), headers=como(ANALISTA)).json()
    assert [x["descripcion"] for x in estructura["lineas"]] == ["Ingreso", "Gasto", "Total"]
    versiones = cliente.get(url(mundo, "mapeo/versiones"), headers=como(ANALISTA)).json()
    assert [(v["numero"], v["estado"]) for v in versiones] == [(1, "borrador")]
    detalle = cliente.get(
        url(mundo, f"mapeo/versiones/{versiones[0]['id']}"), headers=como(ANALISTA)
    ).json()
    assert {(v["cuenta_codigo"], v["signo"]) for v in detalle["vinculos"]} == {
        ("4001002.001", -1),
        ("6009001.001", 1),
    }
    assert {"estructura.creada", "reglas_suma.importadas", "mapeo.borrador_creado"} <= set(
        acciones(db)
    )


def test_reimportar_pide_confirmacion_sin_escribir(cliente: TestClient, mundo: Mundo) -> None:
    assert importar(cliente, mundo).status_code == 201
    r = importar(cliente, mundo)  # hay borrador abierto
    assert r.status_code == 409 and r.json()["resolucion"] == "requiere_confirmacion"
    assert len(cliente.get(url(mundo, "mapeo/versiones"), headers=como(ANALISTA)).json()) == 1
    r = importar(cliente, mundo, confirmar_reemplazo="true")
    assert r.status_code == 201 and r.json()["mapeo_numero"] == 2
    estados = [
        v["estado"]
        for v in cliente.get(url(mundo, "mapeo/versiones"), headers=como(ANALISTA)).json()
    ]
    assert estados == ["descartada", "borrador"]


@pytest.mark.parametrize(
    ("nombre", "contenido", "estado"),
    [
        ("libro.xlsm", b"x", 415),
        ("libro.csv", b"a,b", 415),
        ("falso.xlsx", b"esto no es un excel", 415),
    ],
)
def test_archivos_invalidos_se_rechazan_y_se_auditan(
    cliente: TestClient, mundo: Mundo, db: Session, nombre: str, contenido: bytes, estado: int
) -> None:
    r = importar(cliente, mundo, contenido, nombre=nombre)
    assert r.status_code == estado
    assert "estructura.archivo_rechazado" in acciones(db)
    assert cliente.get(url(mundo, "reporte/estructura"), headers=como(ANALISTA)).status_code == 404


def test_libro_sin_hoja_ago_es_422(cliente: TestClient, mundo: Mundo) -> None:
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    assert importar(cliente, mundo, buf.getvalue()).status_code == 422


# --- Edición y aprobación por API -------------------------------------------------------------


def preparar(cliente: TestClient, mundo: Mundo) -> tuple[int, dict[str, int]]:
    assert importar(cliente, mundo).status_code == 201
    version = cliente.get(url(mundo, "mapeo/versiones"), headers=como(ANALISTA)).json()[0]["id"]
    lineas = cliente.get(url(mundo, "reporte/estructura"), headers=como(ANALISTA)).json()["lineas"]
    return version, {x["descripcion"]: x["id"] for x in lineas}


def test_flujo_completo_editar_aprobar_y_bloquear(
    cliente: TestClient, mundo: Mundo, db: Session
) -> None:
    cargar_movimientos(
        db, mundo, [("4001002.001", "0", "100"), ("6009001.001", "30", "0"), ("6099999", "9", "0")]
    )
    version, lineas = preparar(cliente, mundo)
    base = url(mundo, f"mapeo/versiones/{version}")

    v = cliente.get(f"{base}/validacion?periodo=2026-08", headers=como(ANALISTA)).json()
    assert v["aprobable"] is False and [c["cuenta_codigo"] for c in v["sin_mapear"]] == ["6099999"]

    # El analista edita pero no aprueba (403)
    r = cliente.put(
        f"{base}/vinculos",
        json={"linea_id": lineas["Gasto"], "cuenta_codigo": "6099999", "signo": 1},
        headers=como(ANALISTA),
    )
    assert r.status_code == 200 and r.json()["origen"] == "manual"
    assert cliente.post(f"{base}/aprobar", headers=como(ANALISTA)).status_code == 403

    # Un vínculo a una línea que no es de detalle se rechaza (422)
    r = cliente.put(
        f"{base}/vinculos",
        json={"linea_id": lineas["Total"], "cuenta_codigo": "1", "signo": 1},
        headers=como(ANALISTA),
    )
    assert r.status_code == 422
    r = cliente.put(
        f"{base}/vinculos",
        json={"linea_id": lineas["Gasto"], "cuenta_codigo": "1", "signo": 5},
        headers=como(ANALISTA),
    )
    assert r.status_code == 422

    # Quitarlo vuelve a bloquear la aprobación (409 con motivos)
    r = cliente.delete(
        f"{base}/vinculos?linea_id={lineas['Gasto']}&cuenta_codigo=6099999", headers=como(ANALISTA)
    )
    assert r.status_code == 204
    r = cliente.post(f"{base}/aprobar", headers=como(CONTRALOR))
    assert r.status_code == 409 and "no están mapeadas" in r.json()["detail"]["motivos"][0]

    # Excluir a propósito (con motivo) permite aprobar
    r = cliente.put(
        f"{base}/exclusiones",
        json={"cuenta_codigo": "6099999", "motivo": "No operativa"},
        headers=como(ANALISTA),
    )
    assert r.status_code == 200
    r = cliente.post(f"{base}/aprobar", json={"notas": "Primera"}, headers=como(CONTRALOR))
    assert r.status_code == 200 and r.json()["estado"] == "aprobada"

    # Aprobada = inmutable (409) en todas las rutas de edición
    for llamada in (
        cliente.put(
            f"{base}/vinculos",
            json={"linea_id": lineas["Gasto"], "cuenta_codigo": "7", "signo": 1},
            headers=como(ANALISTA),
        ),
        cliente.delete(
            f"{base}/vinculos?linea_id={lineas['Gasto']}&cuenta_codigo=6009001.001",
            headers=como(ANALISTA),
        ),
        cliente.put(
            f"{base}/exclusiones",
            json={"cuenta_codigo": "8", "motivo": "x"},
            headers=como(ANALISTA),
        ),
        cliente.delete(f"{base}/exclusiones?cuenta_codigo=6099999", headers=como(ANALISTA)),
        cliente.post(f"{base}/descartar", headers=como(ANALISTA)),
        cliente.post(f"{base}/aprobar", headers=como(CONTRALOR)),
    ):
        assert llamada.status_code == 409

    # Para cambiar: nuevo borrador copiado de la aprobada
    r = cliente.post(url(mundo, "mapeo/versiones"), headers=como(ANALISTA))
    assert r.status_code == 201 and r.json()["origen"] == "copia de la versión 1"
    assert cliente.post(url(mundo, "mapeo/versiones"), headers=como(ANALISTA)).status_code == 422
    assert {"mapeo.aprobado", "mapeo.editado"} <= set(acciones(db))


def test_descartar_borrador_por_api(cliente: TestClient, mundo: Mundo) -> None:
    version, _ = preparar(cliente, mundo)
    r = cliente.post(url(mundo, f"mapeo/versiones/{version}/descartar"), headers=como(ANALISTA))
    assert r.status_code == 204
    estado = cliente.get(url(mundo, f"mapeo/versiones/{version}"), headers=como(ANALISTA)).json()
    assert estado["estado"] == "descartada"


def test_validacion_exige_periodo_valido(cliente: TestClient, mundo: Mundo) -> None:
    version, _ = preparar(cliente, mundo)
    base = url(mundo, f"mapeo/versiones/{version}/validacion")
    assert cliente.get(base, headers=como(ANALISTA)).status_code == 422
    assert cliente.get(f"{base}?periodo=2026-13", headers=como(ANALISTA)).status_code == 422
    assert cliente.get(f"{base}?periodo=2026-08", headers=como(ANALISTA)).status_code == 200


# --- Matriz rol × endpoint y aislamiento ------------------------------------------------------


def _endpoints(version: int, linea: int) -> list[tuple[str, str, dict[str, Any], str]]:
    """(método, ruta, argumentos, permiso requerido)."""
    v = f"mapeo/versiones/{version}"
    return [
        ("GET", "reporte/estructura", {}, "reporte:leer"),
        ("GET", "mapeo/versiones", {}, "mapeo:leer"),
        ("GET", v, {}, "mapeo:leer"),
        ("GET", f"{v}/validacion?periodo=2026-08", {}, "mapeo:leer"),
        ("POST", "mapeo/versiones", {}, "mapeo:editar"),
        (
            "PUT",
            f"{v}/vinculos",
            {"json": {"linea_id": linea, "cuenta_codigo": "1", "signo": 1}},
            "mapeo:editar",
        ),
        ("DELETE", f"{v}/vinculos?linea_id={linea}&cuenta_codigo=1", {}, "mapeo:editar"),
        (
            "PUT",
            f"{v}/exclusiones",
            {"json": {"cuenta_codigo": "1", "motivo": "m"}},
            "mapeo:editar",
        ),
        ("DELETE", f"{v}/exclusiones?cuenta_codigo=1", {}, "mapeo:editar"),
        ("POST", f"{v}/descartar", {}, "mapeo:editar"),
        ("POST", f"{v}/aprobar", {}, "mapeo:aprobar"),
    ]


PERMISOS_POR_USUARIO = {
    ANALISTA: {"reporte:leer", "mapeo:leer", "mapeo:editar"},
    CONTRALOR: {"reporte:leer", "mapeo:leer", "mapeo:editar", "mapeo:aprobar"},
    RESPONSABLE: {"reporte:leer"},
}


@pytest.mark.parametrize("usuario", list(PERMISOS_POR_USUARIO))
def test_matriz_rol_por_endpoint(cliente: TestClient, mundo: Mundo, usuario: str) -> None:
    version, lineas = preparar(cliente, mundo)
    permitidos = PERMISOS_POR_USUARIO[usuario]
    for metodo, ruta, args, permiso in _endpoints(version, lineas["Gasto"]):
        r = cliente.request(metodo, url(mundo, ruta), headers=como(usuario), **args)
        if permiso in permitidos:
            assert r.status_code != 403, (usuario, metodo, ruta, r.text)
        else:
            assert r.status_code == 403, (usuario, metodo, ruta, r.status_code)
        assert r.status_code != 401


def test_solo_el_contralor_aprueba(cliente: TestClient, mundo: Mundo) -> None:
    version, _ = preparar(cliente, mundo)
    # sin movimientos cargados no hay contra qué validar: 409, pero no 403, para el contralor
    r = cliente.post(url(mundo, f"mapeo/versiones/{version}/aprobar"), headers=como(CONTRALOR))
    assert r.status_code == 409
    r = cliente.post(url(mundo, f"mapeo/versiones/{version}/aprobar"), headers=como(ANALISTA))
    assert r.status_code == 403


def test_sin_sesion_es_401(cliente: TestClient, mundo: Mundo) -> None:
    for ruta in ("reporte/estructura", "mapeo/versiones", "mapeo/versiones/1"):
        assert cliente.get(url(mundo, ruta)).status_code == 401
    r = cliente.post(
        url(mundo, "reporte/importar"), files={"archivo": ("v.xlsx", libro_variance(), XLSX)}
    )
    assert r.status_code == 401


def test_hotel_ajeno_es_404_y_no_se_filtran_datos(cliente: TestClient, mundo: Mundo) -> None:
    version, lineas = preparar(cliente, mundo)
    # Usuario de otra organización: el hotel A no existe para él
    for metodo, ruta, args, _ in _endpoints(version, lineas["Gasto"]):
        r = cliente.request(metodo, url(mundo, ruta), headers=como(CONTRALOR_C), **args)
        assert r.status_code == 404, (metodo, ruta)
    assert importar(cliente, mundo, usuario=CONTRALOR_C).status_code == 404
    # Auditoría solo tiene acceso al hotel B: tampoco ve el A
    assert cliente.get(url(mundo, "mapeo/versiones"), headers=como(AUDITORIA_B)).status_code == 404
    # Una versión del hotel A consultada por el hotel B (donde el contralor sí tiene acceso)
    r = cliente.get(
        url(mundo, f"mapeo/versiones/{version}", hotel=mundo.hotel_b), headers=como(CONTRALOR)
    )
    assert r.status_code == 404
    r = cliente.put(
        url(mundo, f"mapeo/versiones/{version}/vinculos", hotel=mundo.hotel_b),
        json={"linea_id": lineas["Gasto"], "cuenta_codigo": "1", "signo": 1},
        headers=como(CONTRALOR),
    )
    assert r.status_code == 404
    assert (
        cliente.get(
            url(mundo, "mapeo/versiones", hotel=mundo.hotel_b), headers=como(CONTRALOR)
        ).json()
        == []
    )


def test_inactivo_no_entra(cliente: TestClient, mundo: Mundo) -> None:
    r = cliente.get(url(mundo, "mapeo/versiones"), headers=como("inactivo@test.local"))
    assert r.status_code == 401


# --- Excel real --------------------------------------------------------------------------------


@pytest.mark.skipif(not VARIANCE.exists(), reason="No está el Excel real del Variance")
def test_importar_el_variance_real_por_api(cliente: TestClient, mundo: Mundo) -> None:
    r = importar(cliente, mundo, VARIANCE.read_bytes(), nombre=VARIANCE.name)
    assert r.status_code == 201, r.text
    c = r.json()
    assert c["n_detalle"] == 409 and c["n_reglas_suma"] >= 55
    assert c["lineas_vinculadas"] == 263 and c["n_vinculos"] == 304

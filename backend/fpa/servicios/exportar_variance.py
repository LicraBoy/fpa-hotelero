"""Exportación del Variance a Excel, con el mismo orden y bloques que el reporte actual.

Columnas (las del reporte de Contraloría, hoja `Ago`): A Account #, B Account Description,
C-F mes (Actual, Budget, Var, %), G comentarios, H-K acumulado del año fiscal, L comentarios,
M-P presupuesto anual. Las cifras se escriben **exactas** (no redondeadas): el formato de número
las muestra a centavos, igual que la pantalla.

Seguridad: toda descripción viene de un archivo subido por un usuario. Un texto que empieza con
`=`, `+`, `-` o `@` se interpretaría como fórmula al abrirlo en Excel (inyección de fórmulas), así
que **todo texto se escribe como texto**, con tipo `s`, y nunca como fórmula.
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from fpa.calculo.motor import Cifras, LineaVariance
from fpa.modelos.reporte import LineaReporte
from fpa.servicios.variance import SeccionReporte

FORMATO_IMPORTE = "#,##0.00;[Red](#,##0.00)"
FORMATO_PORCENTAJE = "0.0%;[Red](0.0%)"
MESES = [
    "JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
    "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER",
]  # fmt: skip
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_NEGRITA = Font(bold=True)
_RELLENO_ENCABEZADO = PatternFill("solid", fgColor="DDE3EE")
_RELLENO_SECCION = PatternFill("solid", fgColor="E6EEFB")
_RELLENO_TOTAL = PatternFill("solid", fgColor="EEF1F6")
_LINEA = Side(style="thin", color="888888")
_LINEA_GRUESA = Side(style="medium", color="222222")


def limpiar_texto(valor: str) -> str:
    """Quita caracteres de control (no válidos en un libro de Excel)."""
    return _CONTROL.sub("", valor)


def escribir_texto(hoja: Worksheet, fila: int, columna: int, valor: str) -> None:
    """Escribe SIEMPRE como texto: aunque empiece con `=`, no será una fórmula."""
    celda = hoja.cell(row=fila, column=columna)
    celda.value = limpiar_texto(valor)
    celda.data_type = "s"


def nombre_de_mes(periodo: str) -> str:
    return MESES[int(periodo[5:7]) - 1]


def nombre_de_archivo(periodo: str) -> str:
    """`variance_2026-08.xlsx`; el periodo ya viene validado, aun así se limpia."""
    return f"variance_{re.sub(r'[^0-9-]', '', periodo)}.xlsx"


def _cifras(hoja: Worksheet, fila: int, columna: int, c: Cifras, calculada: bool) -> None:
    valores: list[tuple[Decimal | None, str]] = [
        (c.actual if calculada else None, FORMATO_IMPORTE),
        (c.budget, FORMATO_IMPORTE),
        (c.var if calculada else None, FORMATO_IMPORTE),
        (c.pct if calculada else None, FORMATO_PORCENTAJE),
    ]
    for i, (valor, formato) in enumerate(valores):
        celda = hoja.cell(row=fila, column=columna + i)
        if valor is not None:
            celda.value = valor
            celda.number_format = formato
            celda.alignment = Alignment(horizontal="right")


def generar_excel(
    *,
    hotel: str,
    periodo: str,
    ano_fiscal: str,
    lineas: Sequence[LineaReporte],
    secciones: Sequence[SeccionReporte],
    resultado: dict[int, LineaVariance],
    advertencias: Sequence[str],
    capturadas: frozenset[int] = frozenset(),
    version_mapeo: int,
    version_budget: int | None,
    periodos_ytd: Sequence[str],
) -> bytes:
    """`lineas` ya viene filtrada y en el orden del reporte; `secciones` da sus encabezados."""
    libro = Workbook()
    hoja = libro.active
    if hoja is None:  # pragma: no cover - un libro nuevo siempre trae una hoja
        raise RuntimeError("El libro no tiene hoja activa")
    hoja.title = "Variance"

    escribir_texto(hoja, 1, 1, f"{hotel} - Variance {periodo}")
    hoja["A1"].font = Font(bold=True, size=14)
    escribir_texto(
        hoja,
        2,
        1,
        f"Año fiscal {ano_fiscal} · acumulado de {', '.join(periodos_ytd)}",
    )

    # Encabezados: dos filas, como el reporte actual
    encabezado1, encabezado2 = 4, 5
    textos1 = {
        1: "Account #",
        2: "Account Description",
        3: nombre_de_mes(periodo),
        7: "COMMENTS",
        8: "YEAR TO DATE",
        12: "YTD COMMENTS",
        13: "ANNUAL BUDGET",
    }
    for columna, texto in textos1.items():
        escribir_texto(hoja, encabezado1, columna, texto)
    for inicio in (3, 8, 13):
        for i, texto in enumerate(("Actual", "Budget", "Var", "%")):
            escribir_texto(hoja, encabezado2, inicio + i, texto)
            hoja.cell(row=encabezado2, column=inicio + i).alignment = Alignment(horizontal="right")
        hoja.merge_cells(
            start_row=encabezado1, start_column=inicio, end_row=encabezado1, end_column=inicio + 3
        )
        hoja.cell(row=encabezado1, column=inicio).alignment = Alignment(horizontal="center")
    for fila in (encabezado1, encabezado2):
        for columna in range(1, 17):
            celda = hoja.cell(row=fila, column=columna)
            celda.font = _NEGRITA
            celda.fill = _RELLENO_ENCABEZADO
            celda.border = Border(bottom=_LINEA)

    por_seccion = {s.id: s for s in secciones}
    fila = encabezado2 + 1
    ultima_seccion: int | None = -1
    for linea in lineas:
        if linea.seccion_id != ultima_seccion:
            ultima_seccion = linea.seccion_id
            if linea.seccion_id is not None and linea.seccion_id in por_seccion:
                escribir_texto(hoja, fila, 2, por_seccion[linea.seccion_id].descripcion)
                for columna in range(1, 17):
                    celda = hoja.cell(row=fila, column=columna)
                    celda.font = _NEGRITA
                    celda.fill = _RELLENO_SECCION
                fila += 1
        r = resultado[linea.id]
        calculada = linea.tipo != "informativa"
        if linea.codigo_corporativo:
            escribir_texto(hoja, fila, 1, linea.codigo_corporativo)
        escribir_texto(hoja, fila, 2, linea.descripcion)
        _cifras(hoja, fila, 3, r.mes, calculada or linea.id in capturadas)
        _cifras(hoja, fila, 8, r.ytd, calculada)
        _cifras(hoja, fila, 13, r.anual, calculada)
        if linea.tipo in ("total", "subtotal"):
            for columna in range(1, 17):
                celda = hoja.cell(row=fila, column=columna)
                celda.font = _NEGRITA
                celda.fill = _RELLENO_TOTAL
                celda.border = Border(top=_LINEA_GRUESA if linea.tipo == "total" else _LINEA)
        elif linea.tipo == "informativa":
            for columna in range(1, 17):
                hoja.cell(row=fila, column=columna).font = Font(italic=True, color="666666")
        fila += 1

    hoja.freeze_panes = f"C{encabezado2 + 1}"
    hoja.column_dimensions["A"].width = 22
    hoja.column_dimensions["B"].width = 46
    for columna in range(3, 17):
        hoja.column_dimensions[get_column_letter(columna)].width = 15
    hoja.column_dimensions["G"].width = 10
    hoja.column_dimensions["L"].width = 10
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    hoja.print_title_rows = f"{encabezado1}:{encabezado2}"

    # Notas: de dónde salen las cifras y qué avisó el sistema al calcularlas
    notas = libro.create_sheet("Notas")
    renglones = [
        f"Variance de {periodo} · año fiscal {ano_fiscal}",
        f"Mapeo de cuentas: versión {version_mapeo}",
        f"Budget: versión {version_budget}" if version_budget is not None else "Budget: ninguno",
        "Variación favorable = positiva. Ingresos y utilidades: real - presupuesto; gastos: "
        "presupuesto - real.",
        "Las líneas informativas (p. ej. el tipo de cambio) no se calculan; el Actual del mes "
        "solo aparece si fue capturado.",
        f"Generado por la plataforma el {datetime.now().strftime('%Y-%m-%d %H:%M')} (servidor).",
        "",
        "Advertencias al calcular:" if advertencias else "Sin advertencias al calcular.",
        *advertencias,
    ]
    for i, texto in enumerate(renglones, start=1):
        escribir_texto(notas, i, 1, texto)
    notas.column_dimensions["A"].width = 110

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()

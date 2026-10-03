"""Variance exportado con **fórmulas vivas**: el libro se puede auditar de punta a punta.

Hojas (además de `Variance` y `Notas`):

- `Saldos`: Débito − Crédito de cada cuenta mapeada por mes (lo que el motor leyó de la carga).
- `Mapeo`: cada cuenta de cada línea con su signo y su aporte al mes y al acumulado (`SUMIFS` sobre
  `Saldos`).
- `Budget`: el Budget mes a mes de las líneas de detalle (periodo 0 = presupuesto anual).
- `Capturas`: los Actuales escritos a mano (p. ej. el tipo de cambio).

En `Variance` el Actual de una línea es `SUMIFS` sobre `Mapeo` (más sus componentes, con su
coeficiente), los totales suman sus partes, la Variación sigue el sentido que el Excel de
Contraloría trae para esa línea y el % usa la fórmula de su columna H. Solo se escribe una
fórmula donde el motor tiene un valor: lo que el motor deja vacío (sin Budget) queda vacío.

Seguridad: ninguna fórmula lleva texto de usuario; lo único que cambia es el número de fila. Las
cuentas y descripciones se escriben como **texto** (con prefijo `K` en la clave de búsqueda, para
que Excel no confunda `4003001.010` con el número 4003001.01).
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from fpa.calculo.motor import Cifras, LineaVariance
from fpa.calculo.totales import Coeficiente
from fpa.calculo.variance import Convencion
from fpa.modelos.reporte import LineaReporte
from fpa.servicios.exportar_variance import FORMATO_IMPORTE, FORMATO_PORCENTAJE, escribir_texto
from fpa.servicios.variance import InsumosFormulas

BLOQUES = {"mes": 3, "ytd": 8, "anual": 13}  # columna del Actual de cada bloque
FILA_PARAMETROS = 3
CELDA_PERIODO = "Variance!$B$3"
CELDA_INICIO = "Variance!$D$3"


def numero_de_periodo(periodo: str) -> int:
    """`2026-09` -> 202609 (número, para poder comparar con `>=` y `<=` en Excel)."""
    return int(periodo.replace("-", ""))


def _col(columna: int) -> str:
    return get_column_letter(columna)


def _coef(c: Coeficiente) -> str:
    return format(abs(Decimal(c)).normalize(), "f")


def _suma(componentes: Sequence[tuple[int, Coeficiente]], celdas: Mapping[int, str]) -> str:
    """`D12-D13+0.16*D14`: la suma con signo de las celdas de los componentes."""
    terminos = ""
    for linea, c in componentes:
        if linea not in celdas:
            continue
        signo = "-" if c < 0 else "+"
        factor = "" if abs(c) == 1 else f"{_coef(c)}*"
        terminos += f"{signo}{factor}{celdas[linea]}"
    if not terminos:
        return "0"
    return terminos[1:] if terminos[0] == "+" else terminos


class Rangos:
    """Rangos de las hojas auxiliares (con límite exacto: más rápido que columnas enteras)."""

    def __init__(self, n_saldos: int, n_mapeo: int, n_budget: int, n_capturas: int) -> None:
        def rango(hoja: str, col: str, n: int) -> str:
            return f"{hoja}!${col}$2:${col}${max(n, 2)}"

        self.saldo_clave = rango("Saldos", "A", n_saldos)
        self.saldo_periodo = rango("Saldos", "C", n_saldos)
        self.saldo_neto = rango("Saldos", "D", n_saldos)
        self.mapeo_fila = rango("Mapeo", "A", n_mapeo)
        self.mapeo_mes = rango("Mapeo", "G", n_mapeo)
        self.mapeo_acum = rango("Mapeo", "I", n_mapeo)
        self.budget_fila = rango("Budget", "A", n_budget)
        self.budget_periodo = rango("Budget", "C", n_budget)
        self.budget_importe = rango("Budget", "D", n_budget)
        self.cap_fila = rango("Capturas", "A", n_capturas)
        self.cap_periodo = rango("Capturas", "C", n_capturas)
        self.cap_importe = rango("Capturas", "D", n_capturas)


def _acumulado(periodo_rango: str) -> str:
    return f'{periodo_rango},">="&{CELDA_INICIO},{periodo_rango},"<="&{CELDA_PERIODO}'


def _es_total_por_regla(linea_id: int, ins: InsumosFormulas) -> bool:
    return linea_id in ins.reglas and linea_id not in ins.detalle_con_regla


def _encabezados(hoja: Worksheet, textos: Sequence[str]) -> None:
    for i, t in enumerate(textos, 1):
        escribir_texto(hoja, 1, i, t)


def escribir_hojas_auxiliares(
    libro: Workbook,
    lineas: Sequence[LineaReporte],
    filas: Mapping[int, int],
    ins: InsumosFormulas,
) -> Rangos:
    """Crea `Saldos`, `Mapeo`, `Budget` y `Capturas` y devuelve sus rangos."""
    nombres = {x.id: x.descripcion for x in lineas}

    saldos = libro.create_sheet("Saldos")
    _encabezados(saldos, ("Clave", "Cuenta", "Periodo (AAAAMM)", "Neto (Débito − Crédito)"))
    n_saldos = 1
    for (cuenta, p), neto in sorted(ins.saldos.items()):
        n_saldos += 1
        escribir_texto(saldos, n_saldos, 1, f"K{cuenta}")
        escribir_texto(saldos, n_saldos, 2, cuenta)
        saldos.cell(row=n_saldos, column=3, value=numero_de_periodo(p))
        saldos.cell(row=n_saldos, column=4, value=neto).number_format = FORMATO_IMPORTE

    pares = [
        (filas[linea], nombres[linea], cuenta, signo)
        for linea, cs in ins.cuentas.items()
        if linea in filas
        for cuenta, signo in cs
    ]
    pares.sort(key=lambda x: (x[0], x[2]))
    n_mapeo = 1 + len(pares)
    rg = Rangos(n_saldos, n_mapeo, 1, 1)
    mapeo = libro.create_sheet("Mapeo")
    _encabezados(
        mapeo,
        (
            "Fila en Variance",
            "Línea",
            "Cuenta",
            "Signo",
            "Clave",
            "Neto del mes",
            "Aporte al mes",
            "Neto acumulado",
            "Aporte acumulado",
        ),
    )
    for r, (fila, descripcion, cuenta, signo) in enumerate(pares, start=2):
        mapeo.cell(row=r, column=1, value=fila)
        escribir_texto(mapeo, r, 2, descripcion)
        escribir_texto(mapeo, r, 3, cuenta)
        mapeo.cell(row=r, column=4, value=signo)
        escribir_texto(mapeo, r, 5, f"K{cuenta}")
        base = f"SUMIFS({rg.saldo_neto},{rg.saldo_clave},$E{r},"
        mapeo.cell(row=r, column=6, value=f"={base}{rg.saldo_periodo},{CELDA_PERIODO})")
        mapeo.cell(row=r, column=7, value=f"=D{r}*F{r}")
        mapeo.cell(row=r, column=8, value=f"={base}{_acumulado(rg.saldo_periodo)})")
        mapeo.cell(row=r, column=9, value=f"=D{r}*H{r}")
        for c in (6, 7, 8, 9):
            mapeo.cell(row=r, column=c).number_format = FORMATO_IMPORTE

    budget = libro.create_sheet("Budget")
    _encabezados(budget, ("Fila en Variance", "Línea", "Periodo (AAAAMM; 0 = anual)", "Importe"))
    n_budget = 1
    entradas: list[tuple[int, int, Decimal]] = []
    if ins.budget_mensual is not None:
        entradas += [(ln, numero_de_periodo(p), v) for (ln, p), v in ins.budget_mensual.items()]
    entradas += [(ln, 0, v) for ln, v in ins.budget_anual.items()]
    for linea, periodo_n, importe in sorted(entradas):
        if linea not in filas or _es_total_por_regla(linea, ins):
            continue
        n_budget += 1
        budget.cell(row=n_budget, column=1, value=filas[linea])
        escribir_texto(budget, n_budget, 2, nombres[linea])
        budget.cell(row=n_budget, column=3, value=periodo_n)
        budget.cell(row=n_budget, column=4, value=importe).number_format = FORMATO_IMPORTE

    capturas = libro.create_sheet("Capturas")
    _encabezados(capturas, ("Fila en Variance", "Línea", "Periodo (AAAAMM)", "Importe capturado"))
    n_capturas = 1
    for linea, p, importe in sorted(ins.capturas, key=lambda x: (x[0], x[1])):
        if linea not in filas:
            continue
        n_capturas += 1
        capturas.cell(row=n_capturas, column=1, value=filas[linea])
        escribir_texto(capturas, n_capturas, 2, nombres[linea])
        capturas.cell(row=n_capturas, column=3, value=numero_de_periodo(p))
        capturas.cell(row=n_capturas, column=4, value=importe).number_format = FORMATO_IMPORTE

    for hoja, anchos in (
        (saldos, (16, 16, 18, 22)),
        (mapeo, (16, 40, 16, 8, 18, 16, 16, 16, 16)),
        (budget, (16, 40, 26, 16)),
        (capturas, (16, 40, 18, 18)),
    ):
        for i, a in enumerate(anchos, 1):
            hoja.column_dimensions[_col(i)].width = a
        hoja.freeze_panes = "A2"
    return Rangos(n_saldos, n_mapeo, n_budget, n_capturas)


def escribir_parametros(hoja: Worksheet, periodo: str, inicio_ytd: str) -> None:
    """Periodo e inicio del año fiscal: las fórmulas leen estas celdas (no se cambian a mano)."""
    escribir_texto(hoja, FILA_PARAMETROS, 1, "Periodo (AAAAMM):")
    hoja.cell(row=FILA_PARAMETROS, column=2, value=numero_de_periodo(periodo))
    escribir_texto(hoja, FILA_PARAMETROS, 3, "Inicio año fiscal:")
    hoja.cell(row=FILA_PARAMETROS, column=4, value=numero_de_periodo(inicio_ytd))


class EscritorDeFormulas:
    """Escribe las cifras de cada línea de `Variance` como fórmulas."""

    def __init__(
        self,
        hoja: Worksheet,
        filas: Mapping[int, int],
        ins: InsumosFormulas,
        rangos: Rangos,
        resultado: Mapping[int, LineaVariance],
        capturadas: frozenset[int],
    ) -> None:
        self.hoja, self.filas, self.ins, self.rangos = hoja, filas, ins, rangos
        self.resultado, self.capturadas = resultado, capturadas
        self.componentes = {c for rs in ins.reglas.values() for c, _ in rs}

    # --- referencias ------------------------------------------------------------------------
    def _celda(self, bloque: str, linea: int, desplazamiento: int) -> str:
        return f"{_col(BLOQUES[bloque] + desplazamiento)}{self.filas[linea]}"

    def _celdas(self, bloque: str, desplazamiento: int) -> dict[int, str]:
        return {linea: self._celda(bloque, linea, desplazamiento) for linea in self.filas}

    # --- fórmulas ---------------------------------------------------------------------------
    def _actual(self, bloque: str, linea: LineaReporte) -> str:
        ins, rg, fila = self.ins, self.rangos, self.filas[linea.id]
        if bloque == "anual":
            return f"={self._celda('ytd', linea.id, 0)}"
        partes: list[str] = []
        regla = ins.reglas.get(linea.id)
        if linea.id in ins.cuentas and (regla is None or linea.id in ins.aditivas):
            rango = rg.mapeo_mes if bloque == "mes" else rg.mapeo_acum
            partes.append(f"SUMIFS({rango},{rg.mapeo_fila},{fila})")
        if regla is not None:
            partes.append(_suma(regla, self._celdas(bloque, 0)))
        if linea.id not in ins.cuentas and regla is None:
            base = f"{rg.cap_importe},{rg.cap_fila},{fila},"
            if bloque == "mes":
                partes.append(f"SUMIFS({base}{rg.cap_periodo},{CELDA_PERIODO})")
            elif linea.id in ins.no_se_acumulan:
                partes.append(f"AVERAGEIFS({base}{_acumulado(rg.cap_periodo)})")
            else:
                partes.append(f"SUMIFS({base}{_acumulado(rg.cap_periodo)})")
        return "=" + "+".join(partes) if partes else "=0"

    def _budget(self, bloque: str, linea: LineaReporte, valor: Decimal) -> str | Decimal:
        ins, rg, fila = self.ins, self.rangos, self.filas[linea.id]
        if _es_total_por_regla(linea.id, ins):
            return "=" + _suma(ins.reglas[linea.id], self._celdas(bloque, 1))
        if bloque == "anual":
            return f"=SUMIFS({rg.budget_importe},{rg.budget_fila},{fila},{rg.budget_periodo},0)"
        if ins.budget_mensual is None:
            return valor  # el Budget no viene mes a mes: solo se conoce la cifra
        base = f"{rg.budget_importe},{rg.budget_fila},{fila},"
        if bloque == "mes":
            return f"=SUMIFS({base}{rg.budget_periodo},{CELDA_PERIODO})"
        formula = f"SUMIFS({base}{_acumulado(rg.budget_periodo)})"
        if linea.id in ins.no_se_acumulan:
            formula += f"/{ins.meses_ytd}"
        return "=" + formula

    def _var(self, bloque: str, linea: LineaReporte) -> str:
        compuesta = self.ins.compuestas.get(bloque, {}).get(linea.id)
        if compuesta is not None:
            return "=" + _suma(compuesta, self._celdas(bloque, 2))
        actual, budget = self._celda(bloque, linea.id, 0), self._celda(bloque, linea.id, 1)
        if self.ins.sentidos[bloque][linea.id] is Convencion.REAL_MENOS_PRESUPUESTO:
            return f"={actual}-{budget}"
        return f"={budget}-{actual}"

    def _pct(self, bloque: str, linea: LineaReporte) -> str:
        v, a, b = (self._celda(bloque, linea.id, i) for i in (2, 0, 1))
        return f"=IF({v}=0,0,IF({a}=0,{v}/{b},IF({b}=0,{v}/{a},{v}/{b})))"

    # --- escritura --------------------------------------------------------------------------
    def escribir(self, linea: LineaReporte, calculada: bool) -> None:
        r = self.resultado[linea.id]
        for bloque, c in (("mes", r.mes), ("ytd", r.ytd), ("anual", r.anual)):
            self._bloque(bloque, linea, c, calculada)

    def _mostrar_actual(self, bloque: str, linea: LineaReporte, c: Cifras, calculada: bool) -> bool:
        if calculada:
            return True
        if bloque == "mes":
            return linea.id in self.capturadas
        return c.var is not None

    def _bloque(self, bloque: str, linea: LineaReporte, c: Cifras, calculada: bool) -> None:
        fila, base = self.filas[linea.id], BLOQUES[bloque]
        mostrar = self._mostrar_actual(bloque, linea, c, calculada)
        # un total que suma esta línea necesita su Actual aunque la línea no se muestre
        if mostrar or linea.id in self.componentes:
            self._poner(fila, base, self._actual(bloque, linea))
        if c.budget is not None:
            self._poner(fila, base + 1, self._budget(bloque, linea, c.budget))
        if mostrar and c.var is not None:
            self._poner(fila, base + 2, self._var(bloque, linea))
            if c.pct is not None:
                self._poner(fila, base + 3, self._pct(bloque, linea), porcentaje=True)

    def _poner(
        self, fila: int, columna: int, valor: str | Decimal, porcentaje: bool = False
    ) -> None:
        celda = self.hoja.cell(row=fila, column=columna)
        celda.value = valor
        celda.number_format = FORMATO_PORCENTAJE if porcentaje else FORMATO_IMPORTE
        celda.alignment = Alignment(horizontal="right")

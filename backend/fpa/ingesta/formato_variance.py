"""Formatos del Excel de Variance de Contraloría: Mazatlán (MZT) y Puerto Vallarta (VT).

Los dos traen la misma hoja `Base de Datos` (movimientos) y una hoja `Ago` con el reporte, pero se
arman distinto:

- **VT (Vallarta):** `Ago` tiene tres bloques (mes, acumulado y presupuesto anual) y las fórmulas
  que ligan cada línea a sus cuentas están en las columnas **S** (mes) y **T** (acumulado) de esa
  misma hoja; las cuentas de cada línea van en X a AJ. El Budget es un valor pegado (E, J, O).
- **MZT (Mazatlán):** `Ago` solo trae el **mes** (Actual en D, Budget en F, Var en G, % en H) y todo
  está pegado como valores; las cuentas de cada línea van en I a O. Las fórmulas viven en la hoja
  **`Jul-Ago`** (julio y agosto lado a lado: D a H y I a M), en la columna **E** (julio) y **J**
  (agosto), con `SUMIFS` sobre la hoja de cada cuenta. Cada mes trae su propio Budget (F y K).

Se detecta por los nombres de las hojas; si el archivo no trae `Ago`, es un error claro.
"""

from __future__ import annotations

from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

from openpyxl import load_workbook

HOJA_REPORTE = "Ago"
HOJA_DOS_MESES = "Jul-Ago"
HOJA_MOVIMIENTOS = "Base de Datos"


class Formato(StrEnum):
    VT = "vt"
    MZT = "mzt"


@dataclass(frozen=True)
class PerfilFormato:
    formato: Formato
    hoja_reporte: str
    hoja_formulas: str
    col_formula_mes: int  # índice (desde 0) de la columna con la fórmula del mes del reporte
    letra_formula_mes: str
    tiene_acumulado: bool  # el Excel trae su propio bloque de acumulado del año fiscal
    tiene_anual: bool  # el Excel trae Budget anual
    nombre: str


PERFIL_VT = PerfilFormato(
    formato=Formato.VT,
    hoja_reporte=HOJA_REPORTE,
    hoja_formulas=HOJA_REPORTE,
    col_formula_mes=18,
    letra_formula_mes="S",
    tiene_acumulado=True,
    tiene_anual=True,
    nombre="Puerto Vallarta (tres bloques: mes, acumulado y anual)",
)

PERFIL_MZT = PerfilFormato(
    formato=Formato.MZT,
    hoja_reporte=HOJA_REPORTE,
    hoja_formulas=HOJA_DOS_MESES,
    col_formula_mes=9,
    letra_formula_mes="J",
    tiene_acumulado=False,
    tiene_anual=False,
    nombre="Mazatlán (solo el mes; julio y agosto en la hoja Jul-Ago)",
)


def detectar_formato_de_hojas(nombres: list[str]) -> PerfilFormato:
    """`Jul-Ago` identifica a Mazatlán; con solo `Ago` y `Base de Datos` es Vallarta."""
    if HOJA_REPORTE not in nombres:
        raise ValueError(f'El libro no tiene la hoja "{HOJA_REPORTE}"')
    return PERFIL_MZT if HOJA_DOS_MESES in nombres else PERFIL_VT


def detectar_formato(ruta: str | Path) -> PerfilFormato:
    with Path(ruta).open("rb") as archivo:
        wb = load_workbook(archivo, read_only=True)
        try:
            return detectar_formato_de_hojas(list(wb.sheetnames))
        finally:
            wb.close()

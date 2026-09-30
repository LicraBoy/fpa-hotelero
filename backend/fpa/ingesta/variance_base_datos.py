"""Lector de la hoja "Base de Datos" del Excel de Variance (julio–agosto 2026).

Es la misma información del CON130 de Novohit, pero ya filtrada por Contraloría: solo cuentas de
resultados (clases 4, 5 y 6), una fila por movimiento, con encabezado real. Por eso **no** cuadra
Débito = Crédito (es un subconjunto del libro) y `controla_balance` es False.

Columnas (se ubican por nombre, no por posición; toleran encabezados con caracteres dañados):
Mes, Fecha, No Cuenta, Nombre de la Cuenta, Asiento - Salida, Concepto - Nombre Articulo,
Cantidad, Costo, Débito, Crédito.

Nota: en las cuentas de clase 5 (inventarios) "Asiento - Salida" puede ser un número de salida de
almacén, y Cantidad/Costo vienen llenos.
"""

from __future__ import annotations

import re
import unicodedata
from collections import Counter
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook

from fpa.ingesta.contrato import CERO, Cuenta, Fuente, Movimiento, ResultadoLectura
from fpa.ingesta.utilidades import a_decimal_exacto, a_dinero, a_fecha, a_texto

HOJA = "Base de Datos"
MESES = {
    "enero": 1,
    "febrero": 2,
    "marzo": 3,
    "abril": 4,
    "mayo": 5,
    "junio": 6,
    "julio": 7,
    "agosto": 8,
    "septiembre": 9,
    "setiembre": 9,
    "octubre": 10,
    "noviembre": 11,
    "diciembre": 12,
}

# Nombre normalizado del encabezado -> campo. Débito/Crédito también con el carácter dañado.
_ALIAS = {
    "mes": "mes",
    "fecha": "fecha",
    "nocuenta": "cuenta",
    "nombredelacuenta": "nombre",
    "asientosalida": "asiento",
    "asiento": "asiento",
    "conceptonombrearticulo": "concepto",
    "concepto": "concepto",
    "cantidad": "cantidad",
    "costo": "costo",
    "debito": "debito",
    "dbito": "debito",
    "credito": "credito",
    "crdito": "credito",
}
_OBLIGATORIOS = {"mes", "fecha", "cuenta", "nombre", "asiento", "concepto", "debito", "credito"}
_RE_CODIGO = re.compile(r"^[0-9A-Za-z][0-9A-Za-z.\-]*$")


def _normalizar(texto: str) -> str:
    sin_acentos = unicodedata.normalize("NFKD", texto)
    return re.sub(r"[^a-z0-9]", "", sin_acentos.lower())


def _ubicar_columnas(filas: Sequence[Sequence[object]]) -> tuple[int, dict[str, int]]:
    """Devuelve (índice de la fila de encabezado, campo -> índice de columna)."""
    for i, fila in enumerate(filas[:30]):
        columnas: dict[str, int] = {}
        for j, celda in enumerate(fila):
            if isinstance(celda, str):
                campo = _ALIAS.get(_normalizar(celda))
                if campo and campo not in columnas:
                    columnas[campo] = j
        if columnas.keys() >= _OBLIGATORIOS:
            return i, columnas
    raise ValueError(
        'No se encontró el encabezado de la hoja "Base de Datos" '
        "(Mes, Fecha, No Cuenta, Nombre de la Cuenta, Asiento, Concepto, Débito, Crédito)"
    )


def _codigo_cuenta(valor: object) -> str | None:
    """Código de cuenta como texto. Un número con decimales podría haber perdido ceros finales."""
    if isinstance(valor, bool):
        raise ValueError(f"código de cuenta inválido: {valor!r}")
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor))
        raise ValueError(
            f"código de cuenta numérico {valor!r}: Excel pudo quitar ceros finales "
            "(4002001.010 -> 4002001.01); escríbelo como texto"
        )
    texto = a_texto(valor)
    if texto is None:
        return None
    if not _RE_CODIGO.match(texto):
        raise ValueError(f"código de cuenta inválido: {texto!r}")
    return texto


def leer_variance_base_datos(ruta: str | Path, hoja: str = HOJA) -> ResultadoLectura:
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja not in wb.sheetnames:
            raise ValueError(f'El libro no tiene la hoja "{hoja}"')
        filas = [tuple(r) for r in wb[hoja].iter_rows(values_only=True)]
    finally:
        wb.close()
    return leer_variance_base_datos_filas(filas)


def leer_variance_base_datos_filas(filas: Sequence[Sequence[object]]) -> ResultadoLectura:
    res = ResultadoLectura(fuente=Fuente.VARIANCE_BASE_DATOS, controla_balance=False)
    i_enc, cols = _ubicar_columnas(filas)

    def celda(fila: Sequence[object], campo: str) -> object:
        j = cols.get(campo)
        return fila[j] if j is not None and j < len(fila) else None

    mes_distinto = 0
    sin_importe: list[int] = []
    mes_desconocido: Counter[str] = Counter()
    nombres_por_cuenta: dict[str, set[str]] = {}

    for n, fila in enumerate(filas[i_enc + 1 :], start=i_enc + 2):
        if not any(v is not None and str(v).strip() for v in fila):
            continue  # fila vacía
        try:
            codigo = _codigo_cuenta(celda(fila, "cuenta"))
        except ValueError as e:
            res.errores.append(f"Fila {n}: {e}")
            continue
        if codigo is None:
            # Sin cuenta: no es un movimiento (p. ej. notas sueltas); se ignora con aviso
            res.avisos.append(f"Fila {n}: fila con datos pero sin No Cuenta; se ignora")
            continue
        fecha = a_fecha(celda(fila, "fecha"))
        if fecha is None:
            res.errores.append(f"Fila {n}: fecha inválida {celda(fila, 'fecha')!r}")
            continue
        try:
            debito = a_dinero(celda(fila, "debito"))
            credito = a_dinero(celda(fila, "credito"))
            cantidad = a_decimal_exacto(celda(fila, "cantidad"))
            costo = a_decimal_exacto(celda(fila, "costo"))
        except ValueError as e:
            res.errores.append(f"Fila {n}: {e}")
            continue
        if debito is None and credito is None:
            # Existe en los datos reales (p. ej. una línea con el importe en blanco): se conserva
            # como movimiento de importe 0 y se avisa, para no perder la cuenta de filas.
            sin_importe.append(n)

        nombre = a_texto(celda(fila, "nombre")) or ""
        nombres_por_cuenta.setdefault(codigo, set()).add(nombre)

        mes_texto = a_texto(celda(fila, "mes"))
        if mes_texto:
            numero = MESES.get(_normalizar(mes_texto))
            if numero is None:
                mes_desconocido[mes_texto] += 1
            elif numero != fecha.month:
                mes_distinto += 1

        d, c = debito or CERO, credito or CERO
        res.movimientos.append(
            Movimiento(
                cuenta_codigo=codigo,
                cuenta_descripcion=nombre,
                asiento=a_texto(celda(fila, "asiento")) or "",
                fecha=fecha,
                concepto=a_texto(celda(fila, "concepto")),
                debito=d,
                credito=c,
                fila=n,
                cantidad=cantidad,
                costo=costo,
            )
        )
        cuenta = res.cuentas.get(codigo)
        if cuenta is None:
            cuenta = Cuenta(codigo=codigo, descripcion=nombre, fila=n)
            res.cuentas[codigo] = cuenta
        cuenta.suma_debito += d
        cuenta.suma_credito += c

    if sin_importe:
        muestra = ", ".join(str(n) for n in sin_importe[:5])
        res.avisos.append(
            f"{len(sin_importe)} movimientos sin Débito ni Crédito "
            f"(se toman como 0): filas {muestra}"
        )
    if mes_distinto:
        res.avisos.append(f"{mes_distinto} filas con el mes en texto distinto al mes de la fecha")
    for texto, veces in mes_desconocido.items():
        res.avisos.append(f"Mes no reconocido {texto!r} en {veces} filas")
    for codigo, nombres in nombres_por_cuenta.items():
        if len(nombres) > 1:
            res.avisos.append(f"La cuenta {codigo} aparece con {len(nombres)} nombres distintos")
    return res

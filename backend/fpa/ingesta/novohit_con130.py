"""Lector del reporte CON130 "Mayor Analítico (movimientos por cuentas)" de Novohit.

El reporte no es una tabla plana: intercala títulos de jerarquía, encabezados de
cuenta, filas "Saldo anterior", movimientos y filas "Total". Este lector
clasifica cada fila y entrega solo los movimientos, más los controles
(Débito = Crédito y total por cuenta = suma de sus movimientos).

Reglas (ver Blueprint, sección 12.1):
- Las columnas se ubican a partir del encabezado "# ASIENTO"; las columnas
  auxiliares que Contraloría agrega a la izquierda (mes, fecha, código,
  descripción, F) no se usan.
- Una cuenta con movimientos se abre en la fila que trae el encabezado de la
  cuenta en la columna "# ASIENTO" y "Saldo anterior" en Concepto.
- El Saldo anterior no entra al cálculo del Actual.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

CENTAVO = Decimal("0.01")
_RE_CUENTA = re.compile(r"^(\d+(?:\.\d+)*)\s+(.*\S)\s*$")
_RE_FECHA = re.compile(r"^\s*(\d{2})/(\d{2})/(\d{4})\s*$")
CAMPOS = (
    "asiento",
    "fecha",
    "tp_asiento",
    "operacion",
    "tp_doc",
    "numero",
    "fecha_doc",
    "concepto",
    "debito",
    "credito",
    "saldo",
)


@dataclass
class Movimiento:
    cuenta_codigo: str
    cuenta_descripcion: str
    asiento: str
    fecha: date
    tp_asiento: str | None
    operacion: str | None
    tp_doc: str | None
    numero: str | None
    concepto: str | None
    debito: Decimal
    credito: Decimal
    saldo: Decimal | None
    fila: int


@dataclass
class Cuenta:
    codigo: str
    descripcion: str
    fila: int
    saldo_anterior: Decimal | None = None
    suma_debito: Decimal = Decimal("0")
    suma_credito: Decimal = Decimal("0")
    total_debito_reporte: Decimal | None = None
    total_credito_reporte: Decimal | None = None
    # True = el Total del reporte coincide; False = no coincide; None = no trae Total
    total_verificado: bool | None = None


@dataclass
class TotalGrupo:
    descripcion: str
    debito: Decimal
    credito: Decimal
    fila: int


@dataclass
class ResultadoCon130:
    empresa: str | None = None
    periodo_texto: str | None = None
    movimientos: list[Movimiento] = field(default_factory=list)
    cuentas: dict[str, Cuenta] = field(default_factory=dict)
    totales_grupo: list[TotalGrupo] = field(default_factory=list)
    filas_ignoradas: int = 0
    avisos: list[str] = field(default_factory=list)
    errores: list[str] = field(default_factory=list)

    @property
    def total_debito(self) -> Decimal:
        return sum((m.debito for m in self.movimientos), Decimal("0"))

    @property
    def total_credito(self) -> Decimal:
        return sum((m.credito for m in self.movimientos), Decimal("0"))

    @property
    def cuadra_debito_credito(self) -> bool:
        return abs(self.total_debito - self.total_credito) < CENTAVO

    @property
    def es_valido(self) -> bool:
        return not self.errores and self.cuadra_debito_credito


def _dinero(valor: object) -> Decimal | None:
    """Número o texto con comas (ej. '16,189.85') -> Decimal; vacío -> None."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        raise ValueError(f"monto inválido: {valor!r}")
    if isinstance(valor, (int, float)):
        return Decimal(str(valor)).quantize(CENTAVO)
    texto = str(valor).strip().replace(",", "").replace("$", "")
    if texto == "":
        return None
    try:
        return Decimal(texto).quantize(CENTAVO)
    except InvalidOperation:
        raise ValueError(f"monto inválido: {valor!r}") from None


def _fecha(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        m = _RE_FECHA.match(valor)
        if m:
            dia, mes, anio = (int(x) for x in m.groups())
            try:
                return date(anio, mes, dia)
            except ValueError:
                return None
    return None


def _texto(valor: object) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _ubicar_encabezado(filas: Sequence[Sequence[object]]) -> tuple[int, int]:
    """Devuelve (índice de fila, índice de columna) del encabezado '# ASIENTO'."""
    for i, fila in enumerate(filas[:40]):
        for j, celda in enumerate(fila):
            if isinstance(celda, str) and celda.strip().upper() == "# ASIENTO":
                return i, j
    raise ValueError('No se encontró el encabezado "# ASIENTO": no parece un CON130 de Novohit')


def leer_con130(ruta: str | Path, hoja: str | None = None) -> ResultadoCon130:
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb[hoja] if hoja else wb.worksheets[0]
        filas = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return leer_con130_filas(filas)


def leer_con130_filas(filas: Sequence[Sequence[object]]) -> ResultadoCon130:
    res = ResultadoCon130()
    i_enc, c0 = _ubicar_encabezado(filas)

    # Datos de cabecera (antes del encabezado de columnas)
    for fila in filas[:i_enc]:
        for celda in fila:
            if not isinstance(celda, str):
                continue
            texto = celda.strip()
            if texto.lower().startswith("periodo") and res.periodo_texto is None:
                res.periodo_texto = texto
        if res.empresa is None:
            for celda in fila:
                if (
                    isinstance(celda, str)
                    and celda.strip()
                    and not celda.strip().lower().startswith("novohit")
                ):
                    res.empresa = celda.strip()
                    break

    def col(fila: Sequence[object], campo: str) -> object:
        idx = c0 + CAMPOS.index(campo)
        return fila[idx] if idx < len(fila) else None

    abierta: Cuenta | None = None

    def cerrar_sin_total() -> None:
        nonlocal abierta
        if abierta is not None and abierta.total_verificado is None:
            res.avisos.append(f"Cuenta {abierta.codigo} (fila {abierta.fila}) no trae fila 'Total'")
        abierta = None

    for n, fila in enumerate(filas[i_enc + 1 :], start=i_enc + 2):
        asiento_o_cuenta = _texto(col(fila, "asiento"))
        concepto = _texto(col(fila, "concepto"))
        fecha = _fecha(col(fila, "fecha"))
        try:
            debito = _dinero(col(fila, "debito"))
            credito = _dinero(col(fila, "credito"))
            saldo = _dinero(col(fila, "saldo"))
        except ValueError as e:
            res.errores.append(f"Fila {n}: {e}")
            continue

        # 1) Apertura de cuenta: encabezado de cuenta + "Saldo anterior" en la misma fila
        if concepto and concepto.lower() == "saldo anterior":
            m = _RE_CUENTA.match(asiento_o_cuenta or "")
            if not m:
                res.errores.append(f"Fila {n}: 'Saldo anterior' sin encabezado de cuenta")
                continue
            cerrar_sin_total()
            codigo, descripcion = m.group(1), m.group(2)
            if codigo in res.cuentas:
                res.errores.append(f"Fila {n}: la cuenta {codigo} aparece más de una vez")
            abierta = Cuenta(
                codigo=codigo,
                descripcion=descripcion,
                fila=n,
                saldo_anterior=debito if debito is not None else credito,
            )
            res.cuentas[codigo] = abierta
            continue

        # 2) Fila "Total ..."
        if concepto and concepto.lower().startswith("total"):
            if abierta is not None:
                abierta.total_debito_reporte = debito or Decimal("0")
                abierta.total_credito_reporte = credito or Decimal("0")
                coincide = (
                    abs(abierta.total_debito_reporte - abierta.suma_debito) < CENTAVO
                    and abs(abierta.total_credito_reporte - abierta.suma_credito) < CENTAVO
                )
                abierta.total_verificado = coincide
                if not coincide:
                    res.errores.append(
                        f"Fila {n}: el Total de {abierta.codigo} ({abierta.total_debito_reporte}/"
                        f"{abierta.total_credito_reporte}) no coincide con la suma de sus "
                        "movimientos "
                        f"({abierta.suma_debito}/{abierta.suma_credito})"
                    )
                abierta = None
            else:
                res.totales_grupo.append(
                    TotalGrupo(concepto, debito or Decimal("0"), credito or Decimal("0"), n)
                )
            continue

        # 3) Movimiento
        if fecha is not None:
            if abierta is None:
                res.errores.append(f"Fila {n}: movimiento sin cuenta abierta")
                continue
            if debito is None and credito is None:
                res.errores.append(f"Fila {n}: movimiento sin Débito ni Crédito")
                continue
            d, c = debito or Decimal("0"), credito or Decimal("0")
            res.movimientos.append(
                Movimiento(
                    cuenta_codigo=abierta.codigo,
                    cuenta_descripcion=abierta.descripcion,
                    asiento=asiento_o_cuenta or "",
                    fecha=fecha,
                    tp_asiento=_texto(col(fila, "tp_asiento")),
                    operacion=_texto(col(fila, "operacion")),
                    tp_doc=_texto(col(fila, "tp_doc")),
                    numero=_texto(col(fila, "numero")),
                    concepto=concepto,
                    debito=d,
                    credito=c,
                    saldo=saldo,
                    fila=n,
                )
            )
            abierta.suma_debito += d
            abierta.suma_credito += c
            continue

        # 4) Títulos de jerarquía, filas vacías y todo lo demás
        if any(v is not None and str(v).strip() for v in fila):
            res.filas_ignoradas += 1

    cerrar_sin_total()

    if not res.cuadra_debito_credito:
        res.errores.append(
            f"Débito total {res.total_debito} distinto de Crédito total {res.total_credito}"
        )
    return res

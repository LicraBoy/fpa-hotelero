"""Filas crudas de un reporte de Novohit exportado como .xlsx, .xls o .csv.

Los lectores de reportes (`leer_con130_filas`) trabajan sobre una lista de filas; de dónde salen
esas filas depende del formato del archivo. Aquí se decide **por el contenido**, no por la extensión
(la extensión ya la validó la capa de subida):

- `.xlsx` (ZIP): openpyxl, primera hoja, valores calculados.
- `.xls` (OLE2, el Excel antiguo que exporta Novohit): xlrd; las celdas de fecha se vuelven fechas.
- `.csv` / texto: separador detectado (`,` `;` tabulador `|`) y codificación UTF-8 con o sin BOM,
  o Windows-1252 si no es UTF-8. Todo llega como texto: los importes con comas y las fechas
  `dd/mm/aaaa` ya los entienden `a_dinero` y `a_fecha`.

Nunca se ejecuta ni se evalúa el contenido del archivo.
"""

from __future__ import annotations

import contextlib
import csv
import io
from pathlib import Path

import xlrd
from openpyxl import load_workbook

FIRMA_ZIP = b"PK\x03\x04"
FIRMA_OLE2 = bytes.fromhex("D0CF11E0A1B11AE1")
_SEPARADORES = ",;\t|"
# Una celda de CSV más larga que esto no es un reporte contable: se rechaza el archivo
_MAX_CELDA = 10_000

Fila = tuple[object, ...]


def formato_por_contenido(ruta: Path) -> str:
    """`xlsx`, `xls` o `csv` según los primeros bytes del archivo."""
    with ruta.open("rb") as f:
        inicio = f.read(8)
    if inicio.startswith(FIRMA_ZIP):
        return "xlsx"
    if inicio == FIRMA_OLE2:
        return "xls"
    return "csv"


def _filas_xlsx(ruta: Path) -> list[Fila]:
    with ruta.open("rb") as archivo:
        wb = load_workbook(archivo, read_only=True, data_only=True)
        try:
            return [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()


def _filas_xls(ruta: Path) -> list[Fila]:
    try:
        libro = xlrd.open_workbook(file_contents=ruta.read_bytes())
    except Exception as e:  # xlrd lanza errores de varios tipos ante archivos dañados
        raise ValueError(
            f"No se pudo leer el archivo como un libro de Excel .xls ({type(e).__name__})"
        ) from None
    hoja = libro.sheet_by_index(0)
    filas: list[Fila] = []
    for r in range(hoja.nrows):
        fila: list[object] = []
        for c in range(hoja.ncols):
            tipo = hoja.cell_type(r, c)
            valor = hoja.cell_value(r, c)
            if tipo == xlrd.XL_CELL_DATE:
                with contextlib.suppress(xlrd.XLDateError, OverflowError, ValueError):
                    valor = xlrd.xldate_as_datetime(valor, libro.datemode)
            elif tipo in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or (
                tipo == xlrd.XL_CELL_TEXT and isinstance(valor, str) and valor == ""
            ):
                valor = None
            fila.append(valor)
        filas.append(tuple(fila))
    return filas


def _decodificar(contenido: bytes) -> str:
    if b"\x00" in contenido[:4096]:
        raise ValueError("El archivo no es un CSV de texto")
    try:
        return contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        return contenido.decode("cp1252", errors="replace")


def _separador(texto: str) -> str | None:
    """El separador que más se repite en la línea de encabezados del reporte (`# ASIENTO`).

    Los importes traen comas (`1,030.50`) aunque el separador sea otro, así que el análisis
    automático de `csv.Sniffer` se confunde: se mira solo la línea de encabezados, que no las trae.
    """
    for linea in texto.splitlines()[:80]:
        if "asiento" in linea.lower() and "fecha" in linea.lower():
            cuenta = {sep: linea.count(sep) for sep in _SEPARADORES}
            mejor = max(cuenta, key=lambda sep: cuenta[sep])
            return mejor if cuenta[mejor] >= 2 else None
    return None


def _filas_csv(ruta: Path) -> list[Fila]:
    texto = _decodificar(ruta.read_bytes())
    separador = _separador(texto)
    if separador is not None:
        lector = csv.reader(io.StringIO(texto, newline=""), csv.excel, delimiter=separador)
    else:
        muestra = chr(10).join(texto.splitlines()[:50])
        try:
            dialecto: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
                muestra, delimiters=_SEPARADORES
            )
        except csv.Error:
            dialecto = csv.excel  # sin pistas claras: coma
        lector = csv.reader(io.StringIO(texto, newline=""), dialecto)
    filas: list[Fila] = []
    for fila in lector:
        if any(len(c) > _MAX_CELDA for c in fila):
            raise ValueError("El archivo no parece un reporte: trae celdas demasiado largas")
        filas.append(tuple(c.strip() if c.strip() else None for c in fila))
    return filas


def leer_filas(ruta: str | Path) -> list[Fila]:
    """Todas las filas de la primera hoja (o del CSV). `ValueError` si no se puede leer."""
    ruta = Path(ruta)
    formato = formato_por_contenido(ruta)
    if formato == "xlsx":
        return _filas_xlsx(ruta)
    if formato == "xls":
        return _filas_xls(ruta)
    return _filas_csv(ruta)

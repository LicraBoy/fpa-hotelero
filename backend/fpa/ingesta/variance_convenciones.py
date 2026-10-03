"""Sentido de la Variación de cada línea, leído de las fórmulas del Excel de Contraloría.

En `02. MZT August 2026 Variance.xlsx` (hoja `Variance August 2026`) la columna **G** (mes), la
**L** (acumulado) y la **Q** (anual) traen, línea por línea, cómo se calcula la Variación:

    =+F12-E12    Presupuesto − Actual  (gastos)       =+E12-F12    Actual − Presupuesto (ingresos)
    =-F12+E12    Actual − Presupuesto                 =+K36-J36    Presupuesto − Actual (acumulado)

En vez de deducir el sentido por el tipo de cuenta, se lee tal cual está escrito. Una fórmula que no
es una resta de Actual y Presupuesto (un total que suma otras Variaciones, un `SUM`) no da sentido:
esa línea sigue la deducción habitual. **Nunca se evalúa** el texto del Excel: solo se reconocen
restas de las dos celdas del bloque.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from fpa.calculo.variance import Convencion
from fpa.ingesta.variance_reglas import analizar_formula

HOJA = "Variance August 2026"
# (columna de la Variación, columna del Actual, columna del Presupuesto), índices desde 0
BLOQUES = {"mes": (6, "E", "F"), "ytd": (11, "J", "K"), "anual": (16, "O", "P")}
_LETRA_VARIACION = {"mes": "G", "ytd": "L", "anual": "Q"}
_TERMINO = re.compile(r"([+-]?)\s*\$?([A-Z]{1,2})\$?(\d+)")


@dataclass
class ConvencionesLeidas:
    # fila del Excel -> bloque -> sentido de la Variación
    por_fila: dict[int, dict[str, Convencion]] = field(default_factory=dict)
    # fila -> bloque -> [(desde, hasta, signo)]: la Variación es la suma de otras Variaciones
    # (`=G613-G614`, `=SUM(G11:G31)`): el Excel no la calcula como Actual − Presupuesto
    compuestas: dict[int, dict[str, list[tuple[int, int, int]]]] = field(default_factory=dict)
    avisos: list[str] = field(default_factory=list)

    @property
    def n_lineas(self) -> int:
        return len(set(self.por_fila) | set(self.compuestas))


def sentido_de_formula(formula: object, col_actual: str, col_budget: str) -> Convencion | None:
    """`=+F12-E12` -> Presupuesto − Actual; `=+E12-F12` -> Actual − Presupuesto; otra -> None."""
    if not (isinstance(formula, str) and formula.startswith("=")):
        return None
    cuerpo = formula[1:].strip()
    terminos = _TERMINO.findall(cuerpo)
    # la fórmula debe ser SOLO esos dos términos (nada más: ni SUM, ni números, ni otras celdas)
    if len(terminos) != 2 or _TERMINO.sub("", cuerpo).strip():
        return None
    signos: dict[str, int] = {}
    filas = set()
    for signo, columna, fila in terminos:
        signos[columna] = -1 if signo == "-" else 1
        filas.add(fila)
    if len(filas) != 1 or set(signos) != {col_actual, col_budget}:
        return None
    if signos[col_actual] == 1 and signos[col_budget] == -1:
        return Convencion.REAL_MENOS_PRESUPUESTO
    if signos[col_actual] == -1 and signos[col_budget] == 1:
        return Convencion.PRESUPUESTO_MENOS_REAL
    return None


def leer_convenciones_filas(filas: Sequence[Sequence[object]]) -> ConvencionesLeidas:
    """Sentido de cada línea y bloque a partir de las filas de la hoja (con fórmulas)."""
    res = ConvencionesLeidas()
    for numero, fila in enumerate(filas, start=1):
        for bloque, (col_var, col_actual, col_budget) in BLOQUES.items():
            valor = fila[col_var] if len(fila) > col_var else None
            sentido = sentido_de_formula(valor, col_actual, col_budget)
            if sentido is not None:
                res.por_fila.setdefault(numero, {})[bloque] = sentido
                continue
            if isinstance(valor, str):
                terminos = analizar_formula(valor, _LETRA_VARIACION[bloque])
                if terminos:
                    res.compuestas.setdefault(numero, {})[bloque] = [
                        (desde, hasta, -1 if signo == "-" else 1)
                        for signo, desde, hasta in terminos
                    ]
    if not res.por_fila and not res.compuestas:
        res.avisos.append("La hoja no trae fórmulas de Variación (columnas G, L y Q)")
    return res


def leer_convenciones(ruta: str | Path, hoja: str = HOJA) -> ConvencionesLeidas:
    with Path(ruta).open("rb") as archivo:
        wb = load_workbook(archivo, read_only=True, data_only=False)
        try:
            if hoja not in wb.sheetnames:
                raise ValueError(f'El libro no tiene la hoja "{hoja}"')
            filas = [tuple(r) for r in wb[hoja].iter_rows(max_col=17, values_only=True)]
        finally:
            wb.close()
    return leer_convenciones_filas(filas)

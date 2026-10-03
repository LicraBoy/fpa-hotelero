"""Lector del Budget del año fiscal de Mazatlán (`Budget MZT 26-27.xlsx`).

Una sola hoja (`Budget 26-27 MZT`, solo valores) con el mismo orden de filas que el reporte:

    B: cuenta corporativa   C: descripción   D..O: JUL … JUN   P: TOTAL del año

La fila 8 es el **tipo de cambio** del Budget (18 cada mes; en el reporte es la fila 9). El año
fiscal sale del título (`FY 26/27 Budget Performance` -> `2026/27`, de julio de 2026 a junio de
2027). Se alinea con las líneas del reporte por el **número de fila y el nombre** (con
tolerancia: el Budget escribe `Holydays` donde el reporte dice `Holidays`); lo que no se pueda
alinear se avisa, no se inventa.
"""

from __future__ import annotations

import re
from collections.abc import Collection, Sequence
from dataclasses import dataclass, field
from decimal import Decimal
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

from fpa.ingesta.variance_mzt import ValoresMes
from fpa.ingesta.variance_valores import a_decimal

HOJA = "Budget 26-27 MZT"
MESES = ["JUL", "AUG", "SEP", "OCT", "NOV", "DEC", "JAN", "FEB", "MAR", "APR", "MAY", "JUN"]
COL_PRIMER_MES = 3  # D
COL_TOTAL = 15  # P
_RE_ANIO = re.compile(r"FY\s*(\d{2})\s*/\s*(\d{2})", re.IGNORECASE)
_UMBRAL_PARECIDO = 0.6


@dataclass
class FilaBudget:
    fila: int  # número de fila en el Excel
    descripcion: str
    meses: list[Decimal | None]  # 12 valores, julio a junio
    total: Decimal | None


@dataclass
class BudgetAnualLeido:
    ano_fiscal: str  # 2026/27
    periodos: list[str]  # 2026-07 ... 2027-06
    filas: list[FilaBudget] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


@dataclass
class BudgetAlineado:
    """El Budget ya puesto sobre las líneas del reporte (por la fila de la línea en `Ago`)."""

    ano_fiscal: str
    por_periodo: dict[str, dict[int, ValoresMes]]
    anual: dict[int, Decimal]  # fila de la línea -> Budget del año completo
    avisos: list[str] = field(default_factory=list)


def _normalizar(texto: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(texto or "").lower())


def es_budget_anual_mzt(ruta: str | Path) -> bool:
    """¿El libro es el Budget del año de Mazatlán (trae la hoja `Budget 26-27 MZT`)?"""
    with Path(ruta).open("rb") as archivo:
        wb = load_workbook(archivo, read_only=True)
        try:
            return HOJA in wb.sheetnames
        finally:
            wb.close()


def periodos_del_titulo(titulo: object) -> tuple[str, list[str]]:
    """`FY 26/27 Budget Performance` -> (`2026/27`, [`2026-07`, …, `2027-06`])."""
    m = _RE_ANIO.search(str(titulo or ""))
    if m is None:
        raise ValueError("No se encontró el año fiscal en el título (p. ej. «FY 26/27 …»)")
    inicio, fin = 2000 + int(m.group(1)), 2000 + int(m.group(2))
    if fin % 100 != (inicio + 1) % 100:
        raise ValueError(f"El año fiscal «{m.group(0)}» no es de años consecutivos")
    periodos = [f"{inicio}-{mes:02d}" for mes in range(7, 13)]
    periodos += [f"{fin if fin > inicio else inicio + 1}-{mes:02d}" for mes in range(1, 7)]
    return f"{inicio}/{fin % 100:02d}", periodos


def leer_budget_anual_filas(filas: Sequence[Sequence[object]]) -> BudgetAnualLeido:
    """Lee las filas de la hoja ya cargadas (sin tocar el disco): para probar sin un Excel."""
    titulo = next((c for f in filas[:5] for c in f if isinstance(c, str) and "FY" in c), None)
    ano, periodos = periodos_del_titulo(titulo)
    encabezado = next(
        (
            i
            for i, f in enumerate(filas[:15])
            if len(f) > COL_TOTAL
            and [str(c or "").strip().upper() for c in f[COL_PRIMER_MES : COL_PRIMER_MES + 12]]
            == MESES
        ),
        None,
    )
    if encabezado is None:
        raise ValueError("No se encontró el encabezado JUL … JUN en las columnas D a O")
    res = BudgetAnualLeido(ano_fiscal=ano, periodos=periodos)
    for n, fila in enumerate(filas[encabezado + 1 :], start=encabezado + 2):
        celdas = list(fila) + [None] * (COL_TOTAL + 1 - len(fila))
        meses = [a_decimal(celdas[COL_PRIMER_MES + i]) for i in range(12)]
        total = a_decimal(celdas[COL_TOTAL])
        if all(m is None for m in meses) and total is None:
            continue
        descripcion = str(celdas[2] if celdas[2] not in (None, "") else celdas[1] or "").strip()
        res.filas.append(FilaBudget(n, descripcion, meses, total))
    if not res.filas:
        res.avisos.append("La hoja no trae cifras de Budget")
    return res


def leer_budget_anual(ruta: str | Path) -> BudgetAnualLeido:
    with Path(ruta).open("rb") as archivo:
        wb = load_workbook(archivo, read_only=True, data_only=True)
        try:
            if HOJA not in wb.sheetnames:
                raise ValueError(f'El libro no tiene la hoja "{HOJA}"')
            filas = [tuple(r) for r in wb[HOJA].iter_rows(max_col=COL_TOTAL + 1, values_only=True)]
        finally:
            wb.close()
    return leer_budget_anual_filas(filas)


def _parecido(a: str, b: str) -> float:
    return SequenceMatcher(None, _normalizar(a), _normalizar(b)).ratio()


def alinear_con_lineas(
    leido: BudgetAnualLeido,
    lineas: Sequence[tuple[int, str]],
    filas_de_detalle: Collection[int] = frozenset(),
) -> BudgetAlineado:
    """Pone cada fila del Budget en la línea del reporte que le corresponde.

    `lineas`: (fila de la línea en el reporte, su descripción). Una fila del Budget cae en la
    línea de **su mismo número de fila** si los nombres se parecen; el tipo de cambio se busca por
    nombre (en el Budget está una fila arriba). Lo que no se pueda alinear queda en los avisos.

    `filas_de_detalle`: las líneas de detalle del reporte. En el Excel de Contraloría cada una toma
    su Budget de una celda del archivo y **una celda en blanco vale 0**: una línea de detalle sin
    cifras en el Budget queda con 0 (así la Variación sale como en el Excel, no en «—»).
    """
    por_fila = dict(lineas)
    por_nombre = {_normalizar(d): f for f, d in lineas}
    salida = BudgetAlineado(
        ano_fiscal=leido.ano_fiscal, por_periodo={p: {} for p in leido.periodos}, anual={}
    )
    sin_linea: list[str] = []
    aproximadas = 0
    for fila in leido.filas:
        destino: int | None = None
        if _normalizar(fila.descripcion) == _normalizar("Exchange Rate"):
            destino = por_nombre.get(_normalizar("Exchange Rate:")) or por_nombre.get(
                _normalizar("Exchange Rate")
            )
        elif fila.fila in por_fila:
            parecido = _parecido(fila.descripcion, por_fila[fila.fila])
            if parecido >= _UMBRAL_PARECIDO:
                destino = fila.fila
                aproximadas += 1 if parecido < 1 else 0
        if destino is None:
            if fila.descripcion:
                sin_linea.append(f"{fila.fila} ({fila.descripcion[:30]})")
            continue
        for periodo, valor in zip(leido.periodos, fila.meses, strict=True):
            if valor is not None:
                salida.por_periodo[periodo][destino] = ValoresMes(destino, budget=valor)
        if fila.total is not None:
            salida.anual[destino] = fila.total
    en_cero = 0
    for fila_linea in sorted(filas_de_detalle):
        for periodo in leido.periodos:
            if fila_linea not in salida.por_periodo[periodo]:
                salida.por_periodo[periodo][fila_linea] = ValoresMes(fila_linea, budget=Decimal(0))
                en_cero += 1
        if fila_linea not in salida.anual:
            salida.anual[fila_linea] = Decimal(0)
    if en_cero:
        salida.avisos.append(
            f"{en_cero} celdas de líneas de detalle sin cifra en el Budget se toman como 0 "
            "(en el Excel una celda en blanco vale 0)"
        )
    if aproximadas:
        salida.avisos.append(
            f"{aproximadas} líneas se alinearon por la fila aunque el nombre se escribe distinto "
            "(p. ej. «Holydays» / «Holidays»)"
        )
    if sin_linea:
        salida.avisos.append(
            f"{len(sin_linea)} filas del Budget sin línea en el reporte se ignoran: "
            + ", ".join(sin_linea[:6])
        )
    return salida


__all__ = [
    "HOJA",
    "BudgetAlineado",
    "BudgetAnualLeido",
    "alinear_con_lineas",
    "es_budget_anual_mzt",
    "leer_budget_anual",
    "leer_budget_anual_filas",
    "periodos_del_titulo",
]

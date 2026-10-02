"""Reglas de suma del reporte Variance: cómo se arma cada total y subtotal.

Están escondidas en las fórmulas de la columna S de la hoja `Ago` (`=SUM(S29:S68)`, `=S69+S70`,
`=S348-S349`). Se leen con un analizador propio de suma/resta de referencias: **nunca se evalúa
texto del Excel**. Cada regla se guarda ya expandida: un `SUM(S29:S68)` se vuelve la lista de
líneas cuya fila cae en ese rango (todas menos las secciones).
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from fpa.ingesta.variance_estructura import (
    HOJA,
    LineaLeida,
    ResultadoEstructura,
    TipoLinea,
    leer_variance_estructura_filas,
)

COLUMNA_S = 18  # índice (desde 0) de la columna S (Vallarta); en Mazatlán es la J (índice 9)


@lru_cache(maxsize=8)
def _regex_termino(letra: str) -> re.Pattern[str]:
    """Término de una fórmula de suma sobre la columna `letra`: `S12`, `-S12` o `SUM(S1:S5)`."""
    return re.compile(
        rf"([+-]?)\s*(?:SUM\(\s*{letra}(\d+)\s*:\s*{letra}(\d+)\s*\)|{letra}(\d+))",
        re.IGNORECASE,
    )


_SUMABLES = {TipoLinea.DETALLE, TipoLinea.SUBTOTAL, TipoLinea.TOTAL, TipoLinea.INFORMATIVA}


@dataclass
class ReglaLeida:
    orden: int  # línea que es el total o subtotal
    fila_origen: int
    # (orden de la línea componente, signo +1 o -1)
    componentes: list[tuple[int, int]] = field(default_factory=list)
    formula: str = ""


@dataclass
class ResultadoReglas:
    reglas: list[ReglaLeida] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    errores: list[str] = field(default_factory=list)

    @property
    def es_valido(self) -> bool:
        return not self.errores


def analizar_formula(formula: str, letra: str = "S") -> list[tuple[str, int, int]] | None:
    """`=S69+S70` -> [("+", 69, 69), ("+", 70, 70)]; `=SUM(S1:S5)` -> [("+", 1, 5)].

    Devuelve None si la fórmula trae algo distinto de sumas/restas de celdas de la columna S.
    """
    texto = formula.strip()
    if not texto.startswith("="):
        return None
    cuerpo = texto[1:]
    terminos: list[tuple[str, int, int]] = []
    posicion = 0
    while posicion < len(cuerpo):
        while posicion < len(cuerpo) and cuerpo[posicion] == " ":
            posicion += 1
        if posicion >= len(cuerpo):
            break
        m = _regex_termino(letra).match(cuerpo, posicion)
        if m is None or m.end() == posicion:
            return None
        signo = "-" if m.group(1) == "-" else "+"
        if m.group(4) is not None:
            fila = int(m.group(4))
            terminos.append((signo, fila, fila))
        else:
            a, b = int(m.group(2)), int(m.group(3))
            terminos.append((signo, min(a, b), max(a, b)))
        posicion = m.end()
    return terminos or None


def leer_variance_reglas(ruta: str | Path, hoja: str = HOJA) -> ResultadoReglas:
    with Path(ruta).open("rb") as archivo:
        wb = load_workbook(archivo, read_only=True, data_only=False)
        try:
            if hoja not in wb.sheetnames:
                raise ValueError(f'El libro no tiene la hoja "{hoja}"')
            filas = [tuple(r) for r in wb[hoja].iter_rows(max_col=COLUMNA_S + 1, values_only=True)]
        finally:
            wb.close()
    return leer_variance_reglas_filas(filas, leer_variance_estructura_filas(filas))


def leer_variance_reglas_filas(
    filas: Sequence[Sequence[object]],
    estructura: ResultadoEstructura,
    columna: int = COLUMNA_S,
    letra: str = "S",
) -> ResultadoReglas:
    """`columna`/`letra`: dónde están las fórmulas de totales (S en Vallarta, J en Mazatlán)."""
    res = ResultadoReglas()
    if not estructura.es_valido:
        res.errores.append("La estructura del reporte no es válida; no se pueden leer las reglas")
        return res
    por_fila: dict[int, LineaLeida] = {x.fila_origen: x for x in estructura.lineas}
    sin_formula: list[str] = []
    for linea in estructura.lineas:
        if linea.tipo not in (TipoLinea.TOTAL, TipoLinea.SUBTOTAL):
            continue
        fila = filas[linea.fila_origen - 1] if linea.fila_origen <= len(filas) else ()
        valor = fila[columna] if len(fila) > columna else None
        if not (isinstance(valor, str) and valor.startswith("=")):
            sin_formula.append(f"{linea.descripcion} (fila {linea.fila_origen})")
            continue
        terminos = analizar_formula(valor, letra)
        if terminos is None:
            res.avisos.append(
                f"Fila {linea.fila_origen} ({linea.descripcion}): fórmula no reconocida {valor!r}"
            )
            continue
        acumulado: dict[int, int] = {}
        for signo, desde, hasta in terminos:
            factor = -1 if signo == "-" else 1
            if desde == hasta:
                destino = por_fila.get(desde)
                if destino is None or destino.tipo not in _SUMABLES:
                    res.avisos.append(
                        f"Fila {linea.fila_origen} ({linea.descripcion}): la fila {desde} no es "
                        "una línea sumable; se ignora"
                    )
                    continue
                acumulado[destino.orden] = acumulado.get(destino.orden, 0) + factor
            else:
                for otra in estructura.lineas:
                    if desde <= otra.fila_origen <= hasta and otra.tipo in _SUMABLES:
                        acumulado[otra.orden] = acumulado.get(otra.orden, 0) + factor
        if linea.orden in acumulado:
            res.avisos.append(
                f"Fila {linea.fila_origen} ({linea.descripcion}): se incluye a sí misma; se quita"
            )
            del acumulado[linea.orden]
        componentes = sorted((o, s) for o, s in acumulado.items() if s != 0)
        if not componentes:
            res.avisos.append(f"Fila {linea.fila_origen} ({linea.descripcion}): regla vacía")
            continue
        res.reglas.append(
            ReglaLeida(
                orden=linea.orden,
                fila_origen=linea.fila_origen,
                componentes=[(o, 1 if s > 0 else -1) for o, s in componentes],
                formula=valor,
            )
        )
    if sin_formula:
        res.avisos.append(
            f"{len(sin_formula)} totales o subtotales sin fórmula en la columna {letra} "
            f"(quedan sin regla): {', '.join(sin_formula[:5])}"
        )
    return res

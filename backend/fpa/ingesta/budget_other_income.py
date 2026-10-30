"""Lector de la tabla interna *Other Income* del Budget de Mazatlán (`Budget MZT 26-27.xlsx`).

La hoja `Other Income` trae una fila por cuenta (de la fila 9 hasta `Total …`) y, en las columnas
M..X, el importe de JUL … JUN. Cada celda es una de dos cosas, y **la fórmula nunca se evalúa**: se
reconoce su forma y se leen los valores que Excel dejó guardados.

    ='EstAct 25-26 MZT'!D12*1.05   ->  base × factor  (la base es el Est-Act del año previo)
    10000                          ->  monto fijo

Cualquier otra fórmula se toma como monto fijo con su valor guardado y se avisa. Además se compara,
cuenta por cuenta, contra la hoja `Budget 26-27 MZT` (filas de la sección *Other Income*): si el
detalle y el Budget no coinciden se avisa, no se rechaza (p. ej. un valor fijo que el Budget pisa).
"""

from __future__ import annotations

import re
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HOJA = "Other Income"
HOJA_BUDGET = "Budget 26-27 MZT"
COL_PRIMER_MES = 12  # M
COL_CONCEPTO = 1  # B
MESES = ["JUL", "AUG", "SEP", "OCT", "NOV", "DEC", "JAN", "FEB", "MAR", "APR", "MAY", "JUN"]
_RE_BASE_POR_FACTOR = re.compile(
    r"^=\s*\(?\s*'?([^'!]+)'?!\$?([A-Z]{1,3})\$?(\d+)\s*\)?\s*\*\s*(\d+(?:\.\d+)?)\s*$"
)
_TOLERANCIA = Decimal("0.01")
_MICRO = Decimal("0.000001")
_MESES_BUDGET_COL = 3  # D en la hoja Budget


@dataclass(frozen=True)
class CeldaLeida:
    metodo: str  # factor | fijo
    importe: Decimal
    base: Decimal | None = None
    factor: Decimal | None = None
    monto_fijo: Decimal | None = None


@dataclass
class FilaOtherIncome:
    fila: int  # número de fila en el Excel
    concepto: str
    celdas: list[CeldaLeida]  # 12, julio a junio

    @property
    def total(self) -> Decimal:
        return sum((c.importe for c in self.celdas), Decimal(0))


@dataclass
class OtherIncomeLeido:
    filas: list[FilaOtherIncome] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    # Bloque inferior de la hoja (12 valores, julio a junio; `None` si el archivo no lo trae):
    inflacion: list[Decimal] | None = None  # fila «Mthly», como fracción (0.05 = 5 %)
    est_act: list[Decimal] | None = None  # fila «Est Act …» (total de la hoja del año previo)
    budget_previo: list[Decimal] | None = None  # fila «Budget …» (Budget del año previo)
    # Cuentas que el archivo trae sin importes válidos (#REF! o vacías): ya se avisaron una vez
    sin_importes: list[str] = field(default_factory=list)
    # Secciones con cuentas por PARTIDAS (Mgmt & Check In): las hojas de detalle leídas
    # (`budget_partidas.PartidasLeidas`) y, por cuenta, el Est-Act y Budget previo del año previo
    partidas: list[Any] = field(default_factory=list)
    referencias: dict[str, tuple[list[Decimal | None], list[Decimal | None]]] = field(
        default_factory=dict
    )


def normalizar(texto: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(texto or "").lower())


def _dec(valor: object) -> Decimal | None:
    if valor is None or isinstance(valor, bool) or valor == "":
        return None
    try:
        # Excel guarda flotantes (61427.310000000005): se redondea a los 6 decimales de `Medida`.
        return Decimal(str(valor)).quantize(_MICRO)
    except InvalidOperation:
        return None


BaseDe = Callable[[str, str], object]  # (hoja, "D12") -> valor guardado


def leer_other_income_filas(
    formulas: Sequence[Sequence[object]],
    valores: Sequence[Sequence[object]],
    base_de: BaseDe,
    budget: Sequence[Sequence[object]] | None = None,
) -> OtherIncomeLeido:
    """Lee las filas ya cargadas de `Other Income` (fórmulas y valores guardados, misma forma)."""
    res = OtherIncomeLeido()
    encabezado = next(
        (
            i
            for i, f in enumerate(formulas[:20])
            if len(f) > COL_PRIMER_MES + 11
            and [str(c or "").strip().upper() for c in f[COL_PRIMER_MES : COL_PRIMER_MES + 12]]
            == MESES
        ),
        None,
    )
    if encabezado is None:
        raise ValueError("No se encontró el encabezado Jul … Jun en las columnas M a X")
    for n in range(encabezado + 1, len(formulas)):
        celdas_f = list(formulas[n]) + [None] * (COL_PRIMER_MES + 12 - len(formulas[n]))
        celdas_v = list(valores[n]) + [None] * (COL_PRIMER_MES + 12 - len(valores[n]))
        concepto = str(celdas_f[COL_CONCEPTO] or "").strip()
        if not concepto:
            continue
        if normalizar(concepto).startswith("total"):
            _leer_bloque_inferior(res, valores, n + 1)
            break
        meses: list[CeldaLeida] = []
        for i in range(12):
            f = celdas_f[COL_PRIMER_MES + i]
            v = _dec(celdas_v[COL_PRIMER_MES + i])
            meses.append(_celda(f, v, base_de, res, n + 1, concepto, MESES[i]))
        res.filas.append(FilaOtherIncome(n + 1, concepto, meses))
    if not res.filas:
        res.avisos.append("La hoja no trae cuentas de Other Income")
    if budget is not None:
        _comparar_con_budget(res, budget)
    return res


COL_ETIQUETA_BLOQUE = 9  # J: «Mthly», «Cum», «Total», «Est Act …», «Budget …», «Variance»


def _leer_bloque_inferior(
    res: OtherIncomeLeido, valores: Sequence[Sequence[object]], desde: int
) -> None:
    """Lee las filas «Mthly», «Est Act» y «Budget» (etiqueta en la columna J, meses en M..X)."""
    for fila in valores[desde:]:
        celdas = list(fila) + [None] * (COL_PRIMER_MES + 12 - len(fila))
        etiqueta = normalizar(celdas[COL_ETIQUETA_BLOQUE])
        meses = [_dec(celdas[COL_PRIMER_MES + i]) for i in range(12)]
        if etiqueta == "mthly":
            res.inflacion = [m or Decimal(0) for m in meses]
        elif etiqueta.startswith("estact") and any(m is not None for m in meses):
            res.est_act = [m or Decimal(0) for m in meses]
        elif etiqueta.startswith("budget") and any(m is not None for m in meses):
            res.budget_previo = [m or Decimal(0) for m in meses]


def _celda(
    formula: object,
    valor: Decimal | None,
    base_de: BaseDe,
    res: OtherIncomeLeido,
    fila: int,
    concepto: str,
    mes: str,
) -> CeldaLeida:
    if isinstance(formula, str) and formula.startswith("="):
        m = _RE_BASE_POR_FACTOR.match(formula)
        if m is not None:
            hoja, col, num, factor = m.groups()
            base = _dec(base_de(hoja, f"{col}{num}")) or Decimal(0)
            importe = valor if valor is not None else base * Decimal(factor)
            return CeldaLeida("factor", importe, base=base, factor=Decimal(factor))
        res.avisos.append(
            f"Fila {fila} ({concepto[:30]}), {mes}: fórmula no reconocida «{formula[:40]}»; "
            "se toma su valor como monto fijo"
        )
        return CeldaLeida("fijo", valor or Decimal(0), monto_fijo=valor or Decimal(0))
    literal = _dec(formula)
    monto = literal if literal is not None else (valor or Decimal(0))
    return CeldaLeida("fijo", monto, monto_fijo=monto)


def _comparar_con_budget(res: OtherIncomeLeido, budget: Sequence[Sequence[object]]) -> None:
    """Aviso por cada cuenta cuyo detalle no suma lo mismo que su fila de la hoja Budget."""
    por_concepto: dict[str, list[Decimal]] = {}
    dentro = False
    for fila in budget:
        etiqueta = normalizar(fila[1] if len(fila) > 1 else None)
        descripcion = normalizar(fila[2] if len(fila) > 2 else None)
        if etiqueta == "otherincome":
            dentro = True
            continue
        if not dentro:
            continue
        if descripcion.startswith("total"):
            break
        if not descripcion:
            continue
        meses = [
            _dec(fila[_MESES_BUDGET_COL + i] if len(fila) > _MESES_BUDGET_COL + i else None)
            or Decimal(0)
            for i in range(12)
        ]
        por_concepto.setdefault(descripcion, meses)
    for f in res.filas:
        esperado = por_concepto.get(normalizar(f.concepto))
        if esperado is None:
            res.avisos.append(f"«{f.concepto}» no aparece en la hoja Budget; no se pudo cuadrar")
            continue
        total_budget = sum(esperado, Decimal(0))
        if any(abs(c.importe - e) > _TOLERANCIA for c, e in zip(f.celdas, esperado, strict=True)):
            res.avisos.append(
                f"«{f.concepto}»: el detalle suma {f.total:,.2f} y el Budget {total_budget:,.2f} "
                "(el Excel del Budget usa otro valor; se importa el detalle)"
            )


def leer_other_income(ruta: str | Path) -> OtherIncomeLeido:
    """Lee `Other Income` de un libro `.xlsx` (sin evaluar fórmulas)."""
    with Path(ruta).open("rb") as archivo:
        wf = load_workbook(archivo, read_only=True, data_only=False)
        archivo.seek(0)
        wv = load_workbook(archivo, read_only=True, data_only=True)
        try:
            if HOJA not in wf.sheetnames:
                raise ValueError(f'El libro no tiene la hoja "{HOJA}"')
            formulas = [
                tuple(r) for r in wf[HOJA].iter_rows(max_col=COL_PRIMER_MES + 12, values_only=True)
            ]
            valores = [
                tuple(r) for r in wv[HOJA].iter_rows(max_col=COL_PRIMER_MES + 12, values_only=True)
            ]
            budget = (
                [tuple(r) for r in wv[HOJA_BUDGET].iter_rows(max_col=16, values_only=True)]
                if HOJA_BUDGET in wv.sheetnames
                else None
            )
            cache: dict[str, dict[str, object]] = {}

            def base_de(hoja: str, coord: str) -> object:
                if hoja not in cache:
                    cache[hoja] = {}
                    if hoja in wv.sheetnames:
                        for fila in wv[hoja].iter_rows(max_row=60):
                            for c in fila:
                                if c.value is not None and hasattr(c, "coordinate"):
                                    cache[hoja][c.coordinate] = c.value
                return cache[hoja].get(coord)

            return leer_other_income_filas(formulas, valores, base_de, budget)
        finally:
            wf.close()
            wv.close()


__all__ = [
    "HOJA",
    "CeldaLeida",
    "FilaOtherIncome",
    "OtherIncomeLeido",
    "leer_other_income",
    "leer_other_income_filas",
    "normalizar",
]

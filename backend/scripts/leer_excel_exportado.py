"""Lee el Excel exportado del Variance y lo imprime como JSON (lo usa la prueba de navegador).

Uso:  python scripts/leer_excel_exportado.py <archivo.xlsx>
Salida: [{"descripcion": "...", "cifras": ["1000.00", null, ...]}] con las 12 cifras de cada fila
de datos, a centavos (o a 6 decimales los porcentajes), y `null` donde la celda está en blanco.

El libro trae fórmulas vivas y aquí no hay Excel: se recalculan con el evaluador mínimo de las
pruebas (`tests/evaluador_formulas.py`), que entiende justo lo que el exportador escribe.
"""

import json
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))
from evaluador_formulas import Evaluador


def main(ruta: str) -> None:
    libro = load_workbook(ruta)
    hoja = libro["Variance"]
    evaluador = Evaluador(libro)
    salida = []
    for fila in hoja.iter_rows(min_row=6):
        if fila[1].value is None:
            continue
        cifras = []
        for i in (2, 3, 4, 5, 7, 8, 9, 10, 12, 13, 14, 15):
            celda = fila[i]
            if celda.value is None:
                cifras.append(None)
                continue
            valor = (
                evaluador.numero("Variance", celda.column, celda.row)
                if celda.data_type == "f"
                else Decimal(str(celda.value))
            )
            paso = Decimal("0.000001") if i in (5, 10, 15) else Decimal("0.01")
            cifras.append(str(valor.quantize(paso, rounding=ROUND_HALF_UP)))
        salida.append({"descripcion": fila[1].value, "cifras": cifras})
    sys.stdout.write(json.dumps(salida, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1])

"""Crea los Excel SINTÉTICOS que usan las pruebas de navegador (nunca datos reales de Torres).

Uso:  python scripts/crear_archivos_e2e.py <carpeta_destino>
"""

import io
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from xls_sintetico import a_xls, catalogo_basico

RAIZ_REPO = Path(__file__).resolve().parents[2]

ENCABEZADO_VARIANCE = [
    None, "Mes", "Fecha", "No Cuenta", "Nombre de la Cuenta", "Asiento - Salida",
    "Concepto - Nombre Articulo", "Cantidad", "Costo", "Débito", "Crédito",
]  # fmt: skip


def variance(filas: list[tuple]) -> bytes:  # type: ignore[type-arg]
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Base de Datos"
    ws.append([])
    ws.append(ENCABEZADO_VARIANCE)
    for f in filas:
        ws.append(list(f))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fila(
    mes: str, dia: datetime, cuenta: str, nombre: str, debito: float | None, credito: float | None
):  # type: ignore[no-untyped-def]
    return (None, mes, dia, cuenta, nombre, "8.3.1", "MOVIMIENTO", None, None, debito, credito)


def con130(movimientos: list[tuple[str, str, str, float | None, float | None]]) -> bytes:
    """CON130 sintético de Novohit (Mayor Analítico) con el mismo layout que el real.

    `movimientos`: (cuenta con su nombre, fecha `dd/mm/aaaa`, concepto, débito, crédito). Cada
    cuenta lleva su "Saldo anterior" y su fila "Total"; el reporte cuadra (Débito = Crédito) si los
    movimientos cuadran. Nunca contiene datos reales de Torres.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "CON130"
    ws.append([None] * 6 + ["Torres Sintética S.A. de C.V."])
    ws.append([None] * 6 + ["Mayor Analítico (movimientos por cuentas) "])
    ws.append([None] * 6 + ["Periodo Agosto 2026 a Agosto 2026"])
    ws.append(
        [None] * 6
        + ["# ASIENTO", "Fecha", "Tp.Asiento", "Oper.", "Tp.Doc.", "Número", "Fecha", "Concepto"]
        + ["Débito", "Crédito", "Saldo"]
    )
    por_cuenta: dict[str, list[tuple[str, str, float | None, float | None]]] = {}
    for cuenta, fecha, concepto, debito, credito in movimientos:
        por_cuenta.setdefault(cuenta, []).append((fecha, concepto, debito, credito))
    for cuenta, filas in sorted(por_cuenta.items()):
        ws.append([None] * 6 + [cuenta] + [None] * 6 + ["Saldo anterior", 0, 0])
        for n, (fecha, concepto, debito, credito) in enumerate(filas, start=1):
            ws.append(
                [None] * 6
                + [f"8.1.{n}", fecha, "POLIZA", "REGISTRO", "POLIZA", f"P{n}", fecha, concepto]
                + [debito, credito]
            )
        ws.append(
            [None] * 6
            + [None] * 7
            + [f"Total  {cuenta}", sum(f[2] or 0 for f in filas), sum(f[3] or 0 for f in filas)]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_variance() -> bytes:
    """Variance sintético COMPLETO (hojas `Ago` y `Base de Datos`) con cifras fáciles de verificar.

    Agosto: Room Revenue 1,000.00 (Budget 900), Food Revenue 500.50 (600), Salaries 400.00 (380),
    Supplies 150.25 (140). Julio: Room Revenue 700.00 y Salaries 380.00. Los conceptos incluyen
    HTML a propósito: la pantalla debe mostrarlo como texto, nunca interpretarlo.
    """
    wb = Workbook()
    ago = wb.active
    assert ago is not None
    ago.title = "Ago"

    def poner(n: int, b=None, c=None, d=None, e=None, j=None, o=None, s=None):  # type: ignore[no-untyped-def]
        for col, valor in ((2, b), (3, c), (4, d), (5, e), (10, j), (15, o), (19, s)):
            if valor is not None:
                ago.cell(row=n, column=col, value=valor)

    poner(2, "Account #", "Account Description", "AUGUST")
    poner(3, None, None, "Actual")
    poner(4, "R e v e n u e")
    poner(5, "029-0000-1-000", "Room Revenue", 1000.0, 900.0, 1700.0, 10800.0, "=-'4001002.001'!L5")
    poner(6, "029-0000-2-000", "Food Revenue", 500.5, 600.0, 1200.0, 7200.0, "=-'4001005.001'!L5")
    poner(7, None, "T o t a l   R e v e n u e", 1500.5, 1500.0, 2900.0, 18000.0, "=SUM(S5:S6)")
    poner(9, "E x p e n s e s")
    poner(10, "029-0111-60010-000", "Salaries", 400.0, 380.0, 760.0, 4560.0, "='6009001.001'!L9")
    poner(11, "029-0000-3-000", "Supplies", 150.25, 140.0, 280.0, 1680.0, "='6005024'!L3")
    poner(12, None, "T o t a l   E x p e n s e s", 550.25, 520.0, 1040.0, 6240.0, "=SUM(S10:S11)")
    # una línea sin movimientos ni Budget: la pantalla la oculta por omisión y avisa
    poner(13, "029-0000-9-000", "Unused Expense", 0.0, None, None, None, "='6999999'!L2")
    poner(14, None, "N e t   P r o f i t", 950.25, 980.0, 1860.0, 11760.0, "=S7-S12")
    # tipo de cambio: sin cuentas ni fórmula; su Actual se captura a mano (Budget fijo de 18)
    poner(16, None, "Exchange Rate:", None, 18.0, 18.0, 18.0)

    bd = wb.create_sheet("Base de Datos")
    bd.append([])
    bd.append(ENCABEZADO_VARIANCE)
    mov = [
        (
            "Julio",
            datetime(2026, 7, 10),
            "4001002.001",
            "ROOM",
            "7.1.1",
            "Hospedaje julio",
            None,
            700.0,
        ),
        (
            "Julio",
            datetime(2026, 7, 11),
            "6009001.001",
            "SUELDOS",
            "7.1.2",
            "Nómina julio",
            380.0,
            None,
        ),
        (
            "Agosto",
            datetime(2026, 8, 2),
            "4001002.001",
            "ROOM",
            "8.1.1",
            "Hospedaje <b>agosto</b> & más",
            None,
            600.0,
        ),
        (
            "Agosto",
            datetime(2026, 8, 3),
            "4001002.001",
            "ROOM",
            "8.1.2",
            "Hospedaje agosto 2",
            None,
            400.0,
        ),
        ("Agosto", datetime(2026, 8, 4), "4001005.001", "FOOD", "8.1.3", "Alimentos", None, 500.5),
        (
            "Agosto",
            datetime(2026, 8, 5),
            "6009001.001",
            "SUELDOS",
            "8.1.4",
            "Quincena 1",
            150.0,
            None,
        ),
        (
            "Agosto",
            datetime(2026, 8, 6),
            "6009001.001",
            "SUELDOS",
            "8.1.5",
            "Quincena 2",
            150.0,
            None,
        ),
        ("Agosto", datetime(2026, 8, 7), "6009001.001", "SUELDOS", "8.1.6", "Bono", 100.0, None),
        (
            "Agosto",
            datetime(2026, 8, 8),
            "6005024",
            "INSUMOS",
            "8.1.7",
            "Compra de insumos",
            150.25,
            None,
        ),
    ]
    for mes, dia, cuenta, nombre, asiento, concepto, debito, credito in mov:
        bd.append([None, mes, dia, cuenta, nombre, asiento, concepto, None, None, debito, credito])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def golden_real(destino: Path) -> None:
    """Con el Excel real de Mazatlán, guarda cifras de referencia para el e2e (NO se versiona).

    Elige líneas con descripción única que no estén entre las excepciones documentadas de la prueba
    dorada y guarda sus 12 cifras con el orden de la pantalla: mes de agosto tal como lo pegó
    Contraloría, acumulado = julio + agosto (Actual y Budget) y el bloque anual vacío (el Excel de
    Mazatlán no trae presupuesto anual).
    """
    import json
    import os
    import re
    from decimal import Decimal

    from fpa.ingesta.variance_estructura import TipoLinea
    from fpa.ingesta.variance_lectura import leer_variance
    from fpa.ingesta.variance_mzt import leer_hojas_mzt, leer_valores_mensuales_mzt

    carpeta = Path(
        os.environ.get(
            "FPA_DATOS_REALES",
            Path.home() / "Desktop" / "PLATAFORMA FP&A E INTELIGENCIA HOTELERA",
        )
    )
    nombre = "02. MZT August 2026 Variance (Envio).xlsx"
    salida = destino / "golden_real.json"
    excel = next((p for p in (carpeta / nombre,) if p.exists()), None)
    if excel is None:
        salida.unlink(missing_ok=True)
        return
    from test_variance_dorada_mzt import EXCEPCIONES, EXCEPCIONES_VAR

    estructura = leer_variance(excel).estructura
    por_periodo = leer_valores_mensuales_mzt(leer_hojas_mzt(excel)).por_periodo
    ago, jul = por_periodo["2026-08"], por_periodo["2026-07"]
    descripciones = [x.descripcion for x in estructura.lineas if x.tipo is not TipoLinea.SECCION]
    elegidas = []
    for linea in estructura.lineas:
        a, j = ago.get(linea.fila_origen), jul.get(linea.fila_origen)
        if (
            a is None
            or j is None
            or linea.tipo in (TipoLinea.SECCION, TipoLinea.INFORMATIVA)
            or linea.fila_origen in EXCEPCIONES
            or linea.fila_origen in EXCEPCIONES_VAR
            or descripciones.count(linea.descripcion) != 1
            or not re.search(r"[A-Za-z]{4}", linea.descripcion)
            or None in (a.actual, a.budget, a.var, a.pct, j.actual, j.budget)
            or abs(a.actual or 0) < 100  # con ceros no se prueba nada
            or abs(a.budget or 0) < 100
        ):
            continue
        elegidas.append((linea, a, j))
    paso = max(1, len(elegidas) // 14)
    muestra = elegidas[::paso][:14]

    def texto(valor: Decimal | None) -> str | None:
        return None if valor is None else str(valor)

    filas = []
    for linea, a, j in muestra:
        assert a.actual is not None and a.budget is not None
        assert j.actual is not None and j.budget is not None
        filas.append(
            {
                "fila": linea.fila_origen,
                "descripcion": linea.descripcion,
                "tipo": linea.tipo.value,
                "cifras": [
                    texto(a.actual),
                    texto(a.budget),
                    texto(a.var),
                    texto(a.pct),
                    texto(a.actual + j.actual),
                    texto(a.budget + j.budget),
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                ],
            }
        )
    salida.write_text(json.dumps(filas, ensure_ascii=False, indent=1), encoding="utf-8")


def reporte_other_income() -> bytes:
    """Reporte mínimo con la sección *Other Income* (hoja `Ago`): 3 cuentas de detalle y su total.

    Sirve para probar la tabla interna de cada cuenta (Budget por cuenta, parte 1). Reemplaza la
    estructura del reporte: las demás pruebas cargan la suya con `prepararReporteSintetico`.
    """
    wb = Workbook()
    ago = wb.active
    assert ago is not None
    ago.title = "Ago"

    def poner(n: int, b=None, c=None, d=None, e=None, s=None):  # type: ignore[no-untyped-def]
        for col, valor in ((2, b), (3, c), (4, d), (5, e), (19, s)):
            if valor is not None:
                ago.cell(row=n, column=col, value=valor)

    poner(2, "Account #", "Account Description", "AUGUST")
    poner(3, None, None, "Actual")
    poner(4, "O t h e r   I n c o m e")
    poner(5, "028-0000-43080-000", "Rental Income", 100.0, 90.0, "=-'4001002.001'!L5")
    poner(6, "028-0101-45010-000", "Laundry Coins", 20.0, 25.0, "=-'4001005.001'!L5")
    poner(7, "028-0101-45010-000", "Beauty Services", 30.0, 30.0, "=-'4001006.001'!L5")
    poner(8, None, "T o t a l   O t h e r   I n c o m e", 150.0, 145.0, "=SUM(S5:S7)")
    # Mgmt & Check In (cuentas por partidas): Other Fringe y Telephone con hoja de detalle;
    # Insurance sin ella
    poner(10, "M g m t   &   C h e c k   I n")
    poner(11, "028-0111-60055-000", "Other Fringe", 10.0, 10.0, "=-'6009010.004'!L5")
    poner(12, "028-0111-61560-000", "Telephone", 5.0, 5.0, "=-'6009015.001'!L5")
    poner(13, "028-0111-60655-000", "Insurance", 2.0, 2.0, "=-'6009030.001'!L5")
    poner(14, None, "T o t a l   M g m t   &   C h e c k   I n", 17.0, 17.0, "=SUM(S11:S13)")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def budget_other_income() -> bytes:
    """Excel de Budget mínimo: solo la hoja `Other Income`, con montos fijos (el lector real también
    entiende `=EstAct!D12*1.05`, pero openpyxl no guarda valores calculados de fórmulas)."""
    meses = ["Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Other Income"
    ws.append([None, "Item"] + [None] * 10 + meses)
    ws.append([None, "Rental Income"] + [None] * 10 + [1000.0] * 12)
    ws.append([None, "Laundry Coins"] + [None] * 10 + [200.0] * 12)
    ws.append([None, "Beauty Services"] + [None] * 10 + [300.0] * 12)
    ws.append([None, "Total Expense"])
    # Bloque inferior de la hoja: inflación mensual, Est Act y Budget del año previo
    ws.append([None] * 9 + ["Mthly", None, None] + [0.01] * 12)
    ws.append([None] * 9 + ["Est Act 26-27", None, None] + [2000.0] * 12)
    ws.append([None] * 9 + ["Budget 26-27", None, None] + [1500.0] * 12)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def budget_mgmt() -> bytes:
    """Excel de Budget mínimo de *Mgmt & Check In*: la hoja `Budget 26-27 MZT` (sección con
    importes),
    las hojas del año previo y dos hojas de detalle por cuenta (`0111-60055 Other Fringe`, partidas
    mensual y estacional con 5 % de inflación; `0111-61560 Telephone`, sin inflación)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Budget 26-27 MZT"
    ws.append([None, "Account #", "Account Description"])
    ws.append([None, "M g m t   &   C h e c k   I n"])
    ws.append([None, "028-0111-60055-000", "Other Fringe", *([1050.0] * 12)])
    ws.append([None, "028-0111-61560-000", "Telephone", *([4515.0] * 12)])
    ws.append([None, "028-0111-60655-000", "Insurance", *([500.0] * 12)])
    ws.append([None, None, "T o t a l   M g m t   &   C h e c k   I n", *(["#REF!"] * 12)])
    for titulo, base in (("EstAct 25-26 MZT", 1200.0), ("MZT Budget 25-26", 1100.0)):
        hoja = wb.create_sheet(titulo)
        hoja.append([None, "M g m t   &   C h e c k   I n"])
        hoja.append([None, "028-0111-60055-000", "Other Fringe", *([base] * 12)])
        hoja.append([None, "028-0111-61560-000", "Telephone", *([base * 4] * 12)])
        hoja.append([None, None, "T o t a l   M g m t   &   C h e c k   I n", *([base * 5] * 12)])

    def detalle(titulo: str, partidas: list[tuple], inflacion: float) -> None:  # type: ignore[type-arg]
        hoja = wb.create_sheet(titulo)
        for _ in range(7):
            hoja.append([None])
        encabezado = [None] * 26
        encabezado[1], encabezado[3], encabezado[4], encabezado[5], encabezado[6] = (
            "Item", "Qty of Items", "Unit of Measure", "Cost", "Extended Cost",
        )  # fmt: skip
        hoja.append(encabezado)
        total = 0.0
        for item, qty, unidad, costo, marca, meses in partidas:
            fila: list = [None] * 26
            fila[1], fila[3], fila[4], fila[5], fila[6] = item, qty, unidad, costo, qty * costo
            fila[7 if marca == "mensual" else 11] = "X"
            for i, m in enumerate(meses):
                fila[12 + i] = m
            total += qty * costo
            hoja.append(fila)
        fin: list = [None] * 26
        fin[1], fin[6] = "Total Expense", total
        hoja.append(fin)
        infl: list = [None] * 26
        infl[9] = "Mthly"
        for i in range(12):
            infl[12 + i] = inflacion
        hoja.append(infl)

    estacional = [0.0] * 12
    estacional[5] = 12000.0  # diciembre
    detalle(
        "0111-60055 Other Fringe",
        [
            ("PARTY ROOM", 1, "SVC", 12000.0, "estacional", estacional),
            ("EMPLOYEE OF THE MONTH PRIZE", 12, "mos", 100.0, "mensual", [100.0] * 12),
        ],
        0.05,
    )
    detalle(
        "0111-61560 Telephone",
        [("LOCAL CALLS", 12, "ANNUAL", 4515.0, "mensual", [4515.0] * 12)],
        0.0,
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main(destino: Path) -> None:
    destino.mkdir(parents=True, exist_ok=True)
    agosto = [
        fila("Agosto", datetime(2026, 8, 5), "4001002.001", "INGRESOS DE BAR", None, 100.50),
        fila("Agosto", datetime(2026, 8, 6), "6009001.001", "SUELDOS", 40.25, None),
    ]
    corregido = [
        fila("Agosto", datetime(2026, 8, 5), "4001002.001", "INGRESOS DE BAR", None, 120.00),
        fila("Agosto", datetime(2026, 8, 6), "6009001.001", "SUELDOS", 40.25, None),
    ]
    (destino / "variance_reporte.xlsx").write_bytes(reporte_variance())
    (destino / "reporte_other_income.xlsx").write_bytes(reporte_other_income())
    (destino / "budget_other_income.xlsx").write_bytes(budget_other_income())
    (destino / "budget_mgmt.xlsx").write_bytes(budget_mgmt())
    (destino / "variance_agosto.xlsx").write_bytes(variance(agosto))
    (destino / "variance_agosto_corregido.xlsx").write_bytes(variance(corregido))

    # Los mismos movimientos como CON130 de Novohit (con su contrapartida en bancos para que cuadre)
    def movimientos_con130(
        ingreso: float,
    ) -> list[tuple[str, str, str, float | None, float | None]]:
        return [
            ("4001002.001  INGRESOS DE BAR", "05/08/2026", "Venta de bar", None, ingreso),
            ("6009001.001  SUELDOS", "06/08/2026", "Nomina", 40.25, None),
            ("1000001.001  BANCOS", "05/08/2026", "Cobro bar", ingreso, None),
            ("1000001.001  BANCOS", "06/08/2026", "Pago nomina", None, 40.25),
        ]

    (destino / "con130_agosto.xlsx").write_bytes(con130(movimientos_con130(100.50)))
    # El mismo CON130 como .xls (el formato que exporta Novohit) y como .csv, para la subida
    from con130_formatos import con130_csv, con130_xls

    (destino / "con130_novohit.xls").write_bytes(con130_xls(movimientos_con130(100.50)))
    (destino / "con130_novohit.csv").write_bytes(con130_csv(movimientos_con130(100.50)))
    (destino / "con130_agosto_corregido.xlsx").write_bytes(con130(movimientos_con130(120.00)))
    # Catálogo contable (CON020) sintético en .xls, y un falso .xls de texto
    (destino / "catalogo.xls").write_bytes(a_xls(catalogo_basico()))
    (destino / "falso.xls").write_bytes(b"esto no es un libro de Excel")
    # Falso .xlsx: texto plano con extensión de Excel
    (destino / "falso.xlsx").write_bytes(b"esto no es un libro de Excel")
    # Zip que no es un libro de Excel
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hola.txt", "hola")
    (destino / "zip_no_excel.xlsx").write_bytes(buf.getvalue())
    golden_real(destino)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
